#!/usr/bin/env python3
"""generer_empreinte_carbone.py -- activité de SES « Émissions françaises et empreinte carbone : pourquoi deux totaux ? »
(/enseignants/environnement/).

Jeu : émissions de gaz à effet de serre des unités résidentes françaises et empreinte carbone de la France, décomposées.
Source : Insee Première n° 2077 (16/10/2025), fichier de données ARCHIVÉ dans scripts/sources_enseignants/ (empreinte
dans SHA256SUMS) : figure 1 (décomposition de l'année), figure 3 (empreinte et composantes depuis 1990), tableau
complémentaire 2 (émissions et empreinte depuis 1990). Coproduction Insee-SDES. Aucun appel réseau ici.

⚠ SÉRIE RÉVISÉE À CHAQUE ÉDITION (arbitrage ENTRANTE_2026-10-02_Enseignants_Theme3_Exemplaires, T1) : l'empreinte 2023
valait 644 Mt dans l'édition de 2024 et 583 Mt dans celle-ci ; la dernière année est provisoire. On ne mélange jamais
deux éditions : tout vient d'UN fichier.

METTRE À JOUR (édition annuelle, mi-octobre) : archiver le fichier de la nouvelle Insee Première, ajouter son
empreinte, changer FICHIER et EDITION, relancer. Les gardes disent si les constats de la fiche tiennent.

TÉMOINS, bloquants : chaque total se retrouve depuis ses composantes (à l'arrondi) ; la dernière ligne de la série
longue égale la décomposition de l'année ; les deux tableaux de la série longue s'accordent.
Chaque qualificatif de la fiche et du corrigé est une garde. SVG et PNG ensemble ou pas du tout.

Usage : python scripts/generer_empreinte_carbone.py [--check]
Sorties : data/ et static/empreinte_carbone.json (bloc affichage), static/empreinte_carbone.csv,
          static/img/empreinte-carbone-deux-totaux.svg + .png
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import sys
import warnings
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generer_parcours_licence as gl  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "scripts" / "sources_enseignants"
OUT_DATA = ROOT / "data" / "empreinte_carbone.json"
OUT_STATIC = ROOT / "static" / "empreinte_carbone.json"
OUT_CSV = ROOT / "static" / "empreinte_carbone.csv"
OUT_IMG = ROOT / "static" / "img"
FIGURE = "empreinte-carbone-deux-totaux"
PAGE_URL = "stephane-lalut.com/enseignants/environnement/"
FICHIER = "insee_premiere_2077_empreinte_carbone_2024.xlsx"
EDITION = "Insee Première n° 2077, octobre 2025"
fail, log, fr, esc = gl.fail, gl.log, gl.fr, gl.esc

L1 = {"menages": "Émissions directes des ménages",
      "exportee": "Émissions de la production intérieure exportées",
      "interieure": "Émissions de la production intérieure destinées à la demande finale",
      "importee": "Émissions importées destinées à la demande finale",
      "emissions": "Ensemble émissions françaises", "empreinte": "Ensemble empreinte carbone"}
COULEUR = {"menages": (gl.GRIS_CLAIR, gl.INK), "interieure": (gl.BLEU, "#ffffff"), "exportee": (gl.GRIS, "#ffffff"),
           "importee": (gl.ORANGE, gl.INK)}
NOM = {"menages": ("Émissions directes", "des ménages"), "interieure": ("Production en France", "pour la demande française"),
       "exportee": ("Production en France", "exportée"), "importee": ("Importations", "pour la demande française")}


def lire(sums: dict[str, str]) -> dict:
    import openpyxl
    p = SRC / FICHIER
    if not p.is_file():
        fail("fichier absent : %s" % FICHIER)
    if sums.get(FICHIER) != hashlib.sha256(p.read_bytes()).hexdigest():
        fail("%s : empreinte differente de SHA256SUMS (piece remplacee ?)" % FICHIER)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(p, data_only=True)
    for f in ("Figure 1", "Figure 3", "Tableau complémentaire 2"):
        if f not in wb.sheetnames:
            fail("%s : feuille %r absente" % (FICHIER, f))

    def lignes(nom):
        return [list(r) for r in wb[nom].iter_rows(values_only=True) if any(c is not None for c in r)]

    f1 = lignes("Figure 1")
    titre = str(f1[0][0])
    annee = int(titre.strip().split()[-1])
    if f1[1][:4] != ["Type d'émissions", "Émissions françaises", "Émissions importées", "Empreinte carbone"]:
        fail("figure 1 : en-tetes changes : %s" % f1[1][:4])
    par = {str(r[0]): r for r in f1[3:] if isinstance(r[0], str)}
    manq = [v for v in L1.values() if v not in par]
    if manq:
        fail("figure 1 : lignes absentes ou renommees : %s" % " | ".join(manq))

    def val(cle, col):
        v = par[L1[cle]][col]
        if not isinstance(v, (int, float)):
            fail("figure 1, %s : valeur manquante (%r)" % (L1[cle], v))    # « /// » n'est pas zéro
        return float(v)

    d = {"annee": annee, "menages": val("menages", 1), "exportee": val("exportee", 1), "interieure": val("interieure", 1),
         "importee": val("importee", 3), "emissions": val("emissions", 1), "empreinte": val("empreinte", 3)}
    # témoin 1 : les totaux se retrouvent depuis les composantes (valeurs arrondies au Mt)
    if abs(d["menages"] + d["interieure"] + d["exportee"] - d["emissions"]) > 1.01:
        fail("figure 1 : menages + interieure + exportee != emissions francaises")
    if abs(d["menages"] + d["interieure"] + d["importee"] - d["empreinte"]) > 1.01:
        fail("figure 1 : menages + interieure + importee != empreinte")
    f3 = lignes("Figure 3")
    if f3[2] != ["Années", "Émissions directes des ménages", "Émissions intérieures", "Émissions importées", "Empreinte totale"]:
        fail("figure 3 : en-tetes changes : %s" % f3[2])
    serie = [dict(annee=int(r[0]), menages=float(r[1]), interieure=float(r[2]), importee=float(r[3]), empreinte=float(r[4]))
             for r in f3[3:] if isinstance(r[0], (int, float))]
    der = serie[-1]
    # témoin 2 : la série longue se termine sur la décomposition de l'année
    if der["annee"] != annee or any(der[k] != d[k] for k in ("menages", "interieure", "importee", "empreinte")):
        fail("figure 3 : la derniere annee ne correspond pas a la figure 1")
    t2 = lignes("Tableau complémentaire 2")
    em = {int(r[0]): dict(emissions=r[1] / 1000, empreinte=r[4] / 1000, emissions_hab=float(r[7]), empreinte_hab=float(r[10]))
          for r in t2[3:] if isinstance(r[0], (int, float))}
    # témoin 3 : les deux tableaux de la série longue s'accordent, année par année
    for s in serie:
        if s["annee"] not in em or abs(em[s["annee"]]["empreinte"] - s["empreinte"]) > 0.51:
            fail("%d : l'empreinte du tableau complementaire 2 differe de la figure 3" % s["annee"])
        s.update(emissions=round(em[s["annee"]]["emissions"], 1), emissions_hab=em[s["annee"]]["emissions_hab"],
                 empreinte_hab=em[s["annee"]]["empreinte_hab"])
    if abs(em[annee]["emissions"] - d["emissions"]) > 0.51:
        fail("tableau complementaire 2 : emissions %d differentes de la figure 1" % annee)
    return {"decomposition": d, "serie": serie}


def affichage(j: dict) -> tuple[dict, list[str]]:
    d, s = j["decomposition"], j["serie"]
    a0, a1 = s[0], s[-1]
    ent = lambda v: fr(v, 0)                                             # noqa: E731
    evol = lambda x, y: 100 * (y - x) / x                                # noqa: E731
    interieure0, interieure1 = a0["menages"] + a0["interieure"], a1["menages"] + a1["interieure"]
    e_em, e_emp = evol(a0["emissions"], a1["emissions"]), evol(a0["empreinte"], a1["empreinte"])
    e_imp, e_int = evol(a0["importee"], a1["importee"]), evol(interieure0, interieure1)
    A = {"annee": str(d["annee"]), "annee_depart": str(a0["annee"]), "edition": EDITION,
         "emissions": ent(d["emissions"]), "empreinte": ent(d["empreinte"]), "menages": ent(d["menages"]),
         "interieure": ent(d["interieure"]), "exportee": ent(d["exportee"]), "importee": ent(d["importee"]),
         "ecart_totaux": ent(d["empreinte"] - d["emissions"]), "part_importee": ent(100 * d["importee"] / d["empreinte"]),
         "commun": ent(d["menages"] + d["interieure"]),
         "emissions_hab": fr(a1["emissions_hab"]), "empreinte_hab": fr(a1["empreinte_hab"]),
         "evol_emissions": ent(abs(e_em)), "evol_empreinte": ent(abs(e_emp)), "evol_importee": ent(abs(e_imp)),
         "evol_interieure": ent(abs(e_int)),
         "emissions_depart": ent(a0["emissions"]), "empreinte_depart": ent(a0["empreinte"]), "importee_depart": ent(a0["importee"])}
    gardes = [
        ("l'empreinte dépasse les émissions françaises", d["empreinte"] > d["emissions"]),
        ("les émissions importées pour la demande française dépassent « nettement » la production exportée (plus du double)",
         d["importee"] > 2 * d["exportee"]),
        ("les importations font « la moitié » de l'empreinte (entre 45 et 55 %)", 45 < 100 * d["importee"] / d["empreinte"] < 55),
        ("l'écart entre les deux totaux n'est pas égal aux émissions importées, mais aux importations moins les exportations",
         abs((d["empreinte"] - d["emissions"]) - (d["importee"] - d["exportee"])) <= 1.01
         and abs((d["empreinte"] - d["emissions"]) - d["importee"]) > 50),
        ("depuis %d, les émissions françaises ont baissé plus vite que l'empreinte (au moins cinq points d'écart)" % a0["annee"],
         e_em < 0 and e_emp < 0 and e_emp - e_em > 5),
        ("les émissions importées sont « presque stables » depuis %d (moins de 5 %% de variation)" % a0["annee"], abs(e_imp) < 5),
        ("la composante intérieure de l'empreinte, ménages compris, a baissé « d'un tiers environ » (entre 28 et 40 %)", -40 < e_int < -28),
        ("les signes écrits dans le corrigé : émissions, empreinte et composante intérieure en baisse, importations en hausse",
         e_em < 0 and e_emp < 0 and e_int < 0 and e_imp > 0),
    ]
    return A, [nom for nom, ok in gardes if not ok]


def figure(j: dict, A: dict) -> str:
    """Deux barres horizontales à la même échelle : ce que chaque total additionne."""
    d = j["decomposition"]
    W = gl.W
    X0, LARG, TOP, HB, PAS = 168, 478, 78, 62, 92
    barres = [(("Émissions des unités", "résidentes françaises"), ["menages", "interieure", "exportee"], d["emissions"]),
              (("Empreinte carbone", "de la France"), ["menages", "interieure", "importee"], d["empreinte"])]
    NB_CART = 6
    H = TOP + PAS + HB + 16 + 12 * (NB_CART - 3)
    ech = LARG / d["empreinte"]
    titre = "Émissions françaises et empreinte carbone en %s : deux totaux, trois postes chacun" % A["annee"]
    desc = ("Deux barres horizontales à la même échelle, en millions de tonnes équivalent CO2. Émissions des unités résidentes "
            "françaises : %s, dont %s émises directement par les ménages, %s par la production en France destinée à la demande "
            "française et %s par la production en France exportée. Empreinte carbone : %s, dont les mêmes %s et %s, et %s "
            "d'émissions importées pour la demande française."
            % (A["emissions"], A["menages"], A["interieure"], A["exportee"], A["empreinte"], A["menages"], A["interieure"], A["importee"]))
    e = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 %d %d" role="img" font-variant-numeric="tabular-nums" '
         'aria-labelledby="ec-t ec-d" font-family="%s">' % (W, H + 52, gl.FONT),
         '<title id="ec-t">%s</title><desc id="ec-d">%s</desc>' % (esc(titre), esc(desc)),
         '<rect width="%d" height="%d" fill="#ffffff"/>' % (W, H + 52),
         '<text x="0" y="18" font-size="15" font-weight="600" fill="%s">%s</text>' % (gl.INK, esc(titre)),
         '<text x="0" y="38" font-size="11" fill="%s">%s</text>'
         % (gl.INK2, esc("Gaz à effet de serre, en millions de tonnes équivalent CO2 (Mt CO2 éq), France, %s." % A["annee"]))]
    for n, (lib, postes, total) in enumerate(barres):
        y = TOP + PAS * n
        for k, mot in enumerate(lib):
            e.append('<text x="%d" y="%.1f" font-size="11.5" fill="%s" text-anchor="end"%s>%s</text>'
                     % (X0 - 10, y + HB / 2 - 2 + 14 * k, gl.INK, ' font-weight="600"' if k == 0 else "", esc(mot)))
        x = float(X0)
        for cle in postes:
            w = d[cle] * ech
            fond, encre = COULEUR[cle]
            e.append('<rect x="%.1f" y="%d" width="%.1f" height="%d" fill="%s" stroke="#ffffff" stroke-width="1.5"/>' % (x, y, w, HB, fond))
            e.append('<text x="%.1f" y="%.1f" font-size="15" font-weight="600" fill="%s" text-anchor="middle">%s</text>'
                     % (x + w / 2, y + 22, encre, esc(fr(d[cle], 0))))
            for k, mot in enumerate(NOM[cle]):
                e.append('<text x="%.1f" y="%.1f" font-size="%s" fill="%s" text-anchor="middle">%s</text>'
                         % (x + w / 2, y + 37 + 11 * k, "9" if w > 110 else "8", encre, esc(mot)))
            x += w
        e.append('<text x="%.1f" y="%.1f" font-size="15" font-weight="600" fill="%s">= %s</text>' % (x + 8, y + HB / 2 + 5, gl.INK, esc(fr(total, 0))))
    y0 = H + 4 - 12 * (NB_CART - 3)
    e.append('<line x1="0" y1="%.1f" x2="%d" y2="%.1f" stroke="%s"/>' % (y0, W, y0, gl.GRID))
    cart = [("Insee et SDES, %s, figure 1 ; champ : France, CO2, méthane, protoxyde d'azote et gaz fluorés" % EDITION, gl.INK2),
            ("Émissions des unités résidentes : format des comptes d'émissions dans l'air, transport international des résidents compris ; "
             "ce n'est pas le total de l'inventaire national.", gl.INK2),
            ("Empreinte : estimation des émissions associées à la demande finale française, où qu'elles aient lieu ; hors émissions "
             "des produits exportés.", gl.INK2),
            ("Deux conventions de comptabilité : elles ne mesurent ni une délocalisation ni l'effet d'une politique.", gl.INK2),
            ("Estimations révisées à chaque édition ; %s est provisoire. Valeurs arrondies au million de tonnes." % A["annee"], gl.INK2),
            ("Figure Stéphane Lalut, CC BY 4.0 · " + PAGE_URL, gl.MUTED)]
    assert len(cart) == NB_CART
    for k, (txt, col) in enumerate(cart):
        e.append('<text x="0" y="%.1f" font-size="9" fill="%s">%s</text>' % (y0 + 13 + 12 * k, col, esc(txt)))
    e.append("</svg>")
    return "\n".join(e)


def csv_texte(j: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["annee", "emissions_unites_residentes_mt", "empreinte_mt", "dont_menages_mt", "dont_production_interieure_demande_francaise_mt",
                "dont_importations_demande_francaise_mt", "emissions_t_par_habitant", "empreinte_t_par_habitant"])
    for s in j["serie"]:
        w.writerow([s["annee"], "%.1f" % s["emissions"], "%.0f" % s["empreinte"], "%.0f" % s["menages"], "%.0f" % s["interieure"],
                    "%.0f" % s["importee"], "%.1f" % s["emissions_hab"], "%.1f" % s["empreinte_hab"]])
    return buf.getvalue()


def main() -> int:
    check = "--check" in sys.argv[1:]
    if not check:
        try:
            import cairosvg  # noqa: F401
        except ImportError:
            fail("cairosvg absent : SVG et PNG se produisent ensemble ou pas du tout (pip install cairosvg)")
    j = lire(gl.empreintes())
    A, faux = affichage(j)
    if faux:
        fail("la page affirme ce que les donnees ne soutiennent plus : " + " ; ".join(faux))
    log("%s : emissions %s Mt, empreinte %s Mt, dont importations %s ; depuis %s : emissions -%s %%, empreinte -%s %%"
        % (A["annee"], A["emissions"], A["empreinte"], A["importee"], A["annee_depart"], A["evol_emissions"], A["evol_empreinte"]))
    if check:
        log("--check : temoins et gardes passes (%d cles d'affichage), rien ecrit." % len(A))
        return 0
    svg = figure(j, A)
    payload = {"meta": {"page": "https://" + PAGE_URL, "licence": "CC BY 4.0 pour la compilation et la figure ; données Insee-SDES",
                        "source": "Insee et SDES, " + EDITION + ", fichier de données (figure 1, figure 3, tableau complémentaire 2)",
                        "champ": "France ; CO2, CH4, N2O et gaz fluorés",
                        "definitions": {
                            "emissions": "émissions des unités résidentes françaises (inventaire au format des comptes d'émissions dans l'air), transport international des résidents compris",
                            "empreinte": "estimation des émissions associées à la demande finale française, hors exportations, où qu'elles aient lieu",
                            "menages": "émissions directes des ménages (chauffage, carburants)",
                            "interieure": "émissions de la production intérieure destinée à la demande finale française",
                            "exportee": "émissions de la production intérieure exportée (décomposition de l'année seulement)",
                            "importee": "émissions importées destinées à la demande finale française",
                            "avertissement": "série révisée à chaque édition ; la dernière année est provisoire ; 1990-2009 rétropolées"}},
               "decomposition": j["decomposition"], "serie": j["serie"], "affichage": A}
    releve = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    if OUT_DATA.exists():
        try:
            prev = json.loads(OUT_DATA.read_text(encoding="utf-8"))
            p2 = dict(prev)
            ancien = p2.pop("releve_le", None)
            p2.pop("_licence", None)
            if ancien and json.dumps(p2, sort_keys=True, ensure_ascii=False) == json.dumps(payload, sort_keys=True, ensure_ascii=False) \
                    and (OUT_IMG / (FIGURE + ".svg")).exists() and (OUT_IMG / (FIGURE + ".svg")).read_text(encoding="utf-8") == svg \
                    and (OUT_IMG / (FIGURE + ".png")).exists() and OUT_STATIC.exists() \
                    and OUT_CSV.exists() and OUT_CSV.read_text(encoding="utf-8-sig") == csv_texte(j):
                log("Donnees et figure identiques : rien ecrit (releve_le conserve : %s)." % ancien)
                return 0
        except (ValueError, KeyError):
            pass
    payload = {"releve_le": releve, "_licence": "CC BY 4.0 — compilation Stéphane Lalut ; source Insee-SDES", **payload}
    txt = json.dumps(payload, ensure_ascii=False, indent=1)
    OUT_DATA.write_text(txt, encoding="utf-8")
    OUT_STATIC.write_text(txt, encoding="utf-8")
    OUT_CSV.write_text(csv_texte(j), encoding="utf-8-sig", newline="\n")
    import cairosvg
    (OUT_IMG / (FIGURE + ".svg")).write_text(svg, encoding="utf-8", newline="\n")
    cairosvg.svg2png(url=str(OUT_IMG / (FIGURE + ".svg")), write_to=str(OUT_IMG / (FIGURE + ".png")), output_width=1440,
                     background_color="white")
    log("Ecrit : data/ et static/empreinte_carbone.json, static/empreinte_carbone.csv, static/img/%s.svg + .png" % FIGURE)
    return 0


if __name__ == "__main__":
    sys.exit(main())
