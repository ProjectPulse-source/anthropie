#!/usr/bin/env python3
"""generer_parcours_licence.py -- activité de SES « Avons-nous tous le même droit à l'erreur ? » (/enseignants/ecole-et-parcours/).

Jeu : devenir, un an après, des néo-bacheliers inscrits en première année de licence, selon l'origine sociale.
Source : MESRE-SIES, Note Flash « Parcours et réussite en licence », tableaux nationaux, feuille « Devenir cohorte N »
(systèmes d'information SISE, SCOLARITE et SIFA). Les tableurs sont ARCHIVÉS dans scripts/sources_enseignants/ : le
ministère n'offre ni API ni adresse stable, et ses pages refusent les scripts. Aucun appel réseau ici.

Décisions de l'auteur (02/10/2026) : figure propre à /enseignants/ (exception écrite dans CLAUDE.md, « Couche
pédagogique ») ; construction sur la cohorte 2023 sans attendre l'édition suivante.

PIÈCES : téléchargées le 2 octobre 2026. Le SIES corrige parfois ses fichiers après parution (onglets de la cohorte 2021
remplacés le 26/11/2025, selon la méthodologie du tableur) : l'empreinte dit quelle version est lue, cette date dit quand.

METTRE À JOUR (édition annuelle, en novembre) : `python scripts/maj_sources.py integrer` (ou le workflow quotidien
maj-sources.yml) dépose le tableur dans scripts/sources_enseignants/, ajoute son empreinte et sa ligne au registre
data/sources_maj.json, puis relance ce générateur. Les gardes disent si les constats de la fiche tiennent sur la
cohorte nouvelle ; si la définition des colonnes change, arrêt, et rien n'est adopté.

SÉRIE COURTE, ET POURQUOI. Trois cohortes seulement sont comparables (2021, 2022, 2023). Avant, la « réorientation »
ne couvrait que l'université et la dernière colonne s'appelait « sortie de l'université », réorientations vers une STS
ou une école comprises : les tableurs 2018 à 2020 sont archivés comme preuve de cette rupture, ils ne sont pas lus.

TÉMOINS INDÉPENDANTS, bloquants :
  1. chaque taux publié est recalculé depuis les effectifs publiés (écart toléré : l'arrondi au dixième) ;
  2. conservation du périmètre : les quatre origines et la non-réponse somment exactement à l'ensemble ;
  3. les quatre devenirs couvrent les inscrits à RESIDU_MAX près (écart mesuré, écrit dans le jeu, jamais corrigé).

Chaque qualificatif de la fiche et du corrigé est une garde (bloc `gardes`) : si une donnée le dément, rien n'est écrit.
SVG et PNG sont produits ensemble ou pas du tout (cairosvg).

Usage : python scripts/generer_parcours_licence.py [--check]
Sorties : data/ et static/parcours_licence.json (bloc affichage), static/parcours_licence.csv (format long),
          static/img/parcours-licence-devenir.svg + .png
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "scripts" / "sources_enseignants"
OUT_DATA = ROOT / "data" / "parcours_licence.json"
OUT_STATIC = ROOT / "static" / "parcours_licence.json"
OUT_CSV = ROOT / "static" / "parcours_licence.csv"
OUT_IMG = ROOT / "static" / "img"
FIGURE = "parcours-licence-devenir"
PAGE_URL = "stephane-lalut.com/enseignants/ecole-et-parcours/"

# cohorte (année d'entrée en L1) -> (tableur archivé, feuille, référence de la note)
REGISTRE = ROOT / "data" / "sources_maj.json"
MOIS = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre")
LETTRES = {3: "trois", 4: "quatre", 5: "cinq", 6: "six", 7: "sept", 8: "huit", 9: "neuf", 10: "dix"}


def source(ident: str) -> dict:
    """Une ligne du registre des sources (data/sources_maj.json) : c'est lui qui dit quelle édition lire."""
    for x in json.loads(REGISTRE.read_text(encoding="utf-8"))["sources"]:
        if x["id"] == ident:
            return x
    print("ECHEC : source %s absente de data/sources_maj.json" % ident)
    sys.exit(1)


def date_fr(iso: str) -> str:
    a, m, j = (int(x) for x in iso.split("-"))
    return "%d%s %s %d" % (j, "er" if j == 1 else "", MOIS[m - 1], a)


# Cohortes lues dans le registre : une édition nouvelle s'ajoute là (scripts/maj_sources.py), jamais ici.
COHORTES = {int(k): tuple(v) for k, v in source("sies-licence")["cohortes"].items()}
COLONNES = {"inscrits": "Inscrits en L1", "passage": "Passage en L2", "redoublement": "Redoublement en L1",
            "reorientation": "Réorientation", "sortie": "Sortie de l'enseignement supérieur"}
TAUX = {"passage": "Taux de passage en L2", "redoublement": "Taux de redoublement en L1",
        "reorientation": "Taux de réorientation", "sortie": "Taux de sortie de l'enseignement supérieur"}
STRICT = "Dont redoublement strict"          # absent du tableur de la cohorte 2021
ISSUES = ("passage", "redoublement", "reorientation", "sortie")
ORIGINES = [("tf", "Très favorisée"), ("f", "Favorisée"), ("ad", "Assez défavorisée"), ("d", "Défavorisée")]
NON_REPONSE = "Non réponse"
MENTIONS = [("mention_tb", "Très bien"), ("mention_p2", "Passable deuxième groupe")]
RESIDU_MAX = 0.2      # % des inscrits : part tolérée d'étudiants hors des quatre devenirs publiés

W = 720
FONT = "Inter, 'Helvetica Neue', Arial, sans-serif"
# Aucune teinte nouvelle : bleu et orange de la charte, deux gris déjà employés par update_dette_dynamique.py.
# Ordre des gris choisi pour l'impression en noir : deux voisins n'ont jamais la même clarté.
BLEU, ORANGE, GRIS, GRIS_CLAIR = "#184f95", "#eb6834", "#8a8781", "#c9c5c0"
INK, INK2, MUTED, GRID = "#26262f", "#55524f", "#96928f", "#dcd8d3"
COULEUR = {"passage": (BLEU, "#ffffff"), "redoublement": (GRIS, "#ffffff"), "reorientation": (GRIS_CLAIR, INK),
           "sortie": (ORANGE, INK)}
LIBELLE = {"passage": ("Passés en", "2ᵉ année"), "redoublement": ("Recommencent", "une 1ʳᵉ année"),
           "reorientation": ("Réorientés", "hors licence"), "sortie": ("Non", "retrouvés")}
# « Non retrouvés » et non « sortis du supérieur » (contre-expertise du 02/10/2026, F3) : la catégorie du SIES désigne une
# absence des fichiers, et la figure voyage sans la page qui l'explique.


def log(msg: str) -> None:
    print(msg.encode("ascii", "replace").decode("ascii"))


def fail(msg: str) -> None:
    log("ECHEC : " + msg)
    log("Aucun fichier ecrit.")
    sys.exit(1)


def fr(v: float, dec: int = 1) -> str:
    return (("%." + str(dec) + "f") % v).replace("-", "−").replace(".", ",")


def milliers(n: int) -> str:
    return "{:,}".format(n).replace(",", " ")


def esc(s: str) -> str:
    return html.escape(s, quote=False)


# ------------------------------------------------------------------ lecture
def empreintes() -> dict[str, str]:
    f = SRC / "SHA256SUMS"
    if not f.is_file():
        fail("scripts/sources_enseignants/SHA256SUMS absent")
    out = {}
    for ligne in f.read_text(encoding="utf-8").splitlines():
        if ligne.strip():
            h, nom = ligne.split(None, 1)
            out[nom.strip().lstrip("*")] = h
    return out


def lire_cohorte(annee: int, sums: dict[str, str]) -> dict:
    import openpyxl
    import warnings
    fichier, feuille, note = COHORTES[annee]
    p = SRC / fichier
    if not p.is_file():
        fail("tableur absent : %s" % fichier)
    if sums.get(fichier) != hashlib.sha256(p.read_bytes()).hexdigest():
        fail("%s : empreinte differente de SHA256SUMS (piece remplacee ?)" % fichier)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")           # zones de texte de la feuille « Méthodologie », hors sujet ici
        wb = openpyxl.load_workbook(p, data_only=True)
    if feuille not in wb.sheetnames:
        fail("%s : feuille %r absente" % (fichier, feuille))
    rows = list(wb[feuille].iter_rows(values_only=True))
    entetes = next((list(r) for r in rows if r and r[0] == "Discipline en L1"), None)
    if entetes is None:
        fail("%s : ligne d'en-tetes introuvable" % fichier)
    # La définition tient aux libellés : une colonne renommée (« Sortie de l'université », cohortes 2018-2020) arrête tout.
    manquantes = [c for c in list(COLONNES.values()) + list(TAUX.values()) if c not in entetes]
    if manquantes:
        fail("%s : colonnes absentes ou renommees (definition changee ?) : %s" % (fichier, ", ".join(manquantes)))
    ix = {k: entetes.index(v) for k, v in COLONNES.items()}
    it = {k: entetes.index(v) for k, v in TAUX.items()}
    i_strict = entetes.index(STRICT) if STRICT in entetes else None

    def ligne(r) -> dict:
        d = {k: r[i] for k, i in ix.items()}
        if any(not isinstance(v, (int, float)) for v in d.values()):
            fail("%d, %s : valeur manquante dans les effectifs" % (annee, r[2]))   # manquant n'est pas zéro
        d = {k: int(v) for k, v in d.items()}
        d["taux_publies"] = {k: float(r[i]) for k, i in it.items()}
        d["redoublement_strict"] = int(r[i_strict]) if i_strict is not None and isinstance(r[i_strict], (int, float)) else None
        # témoin 1 : le taux publié se retrouve depuis les effectifs
        for k in ISSUES:
            if abs(100 * d[k] / d["inscrits"] - d["taux_publies"][k]) > 0.0501:
                fail("%d, %s, %s : taux publie %.1f != effectifs %.3f" % (annee, r[2], k, d["taux_publies"][k], 100 * d[k] / d["inscrits"]))
        return d

    toutes = [r for r in rows if r and r[0] == "Toutes disciplines"]
    ens = [r for r in toutes if r[1] == "Ensemble"]
    soc = [r for r in toutes if str(r[1]).startswith("Origine sociale") and "nage" not in str(r[1])]
    men = [r for r in toutes if str(r[1]).startswith("Mention")]
    if len(ens) != 1 or len(soc) != 5:
        fail("%d : %d ligne(s) Ensemble et %d d'origine sociale, attendu 1 et 5" % (annee, len(ens), len(soc)))
    par_modalite = {r[2]: ligne(r) for r in soc}
    attendues = [lib for _, lib in ORIGINES] + [NON_REPONSE]
    if sorted(par_modalite) != sorted(attendues):
        fail("%d : modalites d'origine sociale inattendues : %s" % (annee, ", ".join(sorted(par_modalite))))
    out = {"annee": annee, "note": note, "fichier": fichier, "ensemble": ligne(ens[0]),
           "origines": {cle: par_modalite[lib] for cle, lib in ORIGINES}, "non_reponse": par_modalite[NON_REPONSE],
           "mentions": {}}
    for cle, lib in MENTIONS:
        m = [r for r in men if r[2] == lib]
        if len(m) != 1:
            fail("%d : mention %r introuvable" % (annee, lib))
        out["mentions"][cle] = ligne(m[0])
    # témoin 2 : conservation du périmètre -- entré == retenu (quatre origines) + écarté (non-réponse, motif nommé)
    somme = sum(o["inscrits"] for o in out["origines"].values()) + out["non_reponse"]["inscrits"]
    if somme != out["ensemble"]["inscrits"]:
        fail("%d : origines + non-reponse = %d != ensemble %d" % (annee, somme, out["ensemble"]["inscrits"]))
    # témoin 3 : les quatre devenirs couvrent les inscrits (résidu mesuré, jamais redistribué)
    for nom, d in [("ensemble", out["ensemble"])] + list(out["origines"].items()):
        d["residu"] = d["inscrits"] - sum(d[k] for k in ISSUES)
        if abs(d["residu"]) > RESIDU_MAX / 100 * d["inscrits"]:
            fail("%d, %s : %d etudiants hors des quatre devenirs (> %.1f %%)" % (annee, nom, d["residu"], RESIDU_MAX))
    return out


# ------------------------------------------------------------------ calcul
def indicateurs(d: dict) -> dict:
    """Taux sur les inscrits (publiés) et parts parmi les non-passants (calcul propre, depuis les effectifs)."""
    np_ = d["inscrits"] - d["passage"]
    t = d["taux_publies"]
    non_passage_pub = round(100 - t["passage"], 1)
    r = {"non_passants": np_, "taux_non_passage": non_passage_pub,
         "sortie_np": 100 * d["sortie"] / np_, "reorientation_np": 100 * d["reorientation"] / np_,
         "redoublement_np": 100 * d["redoublement"] / np_,
         # ce que l'élève obtient avec les seuls taux arrondis de la figure : c'est ce que le corrigé affiche
         "sortie_np_eleve": round(100 * t["sortie"] / non_passage_pub, 1),
         "reorientation_np_eleve": round(100 * t["reorientation"] / non_passage_pub, 1),
         "redoublement_np_eleve": round(100 * t["redoublement"] / non_passage_pub, 1)}
    if d["redoublement_strict"] is not None:
        r["redoublement_autre"] = 100 * (d["redoublement"] - d["redoublement_strict"]) / d["inscrits"]
    return r


# ------------------------------------------------------------------ affichage et gardes
def affichage(cohortes: dict[int, dict]) -> tuple[dict, list[str]]:
    der = max(cohortes)
    c = cohortes[der]
    o, ens, men = c["origines"], c["ensemble"], c["mentions"]
    ind = {k: indicateurs(v) for k, v in o.items()}
    A = {"cohorte": str(der), "annee_devenir": str(der + 1), "annee_univ": "%d-%d" % (der, der + 1),
         "annee_univ_suivante": "%d-%d" % (der + 1, der + 2), "note": c["note"],
         "inscrits": milliers(ens["inscrits"]), "cohortes_lib": "%d à %d" % (min(cohortes), der),
         "nb_cohortes": str(len(cohortes)), "nb_cohortes_lettres": LETTRES.get(len(cohortes), str(len(cohortes))),
         "Nb_cohortes_lettres": LETTRES.get(len(cohortes), str(len(cohortes))).capitalize(),
         "premiere_cohorte": str(min(cohortes)), "telecharge_le": date_fr(source("sies-licence")["telecharge_le"])}
    for k in ISSUES:
        A["ens_" + k] = fr(ens["taux_publies"][k])
    for cle, d in o.items():
        for k in ISSUES:
            A["%s_%s" % (cle, k)] = fr(d["taux_publies"][k])
        A[cle + "_inscrits"] = milliers(d["inscrits"])
        A[cle + "_non_passage"] = fr(ind[cle]["taux_non_passage"])
        A[cle + "_sortie_np"] = fr(ind[cle]["sortie_np_eleve"])
        A[cle + "_reorientation_np"] = fr(ind[cle]["reorientation_np_eleve"])
        A[cle + "_redoublement_np"] = fr(ind[cle]["redoublement_np_eleve"])
        if "redoublement_autre" in ind[cle]:
            A[cle + "_redoublement_autre"] = fr(ind[cle]["redoublement_autre"])
    A["ecart_passage"] = fr(o["tf"]["taux_publies"]["passage"] - o["d"]["taux_publies"]["passage"])
    A["ecart_sortie"] = fr(o["d"]["taux_publies"]["sortie"] - o["tf"]["taux_publies"]["sortie"])
    A["ecart_sortie_np"] = fr(abs(ind["d"]["sortie_np_eleve"] - ind["tf"]["sortie_np_eleve"]))
    # même écart pour l'élève qui garde ses décimales jusqu'au bout : les deux démarches sont justes (F8)
    t_tf, t_d = o["tf"]["taux_publies"], o["d"]["taux_publies"]
    exact = 100 * t_d["sortie"] / (100 - t_d["passage"]) - 100 * t_tf["sortie"] / (100 - t_tf["passage"])
    A["ecart_sortie_np_decimales"] = fr(abs(exact))
    # décomposition arithmétique : sortie qu'aurait l'origine défavorisée avec la part de sortie des non-passés très favorisés
    contrefactuel = ind["d"]["taux_non_passage"] * ind["tf"]["sortie_np_eleve"] / 100
    part_frequence = (contrefactuel - t_tf["sortie"]) / (t_d["sortie"] - t_tf["sortie"])
    reor = [d["taux_publies"]["reorientation"] for d in o.values()]
    A["reorientation_min"], A["reorientation_max"] = fr(min(reor)), fr(max(reor))
    A["mention_tb_passage"] = fr(men["mention_tb"]["taux_publies"]["passage"])
    A["mention_p2_passage"] = fr(men["mention_p2"]["taux_publies"]["passage"])
    A["non_reponse_part"] = fr(100 * c["non_reponse"]["inscrits"] / ens["inscrits"])
    A["residu"] = milliers(abs(ens["residu"]))

    # Une garde par qualificatif écrit dans /enseignants/ (fiche, corrigé, légende, méthode).
    def chaque(test) -> bool:
        return all(test(x["origines"], {k: indicateurs(v) for k, v in x["origines"].items()}) for x in cohortes.values())

    ecart_mentions = men["mention_tb"]["taux_publies"]["passage"] - men["mention_p2"]["taux_publies"]["passage"]
    ecart_origines = o["tf"]["taux_publies"]["passage"] - o["d"]["taux_publies"]["passage"]
    gardes = [
        ("au moins trois cohortes comparables (le nombre exact est un jeton, la page ne l'écrit plus en dur)", len(cohortes) >= 3),
        ("le passage en 2e année décroît de l'origine très favorisée à la défavorisée, sur chaque cohorte",
         chaque(lambda og, i: og["tf"]["passage"] / og["tf"]["inscrits"] > og["f"]["passage"] / og["f"]["inscrits"]
                > og["ad"]["passage"] / og["ad"]["inscrits"] > og["d"]["passage"] / og["d"]["inscrits"])),
        ("écart de passage entre origines extrêmes « de plus de quinze points »", ecart_origines > 15),
        ("taux de réorientation « presque le même » d'une origine à l'autre (moins d'un point d'écart), sur chaque cohorte",
         chaque(lambda og, i: max(100 * d["reorientation"] / d["inscrits"] for d in og.values())
                - min(100 * d["reorientation"] / d["inscrits"] for d in og.values()) < 1.0)),
        ("« les trois autres devenirs s'écartent » : plus de cinq points entre origines pour chacun",
         all(max(d["taux_publies"][k] for d in o.values()) - min(d["taux_publies"][k] for d in o.values()) > 5
             for k in ("passage", "redoublement", "sortie"))),
        ("la sortie est plus fréquente pour l'origine défavorisée (plus de quatre points sur les inscrits)",
         o["d"]["taux_publies"]["sortie"] - o["tf"]["taux_publies"]["sortie"] > 4),
        ("parmi les non-passants, parts de sortie « voisines » aux deux extrémités (moins de trois points), sur chaque cohorte",
         chaque(lambda og, i: abs(i["d"]["sortie_np"] - i["tf"]["sortie_np"]) < 3.0)),
        ("l'écart sur tous les inscrits correspond « surtout » à la fréquence du non-passage (plus des deux tiers, "
         "décomposition arithmétique)", part_frequence > 2 / 3),
        ("cet écart réduit s'écrit au singulier dans le corrigé, arrondi ou non : moins de deux points",
         abs(ind["d"]["sortie_np_eleve"] - ind["tf"]["sortie_np_eleve"]) < 2 and abs(exact) < 2),
        ("parmi les non-passants, la réorientation est plus fréquente pour l'origine très favorisée (plus de cinq points), "
         "le redoublement pour la défavorisée, sur chaque cohorte",
         chaque(lambda og, i: i["tf"]["reorientation_np"] - i["d"]["reorientation_np"] > 5
                and i["d"]["redoublement_np"] > i["tf"]["redoublement_np"])),
        ("le calcul de l'élève (taux arrondis) rejoint le calcul sur effectifs à 0,3 point près",
         all(abs(ind[k][x + "_np_eleve"] - ind[k][x + "_np"]) < 0.3 for k in ind
             for x in ("sortie", "reorientation", "redoublement"))),
        ("recommencer une L1 ailleurs ou dans une autre discipline : plus fréquent pour l'origine défavorisée",
         "redoublement_autre" in ind["tf"] and ind["d"]["redoublement_autre"] > ind["tf"]["redoublement_autre"]),
        # Les deux écarts ne se comparent pas (découpages différents, tableau non croisé : F6) ; seul « forts écarts » est écrit.
        ("« forts écarts » de passage selon la mention (plus de cinquante points entre les deux mentions citées)",
         ecart_mentions > 50),
    ]
    return A, [nom for nom, ok in gardes if not ok]


# ------------------------------------------------------------------ figure
def figure(c: dict, A: dict) -> str:
    """Cinq barres empilées à 100 % (quatre origines, puis l'ensemble) : quatre devenirs un an après."""
    lignes = [(lib, c["origines"][cle]) for cle, lib in ORIGINES] + [("Ensemble", c["ensemble"])]
    X0, X1, TOP, HB, PAS = 132, W - 8, 96, 30, 44
    H = TOP + PAS * (len(lignes) - 1) + HB + 14 + 36      # 36 : trois lignes de définitions dans le cartouche
    titre = "Un an après une première année de licence : quatre devenirs, selon l'origine sociale"
    o = c["origines"]
    desc = ("Barres empilées à 100 %%, une par origine sociale. Bacheliers %s entrés en première année de licence : "
            "%s %% des étudiants d'origine très favorisée sont passés en deuxième année un an après, contre %s %% de ceux "
            "d'origine défavorisée ; %s %% et %s %% recommencent une première année ; %s %% et %s %% sont réorientés "
            "hors de la licence ; %s %% et %s %% ne sont plus retrouvés dans les fichiers d'inscription."
            % (A["cohorte"], A["tf_passage"], A["d_passage"], A["tf_redoublement"], A["d_redoublement"],
               A["tf_reorientation"], A["d_reorientation"], A["tf_sortie"], A["d_sortie"]))
    e = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 %d %d" role="img" font-variant-numeric="tabular-nums" '
         'aria-labelledby="pl-t pl-d" font-family="%s">' % (W, H + 52, FONT),
         '<title id="pl-t">%s</title><desc id="pl-d">%s</desc>' % (esc(titre), esc(desc)),
         '<rect width="%d" height="%d" fill="#ffffff"/>' % (W, H + 52),
         '<text x="0" y="18" font-size="15" font-weight="600" fill="%s">%s</text>' % (INK, esc(titre)),
         '<text x="0" y="38" font-size="11" fill="%s">%s</text>'
         % (INK2, esc("Bacheliers %s inscrits en 1ʳᵉ année de licence à la rentrée %s : situation à la rentrée %s, en %% des inscrits."
                      % (A["cohorte"], A["cohorte"], A["annee_devenir"])))]
    larg = X1 - X0
    for n, (lib, d) in enumerate(lignes):
        y = TOP + PAS * n
        t = d["taux_publies"]
        total = sum(t[k] for k in ISSUES)
        gras = ' font-weight="600"' if lib == "Ensemble" else ""
        if lib == "Ensemble":
            e.append('<line x1="0" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="0.6"/>' % (y - 7, W, y - 7, GRID))
        e.append('<text x="%d" y="%.1f" font-size="11.5" fill="%s" text-anchor="end"%s>%s</text>'
                 % (X0 - 10, y + HB / 2 + 4, INK, gras, esc(lib)))
        x = float(X0)
        for k in ISSUES:
            w = larg * t[k] / total
            fond, encre = COULEUR[k]
            e.append('<rect x="%.1f" y="%d" width="%.1f" height="%d" fill="%s" stroke="#ffffff" stroke-width="1.5"/>'
                     % (x, y, w, HB, fond))
            e.append('<text x="%.1f" y="%.1f" font-size="12" font-weight="600" fill="%s" text-anchor="middle">%s</text>'
                     % (x + w / 2, y + HB / 2 + 4, encre, esc(fr(t[k]))))
            if n == 0:      # étiquetage direct : le nom de chaque devenir au-dessus de son segment, première barre
                for j, mot in enumerate(LIBELLE[k]):
                    e.append('<text x="%.1f" y="%.1f" font-size="10.5" fill="%s" text-anchor="middle">%s</text>'
                             % (x + w / 2, y - 19 + 12 * j, INK2, esc(mot)))
            x += w
    y0 = H + 4 - 36
    e.append('<line x1="0" y1="%.1f" x2="%d" y2="%.1f" stroke="%s"/>' % (y0, W, y0, GRID))
    cart = [("MESRE-SIES, %s, tableaux nationaux, feuille « Devenir cohorte %s » (SISE, SCOLARITE, SIFA)" % (c["note"], A["cohorte"]), INK2),
            ("Passés : inscrits en 2ᵉ ou 3ᵉ année de licence, ou entrés en études de santé depuis une licence accès santé. "
             "Recommencent : changements de discipline compris.", INK2),
            ("Non retrouvés : absents des fichiers d'inscription (« sortie de l'enseignement supérieur » pour le SIES) ; "
             "certains étudient à l'étranger ou hors du champ suivi.", INK2),
            ("Situation d'inscription un an après, non un diplôme. Le tableau ne croise pas l'origine sociale et le niveau "
             "scolaire à l'entrée.", INK2),
            ("Taux arrondis : leur somme peut différer de 100 %%. Origine non renseignée (%s %% des inscrits) : pas de barre, "
             "comprise dans l'ensemble." % A["non_reponse_part"], INK2),
            ("Figure Stéphane Lalut, CC BY 4.0 · " + PAGE_URL, MUTED)]
    for k, (txt, col) in enumerate(cart):
        e.append('<text x="0" y="%.1f" font-size="9" fill="%s">%s</text>' % (y0 + 13 + 12 * k, col, esc(txt)))
    e.append("</svg>")
    return "\n".join(e)


def csv_texte(cohortes: dict[int, dict]) -> str:
    """Format long : une ligne par cohorte, origine et devenir."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["cohorte", "origine_sociale", "devenir", "effectif", "inscrits", "taux_pct_inscrits", "part_pct_non_passants"])
    noms = dict(ORIGINES, ensemble="Ensemble", non_reponse="Non-réponse")
    for a in sorted(cohortes):
        c = cohortes[a]
        for cle, d in list(c["origines"].items()) + [("non_reponse", c["non_reponse"]), ("ensemble", c["ensemble"])]:
            np_ = d["inscrits"] - d["passage"]
            for k in ISSUES:
                w.writerow([a, noms[cle], k, d[k], d["inscrits"], "%.1f" % d["taux_publies"][k],
                            "" if k == "passage" else "%.1f" % (100 * d[k] / np_)])
    return buf.getvalue()


# ------------------------------------------------------------------ main
def main() -> int:
    check = "--check" in sys.argv[1:]
    if not check:
        try:
            import cairosvg  # noqa: F401
        except ImportError:
            fail("cairosvg absent : SVG et PNG se produisent ensemble ou pas du tout (pip install cairosvg)")
    sums = empreintes()
    cohortes = {a: lire_cohorte(a, sums) for a in sorted(COHORTES)}
    A, faux = affichage(cohortes)
    if faux:
        fail("la page affirme ce que les donnees ne soutiennent plus : " + " ; ".join(faux))
    der = max(cohortes)
    log("Cohorte %d : passage %s / %s, reorientation %s a %s, sortie parmi les non-passants %s / %s"
        % (der, A["tf_passage"], A["d_passage"], A["reorientation_min"], A["reorientation_max"], A["tf_sortie_np"], A["d_sortie_np"]))
    if check:
        log("--check : temoins et gardes passes (%d cles d'affichage), rien ecrit." % len(A))
        return 0
    svg = figure(cohortes[der], A)
    payload = {"meta": {"page": "https://" + PAGE_URL, "licence": "CC BY 4.0 pour la compilation et la figure ; données MESRE-SIES",
                        "champ": "néo-bacheliers inscrits pour la première fois en première année de licence, universités et établissements assimilés, France entière",
                        "source": "MESRE-SIES, Notes Flash « Parcours et réussite en licence », tableaux nationaux (SISE, SCOLARITE, SIFA)",
                        "definitions": {
                            "passage": "inscrit l'année suivante en deuxième année de licence (ou en troisième, ou en filière de santé pour les licences accès santé), quels que soient la mention et l'établissement",
                            "redoublement": "à nouveau inscrit en première année de licence, quels que soient la mention et l'établissement : un changement de discipline à l'intérieur de la licence en fait partie",
                            "reorientation": "inscrit l'année suivante dans une autre formation que la licence, à l'université (BUT, par exemple) ou en dehors (STS, école)",
                            "sortie": "non retrouvé dans les systèmes d'information ; une faible part peut être inscrite à l'étranger ou dans une formation non couverte",
                            "origine_sociale": "catégorie socioprofessionnelle du parent référent, regroupée par le SIES : très favorisée (cadres, enseignants…), favorisée (professions intermédiaires), assez défavorisée (employés…), défavorisée (ouvriers…)",
                            "part_pct_non_passants": "calcul de cette page : effectif rapporté aux inscrits non passés en deuxième année ; non publié par le SIES",
                            "residu": "inscrits moins la somme des quatre devenirs publiés ; mesuré, non redistribué"}},
               "cohortes": [cohortes[a] for a in sorted(cohortes)], "affichage": A}
    releve = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    # « Rien réécrit à données identiques » : le relevé ne bouge que si un chiffre, la figure ou le CSV a bougé.
    if OUT_DATA.exists():
        try:
            prev = json.loads(OUT_DATA.read_text(encoding="utf-8"))
            p2 = dict(prev)
            ancien = p2.pop("releve_le", None)
            p2.pop("_licence", None)
            if ancien and json.dumps(p2, sort_keys=True, ensure_ascii=False) == json.dumps(payload, sort_keys=True, ensure_ascii=False) \
                    and (OUT_IMG / (FIGURE + ".svg")).exists() and (OUT_IMG / (FIGURE + ".svg")).read_text(encoding="utf-8") == svg \
                    and (OUT_IMG / (FIGURE + ".png")).exists() and OUT_STATIC.exists() \
                    and OUT_CSV.exists() and OUT_CSV.read_text(encoding="utf-8-sig") == csv_texte(cohortes):
                log("Donnees et figure identiques : rien ecrit (releve_le conserve : %s)." % ancien)
                return 0
        except (ValueError, KeyError):
            pass
    payload = {"releve_le": releve, "_licence": "CC BY 4.0 — compilation Stéphane Lalut ; source MESRE-SIES", **payload}
    txt = json.dumps(payload, ensure_ascii=False, indent=1)
    OUT_DATA.write_text(txt, encoding="utf-8")
    OUT_STATIC.write_text(txt, encoding="utf-8")
    OUT_CSV.write_text(csv_texte(cohortes), encoding="utf-8-sig", newline="\n")  # BOM : Excel lit les accents
    import cairosvg
    (OUT_IMG / (FIGURE + ".svg")).write_text(svg, encoding="utf-8", newline="\n")
    cairosvg.svg2png(url=str(OUT_IMG / (FIGURE + ".svg")), write_to=str(OUT_IMG / (FIGURE + ".png")), output_width=1440,
                     background_color="white")
    log("Ecrit : data/ et static/parcours_licence.json, static/parcours_licence.csv, static/img/%s.svg + .png" % FIGURE)
    return 0


if __name__ == "__main__":
    sys.exit(main())
