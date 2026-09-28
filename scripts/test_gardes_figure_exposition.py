#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Temoin positif des gardes de la figure « qui-paie-exposition ».

    python scripts/test_gardes_figure_exposition.py

La figure F6 de generer_figures_qui_paie.py repose sur quatre gardes bloquantes.
Une garde qu'on n'a jamais vue MORDRE est reputee absente : elle produit de la
confiance sans en porter. Ce script les eprouve une par une, par mutation reelle
du classeur source dans un dossier temporaire -- le depot n'est jamais touche.

Ce qu'il verifie, et qui ne va pas de soi : que la garde VISEE repond. Une
mutation arretee par une AUTRE garde ne prouve rien sur celle qu'on voulait
eprouver. Deux cas le montrent ici :

  - la cellule atypique (G4) ne s'atteint qu'en mutant les DEUX publications de
    facon coherente, sinon le temoin croise arrete tout le premier ;
  - le U (G3) ne tombe pas quand on aplatit UN poste : rapporte au revenu, un
    poste constant en masse redonne une courbe decroissante tres proche de la
    vraie. Ce qui tue le U, c'est que les trois instruments cessent de differer.

Sortie : 0 si les quatre gardes mordent sur leur propre motif, 1 sinon.

CONDITION DE MORT : ce script disparait avec la figure F6. Il ne se generalise
pas aux autres figures, dont les gardes sont d'une autre nature.
"""
from __future__ import annotations

import hashlib
import importlib.util
import shutil
import sys
import tempfile
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

REPO = Path(__file__).resolve().parent.parent
SRC_CND = REPO / "scripts" / "sources" / "insee_T_CND_101_vingtiemes_2020_2023.xlsx"
SRC_118 = REPO / "scripts" / "sources" / "insee_IA118_comptes_distribues_2023.xlsx"

spec = importlib.util.spec_from_file_location(
    "gfqp", REPO / "scripts" / "generer_figures_qui_paie.py")
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)
XLSX118_ORIG = mod.XLSX

try:
    import openpyxl
except ImportError:
    print("ECHEC : openpyxl absent")
    sys.exit(1)


def _lignes(ws, debuts: dict) -> dict:
    out = {}
    for r in ws.iter_rows(min_col=1, max_col=1):
        v = str(r[0].value or "")
        for cle, debut in debuts.items():
            if v.startswith(debut):
                out[cle] = r[0].row
    return out


def mut_structure(wb) -> None:
    """G1 : un libelle change en 2021 ; la lecture par libelle devient fausse."""
    ws = wb["CND_2021"]
    l = _lignes(ws, {"pen": "Pensions de retraite"})
    ws.cell(row=l["pen"], column=1).value = "Pensions (nouveau libelle)"


def mut_temoin(wb) -> None:
    """G2 : une valeur 2023 diverge de l'autre publication de la meme source."""
    ws = wb["CND_2023"]
    l = _lignes(ws, {"csg": "Contribution sociale"})
    ws.cell(row=l["csg"], column=2).value = -40.0          # -0,5 attendu en V1


def mut_atypique(wb) -> None:
    """G4 : l'Insee corrige la cellule ; l'ecartement du point D1 perd son motif."""
    ws = wb["CND_2023"]
    l = _lignes(ws, {"ir": "Impôt sur le revenu"})
    ws.cell(row=l["ir"], column=2).value = -0.05


def mut_atypique_118(wb) -> None:
    ws = wb["Tableau complémentaire"]
    l = _lignes(ws, {"ir": "Impôt sur le revenu"})
    ws.cell(row=l["ir"], column=3).value = -0.05           # col.2 Ensemble, col.3 V1


def mut_u(wb) -> None:
    """G3 : en 2020, les trois decisions recoivent le MEME profil ; le U doit tomber."""
    ws = wb["CND_2020"]
    l = _lignes(ws, {"ens": "Enseignement", "csg": "Contribution sociale",
                     "ir": "Impôt sur le revenu", "aut": "Autres impôts sur les revenus",
                     "pen": "Pensions de retraite"})
    ens = [ws.cell(row=l["ens"], column=c).value for c in range(2, 22)]
    for c, val in zip(range(2, 22), ens):
        ws.cell(row=l["csg"], column=c).value = -float(val)
        ws.cell(row=l["ir"], column=c).value = 0.0
        ws.cell(row=l["aut"], column=c).value = 0.0
        ws.cell(row=l["pen"], column=c).value = float(val)


CAS = [("structure (G1)", mut_structure, None, "introuvable"),
       ("temoin croise (G2)", mut_temoin, None, "témoin croisé"),
       ("cellule atypique (G4)", mut_atypique, mut_atypique_118, "n'est plus atypique"),
       ("le U (G3)", mut_u, None, "creux")]


def eprouver(bac: Path, nom: str, muter, muter_118) -> tuple[bool, str]:
    cible = bac / ("cnd_" + nom + ".xlsx")
    shutil.copy2(SRC_CND, cible)
    wb = openpyxl.load_workbook(cible)
    muter(wb)
    wb.save(cible)
    mod.XLSX_CND = cible
    mod.XLSX_CND_SHA256 = hashlib.sha256(cible.read_bytes()).hexdigest()
    mod.XLSX = XLSX118_ORIG
    if muter_118:
        c118 = bac / ("ia118_" + nom + ".xlsx")
        shutil.copy2(SRC_118, c118)
        wb118 = openpyxl.load_workbook(c118)
        muter_118(wb118)
        wb118.save(c118)
        mod.XLSX = c118
    vrai_fail = mod.fail
    try:
        mod.fail = lambda m: (_ for _ in ()).throw(RuntimeError(m))
        try:
            mod.lire_cnd({})
            return False, "aucune garde n'a bronche"
        except RuntimeError as e:
            return True, str(e)
    finally:
        mod.fail = vrai_fail
        mod.XLSX = XLSX118_ORIG


def main() -> int:
    if not SRC_CND.is_file():
        print("ECHEC : %s absent" % SRC_CND)
        return 1
    bac = Path(tempfile.mkdtemp(prefix="gardes_f6_"))
    tous = True
    try:
        print("Chaque mutation doit faire echouer la generation, PAR LA GARDE VISEE :")
        for nom, muter, muter118, attendu in CAS:
            mord, msg = eprouver(bac, nom.split()[0], muter, muter118)
            bonne = mord and attendu in msg
            tous = tous and bonne
            etat = ("mord, bonne garde" if bonne else
                    ("MORD MAIS AUTRE GARDE" if mord else "NE MORD PAS"))
            print("  %-24s %-22s %s" % (nom, etat, msg[:110]))
    finally:
        shutil.rmtree(bac, ignore_errors=True)
    print()
    print("OK  les quatre gardes mordent, chacune sur son motif" if tous
          else "ECHEC : au moins une garde n'est pas demontree")
    return 0 if tous else 1


if __name__ == "__main__":
    sys.exit(main())
