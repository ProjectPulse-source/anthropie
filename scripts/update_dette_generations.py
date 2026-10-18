#!/usr/bin/env python3
"""update_dette_generations.py -- prolongement « La dette publique est-elle un fardeau pour les générations futures ? »

Test décisif avant la page (protocole écrit avant le calcul, verdict, pièces) :
D:\\PRO\\06_PROMOTION\\RECHERCHE_GENERATIONS_FUTURES_2026-10-03\\test_decisif\\ (PROTOCOLE.md, VERDICT.md, 03/10/2026).
Les trois arguments qui allègent le fardeau, chacun confronté à la série qui peut le démentir :

  A1. « On transmet aussi des actifs » : à quoi a correspondu le besoin de financement des administrations publiques
      (Eurostat gov_10a_main, S13). Identité des comptes, par année :
          -B9 = -B8N + (P51G - P51C + P52_P53 + NP) + (D9PAY - D9REC)
          besoin de financement = désépargne nette + acquisitions nettes d'actifs non financiers + transferts en capital nets
      TÉMOIN 1 (bloquant pour la France) : les ratios qu'Eurostat PUBLIE en % du PIB (gov_10a_main, PC_GDP), que le
      calcul n'utilise pas ; tolérance TOL ; vu mordre au test (PIB de l'année précédente : 683 rejets sur 810).
      COMPLÉMENT (autre cadre statistique, plus un témoin depuis la phase B du 11/10/2026) : le patrimoine net des APU
      dans les bilans : Insee, tableau 8.204 (pièce archivée ; actifs non financiers N1N et N2N, valeur financière nette),
      témoins Eurostat nama_10_nfa_bs et nasq_10_f_bs (T4) ; la série annuelle nasa_10_f_bs n'est plus lue (17/10/2026). Un stock aux prix courants, dominé par des réévaluations, monte ou
      baisse selon la fenêtre : il ne valide pas une identité de flux.
      La décomposition dit la COMPOSITION COMPTABLE des déficits, non la charge transmise (Auerbach, Gokhale et
      Kotlikoff 1991 : le déficit dépend de l'étiquetage) ; la prose ne dit jamais « financé ».
  A2. « On se la doit à nous-mêmes » : dette des APU par secteur détenteur (BCE, statistiques de finances publiques,
      GFS), 1995-dernière année ; contrôle de cohérence contre Eurostat gov_10dd_ggd (même déclarant : cohérence, pas
      indépendance). TÉMOIN DE PÉRIMÈTRE État : Banque de France, Webstat, série DET (CSV de la page publique, l'API ne
      la sert pas aux scripts) ; dernière observation : relevé daté du graphique de l'AFT (source Banque de France),
      gardé contre une erreur de transcription et contre la péremption (garde de fraîcheur).
  A3. « Le vrai fardeau est le vieillissement » : indicateur S2 de la Commission (Debt Sustainability Monitor), position
      budgétaire initiale contre effet des variations futures des dépenses liées à l'âge. COR : présenté À PART, comme
      une sensibilité aux hypothèses, jamais substitué à une composante de S2 (substitution chiffrée retirée le 11/10/2026,
      contre-expertise de phase A). Pièces sans API, archivées dans
      scripts/sources_generations/ et vérifiées par leur empreinte (SHA256SUMS).

Une garde numérique protège contre une variation des données, jamais contre une interprétation : une phrase dont le
défaut est conceptuel se corrige dans la prose, pas par une garde (même arbitrage).

Chaque qualificatif de la page est une garde (fonction gardes) ; une donnée qui dément une phrase arrête tout, rien
n'est écrit. Bilingue depuis le 07/10/2026 (page anglaise /en/public-debt-future-generations/, sur le modèle de
update_dette_baisse.py) : un calcul, deux blocs (affichage, affichage_en) aux mêmes clés, un tableau ue.tableau_en et
des figures -en au même dessin ; les gardes ne lisent que des nombres, une passe vaut pour les deux langues.

MISE À JOUR. Séries Eurostat et BCE : workflow dette-monde.yml (mêmes notifications d'avril et d'octobre).
Pièces : à la main, à chaque édition — Debt Sustainability Monitor (publication annuelle, début d'année), rapport
annuel du COR (juin), série Webstat DET et relevé de l'AFT (chaque trimestre ; le générateur s'arrête si Webstat
rattrape le relevé, qu'il faut alors remplacer ou retirer). Geste : déposer la nouvelle pièce dans scripts/sources_generations/, mettre à jour SOURCES
ci-dessous et SHA256SUMS, relancer ; si une garde refuse, réécrire la phrase, jamais desserrer la garde.

PHASE B DE LA REFONTE (11/10/2026, arbitrage ARBITRAGE_CE_GENERATIONS_PHASE_A.md) : trois éclairages partiels, aucun n'est
une facture. Le bilan (patrimoine net) n'est plus un témoin de concordance : un complément replié, à deux bornes et
réévaluations dites (gardes « bilan »). S2 : la position de départ contient la dépense actuelle ; le second terme est
l'EFFET DES VARIATIONS FUTURES des dépenses liées à l'âge (définitions relues dans le rapport, Debt Sustainability Monitor
2025, § 3.1). La substitution chiffrée COR -> S2 est retirée (page, données affichées, gardes). Gardes de TEXTE (gardes_texte) :
elles lisent la page FR et EN et refusent une conclusion qui rend le bilan probant, un terme de vieillissement sans
« évolution future », le vocabulaire proscrit. Pied commun et versions téléphone (pied_figure), langue générée
(langue_generee). Mode --hors-ligne : la refonte se fait à données constantes (relevé du 09/10), les séries en ligne sont
relues dans data/dette_generations.json, les pièces archivées sont relues et comparées au jeu.

Usage : python scripts/update_dette_generations.py [--check] [--hors-ligne]
        [--mutation-garde | --mutation=bilan|bilan_mot|vieillissement|proscrit|cote|deux_temps|deux_temps_2]
Sorties : data/ et static/dette_generations.json, static/dette_generations.csv, data/figures_generations.json,
          static/img/dette-generations-{actifs,detention,vieillissement}{,-en}{,-m}.svg + .png
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import sys
import time
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import langue_generee as lg   # mots autour d'un nombre (scripts/langue_generee.py)
import releve_jeu as rj  # noqa: E402  -- relevé daté par les seules données (organe commun, 17/10/2026)
import lecture_tableur as lt  # forme des cellules (organe commun) : année ou nombre écrits en texte
import pied_figure            # pied commun, version téléphone, parité (scripts/pied_figure.py)

ROOT = Path(__file__).resolve().parent.parent
PAGE_FR = ROOT / "content" / "dette-publique-generations-futures" / "_index.md"
PAGE_EN = ROOT / "content" / "dette-publique-generations-futures" / "_index.en.md"
BAISSE = ROOT / "data" / "dette_baisse.json"   # volet 5 : écart au solde stabilisant (lu, jamais réécrit ici)
COLL = ROOT / "data" / "dette_collectivites.json"   # fenêtre de la baisse des transferts aux collectivités (lue, jamais réécrite)
# Europe : deux pays nommés de chaque côté de la France (phrase « ni universel ni exceptionnel », gardée)
CITE_SUP, CITE_INF = ("RO", "SI"), ("IT", "SK")
OUT_DATA = ROOT / "data" / "dette_generations.json"
OUT_STATIC = ROOT / "static" / "dette_generations.json"
OUT_CSV = ROOT / "static" / "dette_generations.csv"
OUT_FIGURES = ROOT / "data" / "figures_generations.json"
OUT_IMG = ROOT / "static" / "img"
PIECES = ROOT / "scripts" / "sources_generations"
PAGE_URL = "stephane-lalut.com/dette-publique-generations-futures/"
PAGE_URL_EN = "stephane-lalut.com/en/public-debt-future-generations/"

# Pièces sans API : fichier, libellé de l'édition (dit dans la page par jeton, jamais saisi dans la prose)
SOURCES = {
    "dsm": dict(fichier="dsm2025_country_fiches.xlsx", lib="Debt Sustainability Monitor 2025", publie="2026-02",
                producteur="Commission européenne, DG ECFIN, Institutional Paper 332, février 2026",
                url="https://economy-finance.ec.europa.eu/publications/debt-sustainability-monitor-2025_en"),
    "cor": dict(fichier="cor_ra2026_synthese.xlsx", lib="rapport annuel de juin 2026", lib_en="annual report of June 2026", publie="2026-06",
                producteur="Conseil d'orientation des retraites, données de la synthèse, scénario de référence",
                url="https://www.cor-retraites.fr/rapports-du-cor/rapport-annuel-cor-juin-2026-evolutions-perspectives-retraites-france"),
    # Témoin de périmètre État (règle constante de l'auteur, 03/10/2026 : anti-robot -> changer de canal, jamais abandonner
    # la donnée). L'API Webstat ne sert pas ces observations aux scripts ; le CSV est celui du bouton « Télécharger les
    # données » de la page publique de la série (format long), archivé tel quel. Mise à jour : trimestrielle, à la main.
    "det": dict(fichier="webstat_DET.Q.FR.1315.F33000.M.Z9.8.F.csv", lib="Banque de France, série DET.Q.FR.1315.F33000.M.Z9.8.F",
                lib_en="Banque de France, series DET.Q.FR.1315.F33000.M.Z9.8.F",
                producteur="Banque de France, Webstat : détention par les non-résidents de la dette négociable de l'État (en %), valeur de marché",
                url="https://webstat.banque-france.fr/fr/catalogue/det/DET.Q.FR.1315.F33000.M.Z9.8.F"),
    # Dernière observation (arbitrage PRO-20261003-195656, B3) : relevé daté du graphique de l'AFT, source Banque de
    # France ; l'image n'est pas servie aux scripts, ses deux étiquettes sont transcrites dans la pièce avec son empreinte.
    # Patrimoine des APU (17/10/2026) : le producteur, et non plus la série annuelle Eurostat nasa_10_f_bs, qui donnait en
    # octobre 2026 une valeur financière nette de -2 134 Md€ fin 2025 contre -2 195 à l'Insee ET dans la série trimestrielle
    # d'Eurostat : le patrimoine net « montait » en 2025 sur la page alors qu'il baisse. Mise à jour : annuelle, à la main
    # (registre data/sources_maj.json, id « insee-patrimoine-apu »).
    "pat": dict(fichier="insee_T_8204_patrimoine_apu_2026-07-30.xlsx", lib="Insee, tableau 8.204 (30 juillet 2026)",
                lib_en="Insee, table 8.204 (30 July 2026)",
                producteur="Insee, comptes nationaux annuels (base 2020), compte de patrimoine des administrations publiques (S13), corrigé le 30/07/2026",
                url="https://www.insee.fr/fr/statistiques/8988859"),
    "aft": dict(fichier="aft_non_residents_releve_2026T1.json", lib="Agence France Trésor, d'après la Banque de France",
                producteur="Agence France Trésor, page « Principaux chiffres », graphique des non-résidents (source Banque de France)",
                url="https://www.aft.gouv.fr/fr/principaux-chiffres-dette"),
}

TOL = 0.11            # ratios publiés à une décimale
TOL_BOUCLAGE = 0.05   # identité des comptes, par année, en points de PIB
SEUIL_BESOIN = 30.0   # contre-population : besoin de financement cumulé minimal pour comparer la part en actifs
API = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/%s?format=JSON&%s"
ECB = "https://data-api.ecb.europa.eu/service/data/GFS/%s?format=csvdata&detail=dataonly"

W = 720
FONT = "Inter, 'Helvetica Neue', Arial, sans-serif"
BLEU, ORANGE, GRIS, GRIS_CLAIR = "#184f95", "#eb6834", "#8a8781", "#c9c5c0"
INK, INK2, MUTED, GRID = "#26262f", "#55524f", "#96928f", "#dcd8d3"
LICENCES = {"fr": "Compilation Stéphane Lalut, CC BY 4.0 · " + PAGE_URL, "en": "Compiled by Stéphane Lalut, CC BY 4.0 · " + PAGE_URL_EN}

PAYS = {"AT": ("Autriche", "l'Autriche"), "BE": ("Belgique", "la Belgique"), "BG": ("Bulgarie", "la Bulgarie"),
        "CY": ("Chypre", "Chypre"), "CZ": ("Tchéquie", "la Tchéquie"), "DE": ("Allemagne", "l'Allemagne"),
        "DK": ("Danemark", "le Danemark"), "EE": ("Estonie", "l'Estonie"), "EL": ("Grèce", "la Grèce"),
        "ES": ("Espagne", "l'Espagne"), "FI": ("Finlande", "la Finlande"), "FR": ("France", "la France"),
        "HR": ("Croatie", "la Croatie"), "HU": ("Hongrie", "la Hongrie"), "IE": ("Irlande", "l'Irlande"),
        "IT": ("Italie", "l'Italie"), "LT": ("Lituanie", "la Lituanie"), "LU": ("Luxembourg", "le Luxembourg"),
        "LV": ("Lettonie", "la Lettonie"), "MT": ("Malte", "Malte"), "NL": ("Pays-Bas", "les Pays-Bas"),
        "PL": ("Pologne", "la Pologne"), "PT": ("Portugal", "le Portugal"), "RO": ("Roumanie", "la Roumanie"),
        "SE": ("Suède", "la Suède"), "SI": ("Slovénie", "la Slovénie"), "SK": ("Slovaquie", "la Slovaquie")}
TRIM = ["premier trimestre", "deuxième trimestre", "troisième trimestre", "quatrième trimestre"]
MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
LETTRES = ["zéro", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit", "neuf", "dix", "onze", "douze",
           "treize", "quatorze", "quinze", "seize"]
PAYS_EN = {"AT": "Austria", "BE": "Belgium", "BG": "Bulgaria", "CY": "Cyprus", "CZ": "Czechia", "DE": "Germany", "DK": "Denmark",
           "EE": "Estonia", "EL": "Greece", "ES": "Spain", "FI": "Finland", "FR": "France", "HR": "Croatia", "HU": "Hungary",
           "IE": "Ireland", "IT": "Italy", "LT": "Lithuania", "LU": "Luxembourg", "LV": "Latvia", "MT": "Malta",
           "NL": "the Netherlands", "PL": "Poland", "PT": "Portugal", "RO": "Romania", "SE": "Sweden", "SI": "Slovenia",
           "SK": "Slovakia"}
TRIM_EN = ["first quarter of", "second quarter of", "third quarter of", "fourth quarter of"]
MOIS_EN = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
LETTRES_EN = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten", "eleven", "twelve",
              "thirteen", "fourteen", "fifteen", "sixteen"]


def log(msg: str) -> None:
    print(msg.encode("ascii", "replace").decode("ascii"))


def fail(msg: str) -> None:
    log("ECHEC : " + msg)
    log("Aucun fichier ecrit.")
    sys.exit(1)


def fr(v: float, dec: int = 1) -> str:
    s = ("%." + str(dec) + "f") % v
    if float(s) == 0:
        s = s.replace("-", "")
    return s.replace("-", "−").replace(".", ",")


def signe(v: float, dec: int = 1) -> str:
    s = fr(v, dec)
    return s if s.startswith("−") else "+" + s


def lettres(n: int) -> str:
    return LETTRES[n] if 0 <= n < len(LETTRES) else str(n)


def enumere(noms: list[str]) -> str:
    return noms[0] if len(noms) == 1 else ", ".join(noms[:-1]) + " et " + noms[-1]


def maj(s: str) -> str:
    return s[:1].upper() + s[1:]


# ---- Deux langues, un seul calcul (modèle : update_dette_baisse.py) : use(lang) rebranche les formats ; les gardes, elles,
# ne lisent que des nombres. Les chaînes anglaises : point décimal, pas de séparateur insécable devant %, noms anglais.
FR_FMT = dict(fr=fr, signe=signe, lettres=lettres, enumere=enumere)


def _en_nb(v: float, dec: int = 1) -> str:
    s_ = ("%." + str(dec) + "f") % v
    if float(s_) == 0:
        s_ = s_.replace("-", "")
    return s_.replace("-", "−")


def _en_signe(v: float, dec: int = 1) -> str:
    s_ = _en_nb(v, dec)
    return s_ if s_.startswith("−") else "+" + s_


EN_FMT = dict(fr=_en_nb, signe=_en_signe,
              lettres=lambda n: LETTRES_EN[n] if 0 <= n < len(LETTRES_EN) else str(n),
              enumere=lambda noms: noms[0] if len(noms) == 1 else ", ".join(noms[:-1]) + " and " + noms[-1])
LANG = "fr"
SUFFIXE = {"fr": "", "en": "-en"}


def use(lang: str) -> None:
    global LANG
    LANG = lang
    globals().update(FR_FMT if lang == "fr" else EN_FMT)


def art(code: str) -> str:
    """Nom de pays dans la prose (avec article en français)."""
    return PAYS[code][1].replace("'", "’") if LANG == "fr" else PAYS_EN[code]  # apostrophe typographique : un jeton échappe au traitement de Hugo


def nom_pays(code: str) -> str:
    """Nom de pays seul (tableau)."""
    n = PAYS[code][0] if LANG == "fr" else PAYS_EN[code]
    return n[:1].upper() + n[1:]


def lib(cle: str) -> str:
    """Libellé d'une pièce dans la langue courante (lib_en s'il existe : un titre anglais d'origine n'a pas de traduction)."""
    return SOURCES[cle].get("lib_en", SOURCES[cle]["lib"]) if LANG == "en" else SOURCES[cle]["lib"]


def date_longue(iso: str) -> str:
    """AAAA-MM-JJ -> « 6 octobre 2026 » / « 6 October 2026 »."""
    m = (MOIS if LANG == "fr" else MOIS_EN)[int(iso[5:7]) - 1]
    return iso[8:10].lstrip("0") + " " + m + " " + iso[:4]


def pc() -> str:
    """Signe pourcent après un nombre : insécable en français, collé en anglais."""
    return " %" if LANG == "fr" else "%"


# ------------------------------------------------------------------ données
def fetch(url: str, essais: int = 4) -> bytes:
    for k in range(essais):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "stephane-lalut.com (stephane@stephane-lalut.com)"})
            with urllib.request.urlopen(req, timeout=240) as r:
                return r.read()
        except Exception as e:  # noqa: BLE001
            if k == essais - 1:
                fail("reseau : %s (%s)" % (url[:120], e))
            time.sleep(15 * (k + 1))   # le portail de la BCE répond parfois 504 : on attend, on ne contourne pas
    return b""


def eurostat(ds: str, **f) -> dict[tuple[str, int], float]:
    """Tous les pays d'un jeu, une seule valeur par (pays, année) : toute autre dimension doit être fixée."""
    d = json.loads(fetch(API % (ds, "&".join("%s=%s" % kv for kv in f.items()))))
    if "dimension" not in d:
        fail("Eurostat %s : reponse sans dimensions" % ds)
    ids, size = d["id"], d["size"]
    for dim, n in zip(ids, size):
        if dim not in ("geo", "time") and n != 1:
            fail("Eurostat %s : dimension %s non fixee (%d valeurs)" % (ds, dim, n))
    if ids[-2:] != ["geo", "time"]:
        fail("Eurostat %s : ordre des dimensions inattendu (%s)" % (ds, ids))
    nt = size[-1]
    out = {}
    for g, ig in d["dimension"]["geo"]["category"]["index"].items():
        for t, it in d["dimension"]["time"]["category"]["index"].items():
            v = d["value"].get(str(ig * nt + it))
            if v is not None:
                out[(g, int(t))] = float(v)
    return out


def bce(cle: str) -> dict[tuple[str, str], dict[int, float]]:
    """GFS de la BCE : {(pays, secteur détenteur): {année: % du PIB}}."""
    raw = fetch(ECB % cle).decode("utf-8")
    if not raw.startswith("KEY"):
        fail("BCE %s : reponse qui n'est pas un CSV de donnees" % cle)
    out: dict = {}
    for x in csv.DictReader(io.StringIO(raw)):
        if x["OBS_VALUE"]:
            out.setdefault((x["REF_AREA"], x["COUNTERPART_SECTOR"]), {})[int(x["TIME_PERIOD"])] = float(x["OBS_VALUE"])
    return out


def piece(cle: str) -> Path:
    """Une pièce archivée n'est lue que si son empreinte est celle de SHA256SUMS."""
    f = PIECES / SOURCES[cle]["fichier"]
    sums = dict(reversed(l.split(None, 1)) for l in (PIECES / "SHA256SUMS").read_text(encoding="utf-8").split("\n") if l.strip())
    if not f.is_file():
        fail("piece absente : %s" % f)
    h = hashlib.sha256(f.read_bytes()).hexdigest()
    if sums.get(f.name) != h:
        fail("piece %s : empreinte %s differente de SHA256SUMS" % (f.name, h[:16]))
    return f


# ------------------------------------------------------------------ A1 : à quoi ont correspondu les déficits
ITEMS = ("B9", "P51G", "P51C", "B8N", "D9PAY", "D9REC", "D92PAY", "NP", "P52_P53")
PUBLIES = ("B9", "P51G", "P51C", "B8N")


def emplois():
    M = {it: eurostat("gov_10a_main", na_item=it, sector="S13", unit="MIO_NAC") for it in ITEMS}
    pub = {it: eurostat("gov_10a_main", na_item=it, sector="S13", unit="PC_GDP") for it in PUBLIES}
    Y = eurostat("gov_10dd_edpt1", na_item="B1GQ", sector="S1", unit="MIO_NAC")
    R, ecartes, calculables = {}, [], 0
    for (p, a) in sorted(M["B9"]):
        if p not in PAYS or (p, a) not in Y or any((p, a) not in M[it] for it in ("P51G", "P51C", "B8N", "D9PAY", "D9REC")):
            continue
        calculables += 1
        y = Y[(p, a)]
        r = {it: M[it].get((p, a), 0.0) / y * 100 for it in ITEMS}
        if not all((p, a) in pub[it] for it in PUBLIES):
            ecartes.append(dict(pays=p, annee=a, motif="ratio publié absent"))
            continue
        e = {it: r[it] - pub[it][(p, a)] for it in PUBLIES}
        if any(abs(v) > TOL for v in e.values()):
            ecartes.append(dict(pays=p, annee=a, motif="écart au ratio publié", ecarts={k: round(v, 2) for k, v in e.items() if abs(v) > TOL}))
            continue
        B, E = -r["B9"], -r["B8N"]
        N = r["P51G"] - r["P51C"]
        K = N + r["NP"] + r["P52_P53"]          # acquisitions nettes d'actifs non financiers, usure déduite
        T = r["D9PAY"] - r["D9REC"]
        bou = B - (E + K + T)
        if abs(bou) > TOL_BOUCLAGE:
            ecartes.append(dict(pays=p, annee=a, motif="identité des comptes non bouclée", ecart=round(bou, 3)))
            continue
        R.setdefault(p, {})[a] = dict(besoin=B, desepargne=E, investissement_net=N, actifs=K, transferts=T,
                                      aides_investissement=r["D92PAY"], d92_publie=(p, a) in M["D92PAY"],
                                      besoin_meur=-M["B9"][(p, a)], actifs_meur=M["P51G"][(p, a)] - M["P51C"][(p, a)]
                                      + M["NP"].get((p, a), 0.0) + M["P52_P53"].get((p, a), 0.0))
    retenues = sum(len(v) for v in R.values())
    if calculables != retenues + len(ecartes):
        fail("conservation rompue : %d calculables, %d retenues, %d ecartees" % (calculables, retenues, len(ecartes)))
    return R, ecartes, calculables


def cumul(an: dict, a0: int, a1: int):
    if not all(a in an for a in range(a0, a1 + 1)):
        return None
    L = [an[a] for a in range(a0, a1 + 1)]
    s = {k: sum(x[k] for x in L) for k in ("besoin", "desepargne", "investissement_net", "actifs", "transferts", "aides_investissement",
                                            "besoin_meur", "actifs_meur")}
    s.update(debut=a0, fin=a1, annees=len(L))
    if s["besoin"] > 0:
        s.update(part_actifs=100 * s["actifs"] / s["besoin"], part_desepargne=100 * s["desepargne"] / s["besoin"],
                 part_transferts=100 * s["transferts"] / s["besoin"], part_actifs_euros=100 * s["actifs_meur"] / s["besoin_meur"],
                 part_investissement_net=100 * s["investissement_net"] / s["besoin"])
    return s


# ---- Où le recul de l'investissement s'est produit (18/10/2026, piste du test d'abstention de la synthèse du dossier) :
# investissement net (P51G - P51C) des administrations locales (S1313) et des autres administrations, France.
def sous_secteurs(annees: list) -> dict:
    """En % du PIB de la notification. Témoin : l'investissement net S13 relu ici redonne, à 0,02 point près, celui que
    calcule emplois() ; la FBCF locale ne dépasse jamais celle de S13. Une consommation de capital fixe de la dernière
    année identique à celle de l'année précédente (valeur reportée par le producteur) est signalée dans la page."""
    M = {(sec, it): eurostat("gov_10a_main", geo="FR", na_item=it, sector=sec, unit="MIO_NAC")
         for sec in ("S13", "S1313") for it in ("P51G", "P51C")}
    Y = eurostat("gov_10dd_edpt1", geo="FR", na_item="B1GQ", sector="S1", unit="MIO_NAC")
    out = []
    for r in annees:
        a = r["annee"]
        if any(("FR", a) not in M[k] for k in M) or ("FR", a) not in Y:
            fail("sous-secteurs : %d absent d'une serie" % a)
        v = {k: M[k][("FR", a)] / Y[("FR", a)] * 100 for k in M}
        apu, loc = v[("S13", "P51G")] - v[("S13", "P51C")], v[("S1313", "P51G")] - v[("S1313", "P51C")]
        if abs(apu - r["investissement_net"]) > 0.02:
            fail("sous-secteurs : investissement net S13 %d = %.3f, jeu = %.3f (temoin)" % (a, apu, r["investissement_net"]))
        if v[("S1313", "P51G")] > v[("S13", "P51G")]:
            fail("sous-secteurs : FBCF locale superieure a celle des APU en %d" % a)
        out.append(dict(annee=a, net_apu=apu, net_loc=loc, net_autres=apu - loc, fbcf_loc=v[("S1313", "P51G")]))
    fin = annees[-1]["annee"]
    reprise = M[("S13", "P51C")][("FR", fin)] == M[("S13", "P51C")][("FR", fin - 1)]
    return dict(annees=out, ccf_reprise=reprise)


def insee_patrimoine() -> dict[str, dict[int, float]]:
    """Insee, tableau 8.204 (compte de patrimoine des APU), en millions d'euros : {code: {année: valeur}}. Les deux lignes
    « F » (actifs, puis passifs financiers) sont nommées FA et FP ; refus si une ligne attendue manque."""
    import openpyxl
    wb = openpyxl.load_workbook(piece("pat"), read_only=True, data_only=True)
    if "T_8204" not in wb.sheetnames:
        fail("Insee 8.204 : onglet T_8204 absent")
    rows = [r for r in wb["T_8204"].iter_rows(values_only=True)]
    hdr = next((r for r in rows if r and len(r) > 2 and r[2] == "1959"), None)
    if hdr is None:
        fail("Insee 8.204 : ligne des annees introuvable")
    cols = {j: int(c) for j, c in enumerate(hdr) if isinstance(c, str) and c.isdigit()}
    out: dict[str, dict[int, float]] = {}
    for r in rows:
        code, lib = (r[0] or ""), (r[1] or "")
        cle = {"F": "FA" if "actifs" in lib else "FP" if "passifs" in lib else None}.get(code, code)
        if cle in ("N1N", "N2N", "N211N", "FA", "FP", "B90") and cle not in out:
            out[cle] = {a: float(r[j]) * 1000 for j, a in cols.items() if j < len(r) and isinstance(r[j], (int, float))}
    if set(out) != {"N1N", "N2N", "N211N", "FA", "FP", "B90"}:
        fail("Insee 8.204 : lignes manquantes %s" % sorted({"N1N", "N2N", "N211N", "FA", "FP", "B90"} - set(out)))
    return out


def bf90_trimestriel() -> dict[int, float]:
    """Témoin : valeur financière nette des APU au 4e trimestre (Eurostat nasq_10_f_bs, consolidée), {année: M€}."""
    d = json.loads(fetch(API % ("nasq_10_f_bs", "geo=FR&sector=S13&unit=MIO_NAC&na_item=BF90&finpos=LIAB")))
    if "dimension" not in d or any(n != 1 for dim, n in zip(d["id"], d["size"]) if dim != "time"):
        fail("Eurostat nasq_10_f_bs : reponse inattendue")
    t = d["dimension"]["time"]["category"]["index"]
    return {int(k[:4]): float(d["value"][str(i)]) for k, i in t.items() if k.endswith("Q4") and str(i) in d["value"]}


def patrimoine(fin: int):
    Y = eurostat("gov_10dd_edpt1", na_item="B1GQ", sector="S1", unit="MIO_NAC")
    # Tous les actifs non financiers (arbitrage PRO-20261003-195656, B2) : produits N1N (actifs fixes, stocks, objets de
    # valeur) et non produits N2N (terrains surtout) ; les terrains seuls (N211N) gardés pour la phrase sur leur valeur.
    # Source : l'Insee, producteur (tableau 8.204, pièce archivée). Témoins bloquants : les actifs non financiers transmis
    # à Eurostat (nama_10_nfa_bs) et la valeur financière nette trimestrielle d'Eurostat au 4e trimestre. La série annuelle
    # nasa_10_f_bs n'est plus lue : elle s'écartait des deux autres en octobre 2026 (voir SOURCES["pat"]).
    I = insee_patrimoine()
    prod, nonprod, terrains = ({("FR", a): v for a, v in I[k].items()} for k in ("N1N", "N2N", "N211N"))
    bf = {("FR", a): I["FA"][a] - I["FP"][a] for a in I["FA"] if a in I["FP"]}
    for a in I["B90"]:
        if ("FR", a) in bf and ("FR", a) in prod and ("FR", a) in nonprod:
            if abs(prod[("FR", a)] + nonprod[("FR", a)] + bf[("FR", a)] - I["B90"][a]) > 500:
                fail("Insee 8.204 %d : la valeur nette ne boucle pas (actifs non financiers + actifs - passifs financiers)" % a)
    e_prod = eurostat("nama_10_nfa_bs", sector="S13", unit="CP_MNAC", asset10="N1N")
    e_nonprod = eurostat("nama_10_nfa_bs", sector="S13", unit="CP_MNAC", asset10="N2N")
    communs = [a for a in range(fin - 30, fin + 1) if ("FR", a) in e_prod and ("FR", a) in prod]
    if len(communs) < 10:
        fail("temoin Eurostat nama_10_nfa_bs : moins de dix annees communes avec l'Insee")
    for a in communs:
        if abs(e_prod[("FR", a)] - prod[("FR", a)]) > 1000 or abs(e_nonprod[("FR", a)] - nonprod[("FR", a)]) > 1000:
            fail("Insee 8.204 et Eurostat divergent sur les actifs non financiers en %d : millesimes differents, relire" % a)
    q4 = bf90_trimestriel()
    cq = [a for a in q4 if ("FR", a) in bf and a >= fin - 3]
    if len(cq) < 2:
        fail("temoin Eurostat nasq_10_f_bs : moins de deux annees communes recentes")
    for a in cq:
        if abs(q4[a] - bf[("FR", a)]) > 2000:
            fail("Insee 8.204 et Eurostat trimestriel divergent sur la valeur financiere nette %d (%.0f contre %.0f M EUR)" % (a, bf[("FR", a)], q4[a]))
    ans = [a for a in range(fin - 30, fin + 1) if all(("FR", a) in s for s in (Y, prod, nonprod, terrains, bf))]
    if not ans:
        fail("patrimoine : aucune annee complete pour la France")
    out = {a: dict(annee=a, actifs_produits=prod[("FR", a)] / Y[("FR", a)] * 100, actifs_non_produits=nonprod[("FR", a)] / Y[("FR", a)] * 100,
                   terrains=terrains[("FR", a)] / Y[("FR", a)] * 100, valeur_financiere_nette=bf[("FR", a)] / Y[("FR", a)] * 100,
                   patrimoine_net_meur=prod[("FR", a)] + nonprod[("FR", a)] + bf[("FR", a)]) for a in ans}
    for x in out.values():
        x["patrimoine_net"] = x["actifs_produits"] + x["actifs_non_produits"] + x["valeur_financiere_nette"]
        if x["terrains"] > x["actifs_non_produits"] + 1e-9:
            fail("patrimoine %d : terrains superieurs aux actifs non produits" % x["annee"])
    return [out[a] for a in ans]


# ------------------------------------------------------------------ A2 : qui détient la dette
def detention():
    tot = bce("A.N..W0.S13.S1.C.L.LE.GD.T._Z.XDC_R_B1GQ._T.F.V.N._T")
    nr = bce("A.N..W1.S13.S1.C.L.LE.GD.T._Z.XDC_R_B1GQ._T.F.V.N._T")
    res = bce("A.N..W2.S13..C.L.LE.GD.T._Z.XDC_R_B1GQ._T.F.V.N._T")
    fr_ = []
    for a in sorted(tot.get(("FR", "S1"), {})):
        t, n, r, b = tot[("FR", "S1")][a], nr.get(("FR", "S1"), {}).get(a), res.get(("FR", "S1"), {}).get(a), res.get(("FR", "S121"), {}).get(a)
        if None in (n, r, b):
            fail("BCE : detention francaise incomplete en %d" % a)
        if abs(n + r - t) > 0.15:
            fail("BCE : non-residents + residents != total en %d (%.2f)" % (a, n + r - t))
        fr_.append(dict(annee=a, dette=t, non_residents=n, residents=r, banque_centrale=b, autres_residents=r - b,
                        part_non_residents=100 * n / t, part_banque_centrale=100 * b / t, part_autres_residents=100 * (r - b) / t))
    # Cohérence avec Eurostat (même déclarant : un écart signale une clé ou un millésime cassé, pas une contre-preuve)
    ggd = {s: eurostat("gov_10dd_ggd", na_item="GD", sector="S13", unit="PC_GDP", maturity="TOTAL", sector2=s) for s in ("S2", "S121")}
    vus = 0
    for x in fr_:
        a = x["annee"]
        if ("FR", a) in ggd["S2"]:
            vus += 1
            if abs(ggd["S2"][("FR", a)] - x["non_residents"]) > 0.15 or abs(ggd["S121"][("FR", a)] - x["banque_centrale"]) > 0.15:
                fail("BCE et Eurostat divergent sur la detention francaise en %d" % a)
    if vus < 3:
        fail("controle BCE/Eurostat : moins de trois annees communes")
    pays = {}
    for p in PAYS:
        t, n = tot.get((p, "S1"), {}), nr.get((p, "S1"), {})
        if t and n:
            a = max(set(t) & set(n))
            pays[p] = dict(annee=a, part_non_residents=100 * n[a] / t[a])
    return fr_, pays, vus


def etat_negociable():
    """Témoin de périmètre : part des titres négociables de l'État détenue par des non-résidents (Banque de France,
    valeur de marché), fin de chaque année ; autre champ (État, titres négociables, valeur de marché) que la série BCE
    (administrations publiques, toute la dette, valeur nominale)."""
    rows = list(csv.DictReader(io.StringIO(piece("det").read_text(encoding="utf-8-sig")), delimiter=";"))
    if not rows or rows[0].get("series_key") != "DET.Q.FR.1315.F33000.M.Z9.8.F" or rows[0].get("UNIT") != "PC":
        fail("piece Webstat DET : serie ou unite inattendue (le fichier n'est pas celui attendu)")
    obs = {x["time_period"]: float(x["obs_value"].replace(",", ".")) for x in rows if x["obs_value"]}
    dernier = max(obs)
    fins = [dict(annee=int(k[:4]), part=v) for k, v in sorted(obs.items()) if k.endswith("Q4")]
    if not 0 < fins[-1]["part"] < 100 or len(fins) < 20:
        fail("piece Webstat DET : moins de vingt fins d'annee ou valeur hors bornes")
    rel = json.loads(piece("aft").read_text(encoding="utf-8"))
    robs = {k: float(v) for k, v in rel["observations"].items()}
    rder = max(robs)
    # Garde de transcription : la période commune doit égaler Webstat à 0,1 point près.
    communs = [k for k in robs if k in obs]
    if not communs or any(abs(robs[k] - obs[k]) > 0.1 for k in communs):
        fail("releve AFT : la periode commune ne concorde pas avec Webstat (%s)" % {k: (robs[k], obs.get(k)) for k in robs})
    # Garde de fraîcheur : le relevé n'a de raison d'être que s'il est plus récent que Webstat.
    if rder <= dernier:
        fail("releve AFT perime : Webstat publie %s, le releve s'arrete a %s ; remplacer ou retirer le releve" % (dernier, rder))
    return dict(fins=fins, dernier=dernier, valeur_dernier=obs[dernier], maj=rows[0].get("updated_at", "")[:10],
                releve=dict(periode=rder, valeur=robs[rder], lu_le=rel["lu_le"], image_sha256=rel["image_sha256"]))


# ------------------------------------------------------------------ A3 : vieillissement ou position présente
JEUX = ("base", "productivite", "risque", "base_precedente")


def s2():
    import openpyxl
    wb = openpyxl.load_workbook(piece("dsm"), read_only=True, data_only=True)
    out = {}
    for g in wb.sheetnames:
        # Colonnes des quatre jeux repérées par leurs en-têtes, puis lues À LEUR POSITION (contre-expertise
        # anthropie-site-20261016-154235 : `num[:4]` après élimination des cellules illisibles laissait la cinquième
        # valeur prendre la place d'une troisième manquante).
        rows = [lt.ligne(r) for r in wb[g].iter_rows(min_row=180, max_row=200, values_only=True)]
        try:
            ib = next(i for i, r in enumerate(rows) if "Baseline" in r and "Lower productivity scenario" in r)
            cols = (lt.colonnes(rows[ib], ["Baseline", "Lower productivity scenario", "Non-demographic risk scenario"])
                    + lt.colonnes(rows[ib - 1], ["DSM 2024"]))
        except (StopIteration, ValueError) as e:
            fail("DSM, feuille %s : en-tetes des jeux introuvables (%s)" % (g, e))
        lab = {}
        for r in rows[ib + 1:]:
            if "Baseline" in r or "DSM 2024" in r:
                break                                      # en-tête du bloc suivant (S1) : hors sujet
            txt = [v for v in r[:cols[0]] if isinstance(v, str) and v.strip() and lt.nombre(v) is None]
            if not txt or txt[-1].strip() in lab or all(r[c] is None for c in cols if c < len(r)):
                continue
            try:
                lab[txt[-1].strip()] = lt.cellules(r, cols, "DSM %s, %s" % (g, txt[-1].strip()))
            except ValueError as e:
                fail(str(e))
        try:
            out[g] = {k: dict(zip(JEUX, lab[lt.ligne([lib])[0]])) for k, lib in (     # libellés normalisés des deux côtés
                ("S2", "Overall value  (% of GDP)"), ("IBP", "Initial budgetary position"), ("CoA", "Ageing costs"),
                ("pensions", "Pensions"), ("sante", "Health care"), ("dependance", "Long-term care"), ("education", "Education"))}
        except KeyError as e:
            fail("DSM, feuille %s : libelle introuvable %s (la mise en page de la piece a change)" % (g, e))
    if set(out) != set(PAYS):
        fail("DSM : feuilles pays inattendues (%s)" % sorted(set(out) ^ set(PAYS)))
    x = out["FR"]
    for j in JEUX:
        if abs(x["S2"][j] - x["IBP"][j] - x["CoA"][j]) > 0.02 or \
                abs(x["CoA"][j] - x["pensions"][j] - x["sante"][j] - x["dependance"][j] - x["education"][j]) > 0.02:
            fail("DSM France, jeu %s : S2 != position initiale + vieillissement, ou vieillissement != somme des postes" % j)
    return out


def cor():
    import openpyxl
    wb = openpyxl.load_workbook(piece("cor"), read_only=True, data_only=True)
    lignes = {}
    for nom_f, lib in (("Dépenses en %", "Sc. Ref"), ("Solde dépenses ressources", "Ressources"), ("Solde dépenses ressources", "Solde")):
        rows = list(wb[nom_f].iter_rows(values_only=True))
        hdr = next((r for r in rows if {2025, 2070} <= {lt.annee(v) for v in r}), None)
        ref = next((r for r in rows if r and any(isinstance(v, str) and v.strip() == lib for v in r)), None)
        if not hdr or not ref:
            fail("COR, feuille %s : ligne %s introuvable (la mise en page de la piece a change)" % (nom_f, lib))
        d = {lt.annee(hdr[i]): lt.valeur(ref[i]) for i in range(min(len(hdr), len(ref)))
             if lt.annee(hdr[i]) is not None and lt.valeur(ref[i]) is not None}
        lignes[lib] = {a: v * (100 if abs(d[2025]) < 1 and lib != "Solde" else 1) for a, v in d.items()}
    # Les ratios du COR sont en fraction (0,1412) : ramenés en % ; le solde aussi, s'il est en fraction.
    if abs(lignes["Solde"][2025]) < 0.05:
        lignes["Solde"] = {a: v * 100 for a, v in lignes["Solde"].items()}
    dep, ress, solde = lignes["Sc. Ref"], lignes["Ressources"], lignes["Solde"]
    if not 10 < dep[2025] < 20 or not 10 < ress[2025] < 20:
        fail("COR : ordre de grandeur inattendu (depenses %.2f, ressources %.2f)" % (dep[2025], ress[2025]))
    a1 = max(a for a in dep if a in ress and a in solde)
    if abs(ress[a1] - dep[a1] - solde[a1]) > 0.02:
        fail("COR : ressources - depenses != solde en %d" % a1)
    return dict(a0=2025, a1=a1, depenses0=dep[2025], depenses1=dep[a1], ressources0=ress[2025], ressources1=ress[a1],
                solde0=solde[2025], solde1=solde[a1])


def dsm_definitions() -> dict:
    """Définitions de S2 et de S1 telles que la pièce archivée les imprime (feuille France) : année de l'ajustement,
    horizon infini pour S2, cible et date pour S1. Relues dans le rapport le 11/10/2026 (Institutional Paper 332, § 3.1 :
    « In contrast to the debt-reduction indicator S1, neither the level at which debt stabilises nor the timing is
    predefined for the S2 indicator »). La phrase de la page qui les cite se lit ici, jamais saisie."""
    import openpyxl
    wb = openpyxl.load_workbook(piece("dsm"), read_only=True, data_only=True)
    txt = [" ".join(v.split()) for r in wb["FR"].iter_rows(values_only=True) for v in r if isinstance(v, str)]
    s2_ = next((re.match(r"S2 indicator \(effort required in (\d{4}) to stabilise debt over an infinite horizon\)$", t)
                for t in txt if t.startswith("S2 indicator (")), None)
    s1_ = next((re.match(r"S1 indicator \(effort required in (\d{4}) to bring debt to (\d+)% of GDP by (\d{4})\)$", t)
                for t in txt if t.startswith("S1 indicator (")), None)
    if not s2_ or not s1_:
        fail("DSM : definition de S2 ou de S1 introuvable dans la piece (la mise en page a change : relire le rapport)")
    return dict(s2_an=int(s2_.group(1)), s1_an=int(s1_.group(1)), s1_cible=int(s1_.group(2)), s1_horizon=int(s1_.group(3)))


def volet5() -> dict:
    """Volet 5 du dossier (« La dette publique peut-elle baisser ? ») : écart entre le solde primaire observé et le solde
    qui aurait stabilisé le ratio la dernière année. Lu dans data/dette_baisse.json (jamais recalculé ici) ; les chaînes
    affichées par le volet 5 servent de témoin : la page cite la même valeur, pour la même année."""
    if not BAISSE.is_file():
        fail("volet 5 : %s absent" % BAISSE)
    j = json.loads(BAISSE.read_text(encoding="utf-8"))
    last = j["france"]["annees"][-1]
    return dict(annee=last["annee"], ecart=last["ecart"], aff=j["affichage"], aff_en=j["affichage_en"])


def calcul_hors_ligne():
    """Refonte à données constantes : les séries en ligne (Eurostat, BCE) sont relues dans le jeu publié ; les pièces
    archivées sont relues, et doivent redonner exactement ce que le jeu en contient (sinon une pièce a changé)."""
    j = json.loads(OUT_DATA.read_text(encoding="utf-8"))
    total, *dec = j["france"]["periodes"]
    annees, pat, det = j["france"]["annees"], j["france"]["patrimoine"], j["france"]["detention"]
    cp, det_pays = j["ue"]["emplois"], j["ue"]["part_non_residents"]
    S2, C, etat = s2(), cor(), etat_negociable()
    for nom_, neuf, jeu in (("S2", S2, j["ue"]["s2"]), ("COR", C, j["cor"]), ("etat", etat, j["etat_negociable"])):
        if json.loads(json.dumps(neuf, ensure_ascii=False)) != jeu:
            fail("hors ligne : la piece %s ne redonne pas le jeu publie (piece changee depuis le releve)" % nom_)
    fin = max(a["annee"] for a in annees)
    if "sous_secteurs" not in j["france"]:
        fail("hors ligne : le jeu publie n'a pas encore de bloc sous_secteurs (premier releve en ligne)")
    return (fin, total, dec, annees, cp, pat, det, det_pays, j["controles"]["detention_bce_eurostat_annees"], S2, C,
            j["conservation"], etat, j["france"]["sous_secteurs"])


# ------------------------------------------------------------------ calcul
def calcul():
    if "--hors-ligne" in sys.argv[1:]:
        return calcul_hors_ligne()
    R, ecartes, calculables = emplois()
    if "FR" not in R or any(e["pays"] == "FR" for e in ecartes):
        fail("France : annee ecartee par le temoin ou par l'identite (%s)" % [e for e in ecartes if e["pays"] == "FR"])
    FR = R["FR"]
    fin = max(FR)
    a0 = fin - 29
    if not all(a in FR for a in range(a0, fin + 1)):
        fail("France : moins de trente annees (%d-%d)" % (min(FR), fin))
    if not all(FR[a]["d92_publie"] for a in range(a0, fin + 1)):
        fail("France : aides a l'investissement versees (D92PAY) absentes certaines annees")
    total = cumul(FR, a0, fin)
    dec = [cumul(FR, a0 + 10 * k, a0 + 10 * k + 9) for k in range(3)]
    annees = [dict(annee=a, **{k: FR[a][k] for k in ("besoin", "desepargne", "investissement_net", "actifs", "transferts", "aides_investissement")})
              for a in range(a0, fin + 1)]
    # Contre-population : même fenêtre de vingt-cinq ans pour tous les pays, complète
    c0 = fin - 24
    cp = []
    for p in sorted(PAYS):
        s = cumul(R.get(p, {}), c0, fin)
        if s:
            cp.append(dict(pays=p, nom=PAYS[p][0], **s))
    pat = patrimoine(fin)
    det, det_pays, det_controle = detention()
    S2 = s2()
    C = cor()
    cons = dict(annees_pays_calculables=calculables, retenues=sum(len(v) for v in R.values()), ecartees=len(ecartes), detail_ecartees=ecartes)
    return fin, total, dec, annees, cp, pat, det, det_pays, det_controle, S2, C, cons, etat_negociable(), sous_secteurs(annees)


# ------------------------------------------------------------------ affichage et gardes
def affichage(fin, total, dec, annees, cp, pat, det, det_pays, S2, C, cons, etat, ss):
    A = {"fin": str(fin), "a0": str(total["debut"]), "tolerance": fr(TOL, 2),
         "dsm_lib": lib("dsm"), "cor_lib": lib("cor"),
         "proj_an": max(SOURCES[k]["publie"] for k in ("dsm", "cor"))[:4]}   # année de publication des projections
    t = total
    A.update(b=fr(t["besoin"]), k=fr(t["actifs"]), e=fr(t["desepargne"]), t=fr(t["transferts"]), d92=fr(t["aides_investissement"]),
             part_k=fr(t["part_actifs"], 0), part_e=fr(t["part_desepargne"], 0), part_t=fr(t["part_transferts"], 0),
             part_k_eur=fr(t["part_actifs_euros"], 0), part_n=fr(t["part_investissement_net"], 0),
             part_max=fr(100 * (t["actifs"] + t["aides_investissement"]) / t["besoin"], 0))
    for k, d in enumerate(dec, 1):
        p = "d%d_" % k
        A[p + "lib"] = "%d-%d" % (d["debut"], d["fin"])
        A[p + "b"] = fr(d["besoin"]); A[p + "k"] = fr(d["actifs"]); A[p + "e"] = fr(d["desepargne"]); A[p + "t"] = fr(d["transferts"])
        A[p + "part"] = fr(d["part_actifs"], 0); A[p + "k_moy"] = fr(d["actifs"] / d["annees"])
        A[p + "k_moy_pt"] = lg.points(d["actifs"] / d["annees"], 1, LANG)
    pos = [r["annee"] for r in annees if r["desepargne"] < 0]
    A["ep_pos_n"] = lettres(len(pos)); A["ep_pos_lib"] = enumere([str(a) for a in pos]) if pos else ""
    # patrimoine
    p0, p1 = pat[0], pat[-1]
    A.update(s0=str(p0["annee"]), s1=str(p1["annee"]), pn0=fr(p0["patrimoine_net"]), pn1=fr(p1["patrimoine_net"]),
             pn_var=fr(abs(p1["patrimoine_net"] - p0["patrimoine_net"])),
             prod_var=signe(p1["actifs_produits"] - p0["actifs_produits"]), terr_var=signe(p1["terrains"] - p0["terrains"]),
             bf_var_abs=fr(abs(p1["valeur_financiere_nette"] - p0["valeur_financiere_nette"])))
    # Phase B (11/10/2026) : le bilan devient un complément à deux bornes. Plus haut, encore au-dessus du point de départ,
    # fenêtres d'au moins dix ans, et la dernière décennie (mêmes bornes que d3_lib).
    pmax = lg.extremum(pat, lambda x: x["patrimoine_net"], 1)
    psup = max((x for x in pat if x["annee"] > p0["annee"] and lg.depasse(x["patrimoine_net"], p0["patrimoine_net"])), key=lambda x: x["annee"])  # brut-ok: derniere annee d'un ensemble, cle = annee
    pd3 = next(x for x in pat if x["annee"] == dec[-1]["debut"])
    fen = [(x, y) for x in pat for y in pat if y["annee"] - x["annee"] >= 10]
    A.update(pn_haut=fr(pmax["patrimoine_net"]), pn_haut_annee=str(pmax["annee"]), pn_sup=fr(psup["patrimoine_net"]),
             pn_sup_annee=str(psup["annee"]), pn_d3_0=fr(pd3["patrimoine_net"]), pn_d3_an0=str(pd3["annee"]),
             pn_d3_var=lg.sans_signe(p1["patrimoine_net"] - pd3["patrimoine_net"], 1, LANG),
             pn_d3_var_pt=lg.points(abs(p1["patrimoine_net"] - pd3["patrimoine_net"]), 1, LANG),
             pn_fen_n=str(len(fen)), pn_fen_hausse=str(sum(1 for x, y in fen if lg.depasse(y["patrimoine_net"], x["patrimoine_net"]))),
             prod_var_pt=lg.points(p1["actifs_produits"] - p0["actifs_produits"], 1, LANG, signe_=True),
             terr_var_pt=lg.points(p1["terrains"] - p0["terrains"], 1, LANG, signe_=True))
    # contre-population
    A["cp_a0"] = str(cp[0]["debut"]); A["cp_n"] = str(len(cp))
    frc = next(x for x in cp if x["pays"] == "FR")
    # Arbitrage PRO-20261003-195656 : le ratio n'a de sens que si le besoin cumulé est substantiel.
    groupe = [x for x in cp if x["besoin"] >= SEUIL_BESOIN]
    sup = [x for x in groupe if x["part_actifs"] >= 50]
    inf = [x for x in groupe if lg.depasse(frc["part_actifs"], x["part_actifs"], 0)]       # parts affichées à l'unité
    exc = [x for x in cp if x["besoin"] < SEUIL_BESOIN]
    A["cp_fr_part"] = fr(frc["part_actifs"], 0)
    A["cp_sup_n"] = lettres(len(sup)); A["cp_sup_pays"] = enumere([art(x["pays"]) for x in sorted(sup, key=lambda x: -x["part_actifs"])])
    A["cp_inf_n"] = lettres(len(inf)); A["cp_inf_pays"] = enumere([art(x["pays"]) for x in sorted(inf, key=lambda x: -x["part_actifs"])])
    A["cp_exc_pays"] = enumere([art(x["pays"]) for x in exc]) if exc else ""
    A["cp_seuil"] = fr(SEUIL_BESOIN, 0); A["cp_groupe_n"] = str(len(groupe))
    A["cp_exc_pays_maj"] = maj(A["cp_exc_pays"])
    # « ni universel ni exceptionnel » : deux pays nommés de chaque côté, avec leur part (gardes : côté, besoin substantiel)
    part_de = {x["pays"]: x for x in cp}
    nomme = lambda codes: lg.enumere([art(c) + " (" + fr(part_de[c]["part_actifs"], 0) + pc() + ")" for c in codes], LANG)
    A["cp_cite_sup"] = nomme(CITE_SUP); A["cp_cite_inf"] = nomme(CITE_INF)
    A["part_hors_k"] = fr(100 - float(A["part_k"].replace(",", ".")), 0)   # « les 82 % restants » : complément de la part AFFICHÉE
    # détention
    d0, d1 = det[0], det[-1]
    nr_max = lg.extremum(det, lambda x: x["part_non_residents"], 1)
    bc_max = lg.extremum(det, lambda x: x["part_banque_centrale"], 1)
    maj1 = next(x["annee"] for x in det if x["part_non_residents"] > 50)
    apres = [x for x in det if x["annee"] >= maj1]
    A.update(h0=str(d0["annee"]), h_fin=str(d1["annee"]), nr0=fr(d0["part_non_residents"]), nr_fin=fr(d1["part_non_residents"]),
             res_fin=fr(100 - d1["part_non_residents"]),
             nr_max=fr(nr_max["part_non_residents"]), nr_max_annee=str(nr_max["annee"]), nr_maj_premiere=str(maj1),
             nr_maj_n=str(sum(1 for x in apres if x["part_non_residents"] > 50)), nr_maj_tot=str(len(apres)),
             bc0=fr(d0["part_banque_centrale"]), bc_max=fr(bc_max["part_banque_centrale"]), bc_max_annee=str(bc_max["annee"]),
             bc_fin=fr(d1["part_banque_centrale"]), bc_var=fr(bc_max["part_banque_centrale"] - d1["part_banque_centrale"]),
             nr_var_bc=fr(d1["part_non_residents"] - bc_max["part_non_residents"]), nr_bcmax=fr(bc_max["part_non_residents"]),
             dette_h_fin=fr(d1["dette"]))
    # La part « nationale » tenue par la banque centrale : première année au-dessus du dixième, autres résidents avant/après
    bc10 = [x for x in det if x["part_banque_centrale"] > 10]
    avant = max((x for x in det if x["annee"] < bc_max["annee"] and x["part_banque_centrale"] < 3), key=lambda x: x["annee"])  # brut-ok: derniere annee d'un ensemble, cle = annee
    A.update(bc10_premiere=str(bc10[0]["annee"]), an_avant=str(avant["annee"]), ar_avant=fr(avant["part_autres_residents"]),
             ar_fin=fr(d1["part_autres_residents"]))
    # Témoin de périmètre État (Banque de France) : mêmes dates de fin d'année que la série BCE
    ef = {x["annee"]: x["part"] for x in etat["fins"]}
    emax = lg.extremum(etat["fins"], lambda x: x["part"], 1)
    emin_apres = lg.extremum((x for x in etat["fins"] if x["annee"] > emax["annee"]), lambda x: x["part"], 1, maxi=False)
    A.update(etat_a0=str(etat["fins"][0]["annee"]), etat_0=fr(etat["fins"][0]["part"]), etat_fin=fr(ef[d1["annee"]]),
             etat_max=fr(emax["part"]), etat_max_annee=str(emax["annee"]), etat_min=fr(emin_apres["part"]),
             etat_min_annee=str(emin_apres["annee"]), etat_maj=date_longue(etat["maj"]),
             etat_lib=lib("det"))
    rp = etat["releve"]["periode"]
    A.update(etat_der=fr(etat["releve"]["valeur"]), etat_der_trim=(TRIM if LANG == "fr" else TRIM_EN)[int(rp[-1]) - 1] + " " + rp[:4],
             etat_der_lu=date_longue(etat["releve"]["lu_le"]))
    for p, x in det_pays.items():
        A["h_" + p.lower()] = fr(x["part_non_residents"], 0)
    # S2
    x = S2["FR"]
    pts = lambda v, s=False: lg.points(v, 1, LANG, signe_=s)
    A.update(s2=fr(x["S2"]["base"]), ibp=fr(x["IBP"]["base"]), coa=fr(x["CoA"]["base"]),
             s2_pt=pts(x["S2"]["base"]), ibp_pt=pts(x["IBP"]["base"]), coa_pt=pts(x["CoA"]["base"]),
             # composantes de l'effet des variations futures (phase B : par composante, pensions négatives)
             pen_pt=pts(x["pensions"]["base"], True), sante_pt=pts(x["sante"]["base"], True),
             dep_pt=pts(x["dependance"]["base"], True), edu_pt=pts(x["education"]["base"], True),
             ibp_prod=fr(x["IBP"]["productivite"]), coa_prod=fr(x["CoA"]["productivite"]), s2_prod=fr(x["S2"]["productivite"]),
             ibp_risque=fr(x["IBP"]["risque"]), coa_risque=fr(x["CoA"]["risque"]), s2_risque=fr(x["S2"]["risque"]),
             s2_risque_pt=pts(x["S2"]["risque"]), coa_risque_pt=pts(x["CoA"]["risque"]),
             ibp_prec=fr(x["IBP"]["base_precedente"]), coa_prec=fr(x["CoA"]["base_precedente"]), s2_prec=fr(x["S2"]["base_precedente"]),
             dsm_prec_an=str(int(lib("dsm")[-4:]) - 1))
    A["ibp_pdf"] = fr(x["S2"]["base"] - x["CoA"]["base"])   # convention du tableau 3.2 du rapport : S2 - vieillissement
    D = dsm_definitions()
    A.update(s2_an=str(D["s2_an"]), s1_cible=str(D["s1_cible"]), s1_horizon=str(D["s1_horizon"]))
    # COR : sensibilité aux hypothèses, présentée à part ; la substitution chiffrée COR -> S2 est retirée (11/10/2026).
    cor_var = C["depenses1"] - C["depenses0"]
    A.update(cor_a0=str(C["a0"]), cor_a1=str(C["a1"]), cor_dep0=fr(C["depenses0"]), cor_dep1=fr(C["depenses1"]), cor_var=signe(cor_var),
             cor_var_pt=pts(cor_var, True),
             cor_res0=fr(C["ressources0"]), cor_res1=fr(C["ressources1"]), cor_solde1=fr(C["solde1"]),
             cor_deficit1_pt=lg.points(abs(C["solde1"]), 1, LANG))
    # Volet 5 : même valeur, même année que la page « La dette publique peut-elle baisser ? » (gardes)
    V5 = volet5()
    A.update(v5_annee=str(V5["annee"]), v5_ecart=lg.sans_signe(V5["ecart"], 1, LANG), v5_ecart_pt=lg.points(abs(V5["ecart"]), 1, LANG))
    vieux = sorted(g for g, y in S2.items() if y["CoA"]["base"] > y["IBP"]["base"])
    A["pays_coa_n"] = str(len(vieux)); A["pays_n"] = str(len(S2))  # chiffres des deux côtés : « 15 pays sur 27 », jamais lettres et chiffres mêlés
    # Où le recul de l'investissement net s'est produit : deuxième et troisième décennies, administrations locales et autres
    d2, d3 = dec[1], dec[2]
    moy = lambda k, d: sum(x[k] for x in ss["annees"] if d["debut"] <= x["annee"] <= d["fin"]) / d["annees"]
    inv = dict(d2=moy("net_apu", d2), d3=moy("net_apu", d3), loc2=moy("net_loc", d2), loc3=moy("net_loc", d3))
    inv.update(var=inv["d2"] - inv["d3"], loc_var=inv["loc2"] - inv["loc3"])
    inv.update(aut_var=inv["var"] - inv["loc_var"], loc_part=100 * inv["loc_var"] / inv["var"])
    creux = min((x for x in ss["annees"] if d3["debut"] <= x["annee"] <= d3["fin"]), key=lambda x: round(x["net_loc"], 2))  # brut-ok: point bas, egalite tranchee par l'ordre des annees
    der = ss["annees"][-1]
    coll = json.loads(COLL.read_text(encoding="utf-8-sig"))["affichage"]
    inv.update(creux_an=creux["annee"], creux=creux["net_loc"], loc_fin=der["net_loc"], ep_a0=int(coll["ep_a0"]), ep_a1=int(coll["ep_a1"]))
    p2 = lambda v: lg.points(v, 2, LANG)
    A.update(inv_d2_pt=p2(inv["d2"]), inv_d3_pt=p2(inv["d3"]), inv_var_pt=p2(inv["var"]), inv_loc_var_pt=p2(inv["loc_var"]),
             inv_aut_var_pt=p2(inv["aut_var"]), inv_creux_an=str(creux["annee"]), inv_creux_pt=p2(creux["net_loc"]),
             inv_loc_fin_pt=p2(der["net_loc"]))
    A["inv_ccf_note"] = ((" La consommation de capital fixe publiée pour %d est identique à celle de %d : le dernier point reste provisoire."
                          if LANG == "fr" else " The consumption of fixed capital published for %d is identical to that of %d: the last point remains provisional.")
                         % (der["annee"], der["annee"] - 1)) if ss["ccf_reprise"] else ""
    # contrôles et conservation
    A["cons_calculables"] = str(cons["annees_pays_calculables"]); A["cons_ecartees"] = str(cons["ecartees"])
    return A, dict(sup=sup, inf=inf, exc=exc, frc=frc, vieux=vieux, cor_var=cor_var, nr_max=nr_max, bc_max=bc_max, maj1=maj1,
                   etat=etat, emax=emax, emin_apres=emin_apres, pmax=pmax, psup=psup, pd3=pd3, fen=fen, V5=V5, D=D, groupe=groupe,
                   inv=inv)


def gardes(fin, total, dec, annees, cp, pat, det, det_pays, S2, C, aux):
    d1, d2, d3 = dec
    p0, p1 = pat[0], pat[-1]
    x = S2["FR"]
    dl = det[-1]
    cor_var = aux["cor_var"]
    G = [
        # --- A1 : à quoi ont correspondu les déficits
        ("chapo : « seulement » — moins d'un quart du besoin de financement en acquisitions nettes d'actifs", total["part_actifs"] < 25),
        ("verdict du test conservé : investissement net seul sous 25 % du besoin", total["part_investissement_net"] < 25),
        ("la mesure en euros donne le même ordre (à 3 points près)", abs(total["part_actifs_euros"] - total["part_actifs"]) < 3),
        ("« près des deux tiers » en dépenses courantes non couvertes (60 à 70 %)", 60 <= total["part_desepargne"] < 70),
        ("la désépargne est le premier emploi du besoin de financement", total["desepargne"] > max(total["actifs"], total["transferts"])),
        ("« la part investie a baissé de décennie en décennie »", lg.depasse(d1["part_actifs"], d2["part_actifs"], 0) and lg.depasse(d2["part_actifs"], d3["part_actifs"], 0)),
        # Corrections de clôture (11/10/2026, contre-expertise du rendu, § 2) : la baisse 32 -> 19 -> 10 % en deux mouvements,
        # décomposition arithmétique de la part (actifs / besoin), aucune cause. Seuils écrits avant l'écriture de la phrase :
        # « presque stables » / « presque identique » = moins de 5 % d'écart relatif ; « augmente » = plus de 50 %, si bien que
        # le besoin explique l'essentiel du recul (ln 1,5 = 0,41 contre ln 1,05 = 0,05) ; « presque de moitié » = recul de
        # 40 à 55 % (au-delà, « plus de la moitié » ; en deçà, pas « presque de moitié »).
        ("deux temps (1) : de la première à la deuxième décennie, actifs « presque stables » (< 5 %) et besoin en hausse (> 50 %)",
         abs(d2["actifs"] / d1["actifs"] - 1) < 0.05 and d2["besoin"] / d1["besoin"] - 1 > 0.50),
        ("deux temps (2) : de la deuxième à la troisième décennie, besoin « presque identique » (< 5 %) et actifs « presque de moitié » en moins (-40 à -55 %)",
         abs(d3["besoin"] / d2["besoin"] - 1) < 0.05 and -0.55 < d3["actifs"] / d2["actifs"] - 1 < -0.40),
        ("dernière décennie : « que » — moins d'un demi-point de PIB par an d'acquisitions nettes d'actifs", d3["actifs"] / d3["annees"] < 0.5),
        ("recettes courantes supérieures aux dépenses courantes deux à trois années sur trente (phrase au pluriel, années énumérées)", 2 <= sum(1 for r in annees if r["desepargne"] < 0) <= 3),
        ("« aucune année l'investissement n'a égalé le besoin de financement »", all(lg.depasse(r["besoin"], r["actifs"], 1) for r in annees)),
        ("en comptant toutes les aides à l'investissement versées : « toujours moins de la moitié »",
         (total["actifs"] + total["aides_investissement"]) / total["besoin"] < 0.5),
        # --- complément : le bilan (phase B : deux bornes, réévaluations dites ; plus un témoin de concordance)
        ("bilan : « recule » sur la période entière", lg.depasse(p0["patrimoine_net"], p1["patrimoine_net"])),
        ("bilan : « plus haut » en cours de période que les deux bornes (le maximum n'est ni la première ni la dernière année)",
         aux["pmax"]["patrimoine_net"] > max(p0["patrimoine_net"], p1["patrimoine_net"]) and aux["pmax"]["annee"] not in (p0["annee"], p1["annee"])),
        ("bilan : « encore au-dessus de son niveau de départ » une année tardive (au moins vingt ans après le début)",
         aux["psup"]["annee"] - p0["annee"] >= 20),
        ("bilan : « a augmenté » sur la dernière décennie, alors que la part en actifs y est la plus faible",
         lg.depasse(p1["patrimoine_net"], aux["pd3"]["patrimoine_net"]) and lg.depasse(min(d1["part_actifs"], d2["part_actifs"]), d3["part_actifs"], 0)),
        ("bilan : « des fenêtres en hausse, d'autres en baisse » (ni toutes ni aucune)",
         0 < sum(1 for a, b in aux["fen"] if lg.depasse(b["patrimoine_net"], a["patrimoine_net"])) < len(aux["fen"])),
        ("actifs produits : « à peine bougé » (moins de 5 points)", abs(p1["actifs_produits"] - p0["actifs_produits"]) < 5),
        ("terrains : réévaluation en hausse, plus forte que la variation des actifs produits",
         p1["terrains"] > p0["terrains"] and abs(p1["actifs_produits"] - p0["actifs_produits"]) < p1["terrains"] - p0["terrains"]),
        ("patrimoine : la série couvre au moins vingt-cinq ans", p1["annee"] - p0["annee"] >= 25),
        # --- contre-population : « ni universel ni exceptionnel »
        ("« ni exceptionnel » : au moins deux pays plus bas que la France", len(aux["inf"]) >= 2),
        ("« ni universel » : au moins deux pays au-dessus de la moitié", len(aux["sup"]) >= 2),
        ("pays nommés au-dessus de la moitié (%s), besoin cumulé au-delà du seuil" % "/".join(CITE_SUP),
         set(CITE_SUP) <= {y["pays"] for y in aux["sup"]}),
        ("pays nommés sous la France (%s), besoin cumulé au-delà du seuil" % "/".join(CITE_INF),
         set(CITE_INF) <= {y["pays"] for y in aux["inf"]}),
        ("« sur 100 » : la part affichée et son complément font 100", int(round(total["part_actifs"])) + int(round(100 - total["part_actifs"])) == 100),
        ("17 % et 18 % : la France sur la fenêtre européenne arrondit à une valeur proche de la période longue (à 3 points)",
         abs(aux["frc"]["part_actifs"] - total["part_actifs"]) < 3),
        ("contre-population : au moins vingt pays de l'Union comparés sur la même fenêtre", len(cp) >= 20),
        # --- A2 : détention
        ("« un peu plus de la moitié » détenue par des non-résidents la dernière année (50 à 60 %)", 50 < dl["part_non_residents"] < 60),
        ("la part non résidente a monté depuis la première année d'au moins 10 points", dl["part_non_residents"] - det[0]["part_non_residents"] >= 10),
        ("la part non résidente est majoritaire la plupart des années depuis qu'elle a dépassé la moitié",
         sum(1 for y in det if y["annee"] >= aux["maj1"] and y["part_non_residents"] > 50) * 2 > sum(1 for y in det if y["annee"] >= aux["maj1"])),
        ("Banque de France : « détenteur majeur » — plus de 15 % de la dette à son maximum", aux["bc_max"]["part_banque_centrale"] > 15),
        ("Banque de France : moins de 3 % de la dette la première année", det[0]["part_banque_centrale"] < 3),
        ("Banque de France : « ses avoirs refluent » depuis le maximum", dl["part_banque_centrale"] < aux["bc_max"]["part_banque_centrale"] - 2),
        ("« la part des non-résidents est remontée » depuis le maximum de la Banque de France",
         dl["part_non_residents"] > aux["bc_max"]["part_non_residents"] + 2),
        ("la hausse de la Banque de France commence avec les achats de titres publics (2015-2016) : moins de 3 % en 2014, plus de 5 % en 2015",
         next(y for y in det if y["annee"] == 2014)["part_banque_centrale"] < 3 and next(y for y in det if y["annee"] == 2015)["part_banque_centrale"] > 5),
        ("Italie et Suède : prémisse « mieux vérifiée » (au moins 10 points de moins que la France)",
         all(det_pays[p]["part_non_residents"] < dl["part_non_residents"] - 10 for p in ("IT", "SE"))),
        ("Autriche et Belgique : « moins » vérifiée (plus que la France)", all(lg.depasse(det_pays[p]["part_non_residents"], dl["part_non_residents"]) for p in ("AT", "BE"))),
        ("contre-population de détention : même dernière année que la France", all(det_pays[p]["annee"] == dl["annee"] for p in ("IT", "SE", "AT", "BE"))),
        # --- témoin de périmètre État (Banque de France, Webstat) : il doit dire la même chose que la série BCE
        ("État : la série Banque de France couvre la dernière année de la série BCE", dl["annee"] in {x["annee"] for x in aux["etat"]["fins"]}),
        ("État : « majorité non résidente aujourd'hui » (plus de 50 % la dernière année)",
         next(x["part"] for x in aux["etat"]["fins"] if x["annee"] == dl["annee"]) > 50),
        ("État : la part a monté depuis la première fin d'année publiée", aux["etat"]["fins"][-1]["part"] > aux["etat"]["fins"][0]["part"] + 10),
        ("État : « creux au moment des achats de la banque centrale » (à un an près du maximum de la Banque de France)",
         abs(aux["emin_apres"]["annee"] - aux["bc_max"]["annee"]) <= 1),
        ("État : « remonte depuis » (au moins 3 points au-dessus du creux)", aux["etat"]["fins"][-1]["part"] > aux["emin_apres"]["part"] + 3),
        ("État : la dernière observation (relevé AFT) prolonge la remontée", aux["etat"]["releve"]["valeur"] >= aux["etat"]["fins"][-1]["part"]),
        ("les deux séries vont dans le même sens depuis le creux de l'État",
         dl["part_non_residents"] > next(y for y in det if y["annee"] == aux["emin_apres"]["annee"])["part_non_residents"]),
        ("« depuis » la première année au-dessus du dixième, la Banque de France y reste chaque année",
         all(y["part_banque_centrale"] > 10 for y in det if y["annee"] >= next(z["annee"] for z in det if z["part_banque_centrale"] > 10))),
        ("« n'a tenu que par elle » : les autres résidents perdent au moins 8 points de part depuis la veille des achats",
         next(y for y in det if y["annee"] == max(z["annee"] for z in det if z["annee"] < aux["bc_max"]["annee"] and z["part_banque_centrale"] < 3))["part_autres_residents"]
         - dl["part_autres_residents"] >= 8),
        ("la part résidente totale ne baisse pas plus que la part des autres résidents (la banque centrale compense)",
         dl["part_banque_centrale"] > det[0]["part_banque_centrale"] + 8),
        # --- A3 : S2, position de départ et effet des variations futures des dépenses liées à l'âge
        ("« l'essentiel » : position de départ supérieure au double de l'effet des variations futures (scénario de base)", x["IBP"]["base"] > 2 * x["CoA"]["base"]),
        ("la position de départ reste le premier terme dans les deux scénarios de risque et dans l'édition précédente",
         all(x["IBP"][j] > x["CoA"][j] for j in JEUX)),
        ("« nettement plus dans le scénario de risque » : l'effet des variations futures y gagne plus d'un point",
         x["CoA"]["risque"] > x["CoA"]["base"] + 1),
        ("S2 = position de départ + effet des variations futures, dans chaque jeu (total affiché par scénario)",
         all(abs(x["S2"][j] - x["IBP"][j] - x["CoA"][j]) <= 0.02 for j in JEUX)),
        ("composante pensions négative (« dépense de pensions en baisse dans le PIB »)", x["pensions"]["base"] < 0),
        ("santé et dépendance positives, chacune", x["sante"]["base"] > 0 and x["dependance"]["base"] > 0),
        ("éducation négative", x["education"]["base"] < 0),
        ("« un solde » : des composantes de signes contraires", x["pensions"]["base"] < 0 < x["sante"]["base"]),
        ("l'indicateur sait montrer l'inverse : au moins trois pays où l'effet des variations futures pèse davantage", len(aux["vieux"]) >= 3),
        ("« dont la Belgique, l'Espagne et le Luxembourg »", {"BE", "ES", "LU"} <= set(aux["vieux"])),
        ("S2 : ajustement la même année que S1, et S1 vise une cible et une date (la phrase « ni niveau ni date » les cite)",
         aux["D"]["s2_an"] == aux["D"]["s1_an"] and aux["D"]["s1_horizon"] > aux["D"]["s1_an"]),
        ("COR : dépense de retraite en hausse en fin de projection (« là où la Commission retient une baisse »)", cor_var > 0 > x["pensions"]["base"]),
        ("COR : le solde de fin de projection est un déficit", C["solde1"] < 0),
        ("méthode : le tableau 3.2 du rapport imprime une position initiale différente de la composante (sinon la phrase « au lieu de » tombe)",
         fr(x["S2"]["base"] - x["CoA"]["base"]) != fr(x["IBP"]["base"])),
        ("COR : baisse des ressources « autant » que hausse des dépenses (rapport 0,75 à 1,33)",
         0.75 < (C["ressources0"] - C["ressources1"]) / cor_var < 1.33),
        # --- volet 5 : la page cite l'écart au solde stabilisant, même année, même valeur que le volet 5 (tableau à quatre lignes)
        ("volet 5 : même dernière année que cette page", aux["V5"]["annee"] == fin),
        ("volet 5 : solde observé SOUS le solde stabilisant (« écart » à combler)", aux["V5"]["ecart"] < 0),
        ("volet 5 : valeur citée identique à celle que le volet 5 affiche (FR et EN)",
         lg.sans_signe(aux["V5"]["ecart"], 1, "fr") == aux["V5"]["aff"]["ecart_dernier"]
         and lg.sans_signe(aux["V5"]["ecart"], 1, "en") == aux["V5"]["aff_en"]["ecart_dernier"]
         and aux["V5"]["aff"]["annee_fin"] == str(aux["V5"]["annee"])),
        ("« deux nombres voisins » : S2 et l'écart du volet 5 à moins d'un point", abs(x["S2"]["base"] - abs(aux["V5"]["ecart"])) < 1),
        ("détention : « depuis le début des années 2000 » (première année majoritaire de 2000 à 2004)", 2000 <= aux["maj1"] <= 2004),
        # Paragraphe « Deux projections, deux jeux d'hypothèses » (remarque du secrétariat du COR, 06/10/2026) : il décrit
        # les hypothèses de ces deux éditions (Insee nouvelle, Agirc-Arrco 2038 ; rapport de 2024 sur le vieillissement).
        # « Où ce recul s'est produit » (18/10/2026). Seuils écrits avant la phrase : « un peu plus de la moitié » = 50 à 65 %
        # de la baisse ; « guère plus que l'usure » = investissement net local sous 0,15 point de PIB ; « remonte » = au
        # moins le triple du point bas ; le point bas tombe dans la fenêtre de baisse des transferts de la page Collectivités.
        ("« un peu plus de la moitié » du recul de l'investissement net vient des administrations locales (50 à 65 %)",
         aux["inv"]["var"] > 0 and 50 < aux["inv"]["loc_part"] <= 65),
        ("« le reste » : l'investissement net des autres administrations a reculé lui aussi", aux["inv"]["aut_var"] > 0),
        ("« guère plus que l'usure » : investissement net local au point bas sous 0,15 point de PIB", 0 <= aux["inv"]["creux"] < 0.15),
        ("« puis remonte » : investissement net local de la dernière année au moins triple du point bas",
         aux["inv"]["loc_fin"] >= 3 * max(aux["inv"]["creux"], 0.01)),
        ("point bas « pendant la baisse des dotations » : dans la fenêtre de la page Collectivités",
         aux["inv"]["ep_a0"] <= aux["inv"]["creux_an"] <= aux["inv"]["ep_a1"]),
        ("point bas qui « suit les élections municipales de 2014 » : dans les deux années suivantes", 2014 < aux["inv"]["creux_an"] <= 2016),
        ("hypothèses du COR et de la Commission décrites pour ces éditions : une autre édition impose de relire le paragraphe",
         SOURCES["cor"]["lib"] == "rapport annuel de juin 2026" and SOURCES["dsm"]["lib"] == "Debt Sustainability Monitor 2025"),
    ]
    if "--mutation=recul_local" in sys.argv:   # témoin positif : la part locale du recul passe sous la moitié
        i = next(k for k, (nom_, _) in enumerate(G) if "du recul de l'investissement net vient des administrations locales" in nom_)
        G[i] = (G[i][0], 50 < aux["inv"]["loc_part"] - 15 <= 65)
    if "--mutation-garde" in sys.argv:   # témoin positif : part investie de la première décennie mise sous la dernière
        G[5] = (G[5][0], d3["part_actifs"] - 1 > d2["part_actifs"] > d3["part_actifs"])
    faux = [n for n, ok in G if not ok]
    if faux:
        fail("la page affirme ce que les donnees ne soutiennent plus : " + " ; ".join(faux))
    return len(G)


# ------------------------------------------------------------------ gardes de TEXTE (phase B, 11/10/2026)
# Une garde numérique protège contre une variation des données ; ces gardes-ci lisent la PAGE (FR et EN) et les textes des
# figures, et protègent trois décisions de la contre-expertise de phase A qu'une retouche pourrait défaire en silence.
TOK_BILAN = ("pn0", "pn1", "pn_var", "pn_haut", "pn_haut_annee", "pn_sup", "pn_sup_annee", "pn_d3_0", "pn_d3_var",
             "pn_fen_n", "pn_fen_hausse", "prod_var", "prod_var_pt", "terr_var", "terr_var_pt", "bf_var_abs")
TOK_AGE = ("coa", "coa_pt", "coa_prod", "coa_risque", "coa_risque_pt", "coa_prec")
AGE_OK = {"fr": r"évolution (?:future|projetée)|variations? futures?", "en": r"future change|projected change|future variations?"}
AGE_MOT = {"fr": r"coût du vieillissement|vieillissement\s*:", "en": r"cost of ageing|ageing\s*:"}
MOTS_BILAN = {"fr": r"patrimoine|bilan", "en": r"net worth|balance sheet"}
PROSCRIT = {"fr": r"\b(?:financ(?:é|ée|és|ées)|transmise?s?|(?:a|ont|avait|avaient) couvert|(?:a|ont|avait|avaient) monté"
                  r"|ne vaut (?:plus )?que pour|rend le verdict|n'est pas un cas extrême|proches? de l'équilibre)\b",
            "en": r"\b(?:financed|passed on|holds for only|now holds for only|gives the verdict|not an extreme case"
                  r"|close to balance)\b"}


def _tok(k: str) -> str:
    return r'gen-val "%s"|\{gen\.%s\}' % (k, k)


def _phrases(s: str) -> list[str]:
    return [p for p in re.split(r"(?<=[.!?])\s+|\n", s) if p.strip()]


def gardes_texte(textes: dict[str, str], txt_figures: dict[str, str]) -> int:
    """textes : {langue: source markdown de la page}. Rend le nombre de gardes ; arrêt à la première qui refuse."""
    import yaml
    n = 0

    def g(ok, msg):
        nonlocal n
        n += 1
        if not ok:
            fail("garde de texte : " + msg)

    for lang, s in textes.items():
        _, fm, corps = s.split("---", 2)
        meta = yaml.safe_load(fm)
        resumes = [meta.get("chapo", ""), meta.get("description", ""), meta.get("og_image_alt", "")] + \
                  [q["answer"] for q in meta.get("faq", [])]
        # 1. Le bilan est un complément replié : jamais dans le chapô, la description, la FAQ, ni hors d'un repli
        hors_repli = re.sub(r"<details\b.*?</details>", "", corps, flags=re.S)
        for k in TOK_BILAN:
            g(not any(re.search(_tok(k), r) for r in resumes), "%s : le bilan (%s) dans le chapo, la description ou la FAQ" % (lang, k))
            g(not re.search(_tok(k), hors_repli), "%s : le bilan (%s) hors du complement replie" % (lang, k))
        m = re.search(r"\{#(?:retenir|takeaways)\}(.*?)</div>", corps, re.S)
        g(m is not None, "%s : section de conclusion introuvable" % lang)
        g(not re.search(MOTS_BILAN[lang], m.group(1), re.I), "%s : la conclusion invoque le bilan comme une preuve" % lang)
        # 2. Le second terme de S2 se dit « évolution future des dépenses liées à l'âge », jamais « vieillissement » seul
        for p in _phrases(corps) + [x for r in resumes for x in _phrases(r)]:
            if any(re.search(_tok(k), p) for k in TOK_AGE) or re.search(AGE_MOT[lang], p, re.I):
                g(re.search(AGE_OK[lang], p, re.I), "%s : terme de vieillissement sans son evolution future : %s"
                  % (lang, p[:90].encode("ascii", "replace").decode()))
        # 3. Vocabulaire proscrit (verrou PRO-20261003-195656 ; contre-expertise de phase A ; leçon 16)
        for p in _phrases(corps) + [x for r in resumes + [meta.get("title", "")] for x in _phrases(r)]:   # pas les commentaires YAML
            g(not re.search(PROSCRIT[lang], p, re.I), "%s : vocabulaire proscrit : %s" % (lang, p[:90].encode("ascii", "replace").decode()))
    for lang, t in txt_figures.items():
        for p in _phrases(t):
            g(not re.search(PROSCRIT[lang], p, re.I), "figures %s : vocabulaire proscrit : %s" % (lang, p[:90].encode("ascii", "replace").decode()))
            if re.search(AGE_MOT[lang], p, re.I):
                g(re.search(AGE_OK[lang], p, re.I), "figures %s : terme de vieillissement sans son evolution future" % lang)
    return n


def mutation_texte(textes: dict[str, str]) -> dict[str, str]:
    """Témoins positifs des gardes de texte : chaque mutation doit faire échouer le générateur."""
    m = next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith("--mutation=")), None)
    t = dict(textes)
    if m == "bilan":            # une conclusion qui rend le bilan probant
        t["fr"] = t["fr"].replace("{#retenir}", '{#retenir}\n\nLe bilan le confirme : le patrimoine net est tombé à {{< gen-val "pn1" >}}&nbsp;% du PIB.', 1)
    elif m == "vieillissement":  # « vieillissement » sans « évolution future »
        t["fr"] = re.sub(AGE_OK["fr"], "vieillissement", t["fr"])
    elif m == "bilan_mot":      # la conclusion invoque le bilan sans jeton (garde de mots, distincte de la garde de jetons)
        t["en"] = t["en"].replace("{#takeaways}", "{#takeaways}\n\nThe balance sheet confirms it: net worth has fallen.", 1)
    elif m == "proscrit":       # vocabulaire proscrit par le verrou PRO-20261003-195656
        t["fr"] = t["fr"].replace("**Ce qui résiste.**", "**Ce qui résiste.** Les déficits ont financé des dépenses courantes.", 1)
    return t


# ------------------------------------------------------------------ figures
# Textes des figures, par langue. Phase B (11/10/2026) : mentions de méthode DANS chaque export (contre-expertise de
# phase A, § 3), pied commun (pied_figure), version téléphone recomposée pour chaque figure.
WM, FS_MIN = pied_figure.LARGEUR_M, pied_figure.POLICE_M
# Pied téléphone (corrections de clôture, 11/10/2026) : l'URL de la ressource, non plus le domaine. Trop longue pour une ligne de
# 300 px, elle passe sur deux lignes, domaine puis chemin : l'espace avant le chemin n'est qu'un point de coupe pour
# pied_figure.coupe (qui ne coupe qu'aux espaces), et le contrôle de main() vérifie que le chemin tient seul sur sa ligne.
_DOMAINE, _CHEMIN = PAGE_URL.split("/", 1)
_DOMAINE_EN, _CHEMIN_EN = PAGE_URL_EN.split("/", 1)
LICENCES_M = {"fr": "Compilation Stéphane Lalut, CC BY 4.0 · %s /%s" % (_DOMAINE, _CHEMIN),
              "en": "Compiled by Stéphane Lalut, CC BY 4.0 · %s /%s" % (_DOMAINE_EN, _CHEMIN_EN)}
URL_M = {"fr": PAGE_URL, "en": PAGE_URL_EN}
TXT = {
 "fr": dict(
  pc=" %",
  act_titre="France : à quoi ont correspondu les déficits publics",
  act_titre_m="France : à quoi ont correspondu les déficits publics",
  act_desc=("Pour chaque période, le besoin de financement des administrations publiques, ramené à 100, partagé entre trois emplois : "
            "acquisitions nettes d'actifs (investissement moins usure du capital, terrains, stocks), transferts en capital nets vers "
            "d'autres secteurs, dépenses courantes non couvertes par les recettes courantes. "),
  act_desc_per="%s : %s points de PIB de besoin de financement cumulé, dont %s %% en actifs, %s %% en transferts, %s %% en dépenses courantes.",
  act_leg=("actifs (investissement net, terrains)", "transferts en capital", "dépenses courantes non couvertes"),
  act_axe=("besoin de financement cumulé,", "points de PIB"),
  act_part_m="actifs : %s %%",
  act_besoin_m="%s points de PIB cumulés",
  act_src="Eurostat gov_10a_main (B9, B8N, P51G, P51C, NP, P52_P53, D9), gov_10dd_edpt1 (PIB), administrations publiques, France, %s-%s",
  act_src_m="Eurostat (gov_10a_main, gov_10dd_edpt1), administrations publiques, France, %s-%s",
  act_note=("Identité comptable : elle dit à quoi les déficits ont correspondu, non ce qu'ils ont produit. Ces parts ne mesurent ni la "
            "rentabilité des dépenses ni leur utilité pour les générations futures ; l'éducation n'y est pas un actif."),
  det_titre="Qui détient la dette publique française",
  det_titre_m="Qui détient la dette publique française",
  det_desc=("Parts de la dette des administrations publiques détenues par les non-résidents, par la Banque de France et par les autres "
            "résidents, de %s à %s. Non-résidents : %s %% en %s, %s %% au plus haut en %s, %s %% en %s. Banque de France : %s %% en %s, "
            "%s %% en %s, %s %% en %s. Autres résidents : %s %% en %s."),
  det_lib=("non-résidents", "autres résidents", "Banque de France"),
  det_moitie="moitié de la dette",
  det_en_m="en %s :",
  det_src="BCE, statistiques de finances publiques (GFS) : dette de Maastricht des administrations publiques par secteur détenteur, France, %s-%s",
  det_src_m="BCE, statistiques de finances publiques (GFS), France, %s-%s",
  det_note=("Résidence du détenteur enregistré (fonds, dépositaire), non de l'épargnant final ; valeur nominale. La part des non-résidents "
            "ne mesure pas, à elle seule, le transfert net de richesse vers l'étranger."),
  vie_titre="France : l'ajustement qui stabiliserait la dette (indicateur S2) et ses deux composantes",
  vie_titre_m="France : l'ajustement S2 qui stabiliserait la dette",
  vie_groupes=("Scénario de base", "Productivité plus faible", "Risque santé et dépendance", "Édition précédente (%s)"),
  vie_prec="autre jeu de projections",
  vie_total="S2 : %s",
  vie_desc=("Ajustement permanent du solde primaire structurel, à partir de %s, qui stabiliserait la dette sur un horizon infini, en points "
            "de PIB, partagé entre la position budgétaire de départ et l'effet des variations futures des dépenses liées à l'âge, selon la "
            "Commission européenne (%s), dans trois scénarios et dans l'édition précédente. "),
  vie_desc_g="%s : S2 %s, dont position de départ %s et effet des variations futures des dépenses liées à l'âge %s.",
  vie_leg=("position budgétaire de départ (dépenses actuelles comprises, retraites incluses)",
           "effet des variations futures des dépenses liées à l'âge"),
  vie_leg_m=("position de départ (dépenses actuelles, retraites incluses)", "effet des variations futures des dépenses liées à l'âge"),
  vie_axe="points de PIB",
  vie_src="Commission européenne, %s, tableaux par pays : indicateur S2 et ses deux composantes, France",
  vie_src_m="Commission européenne, %s, France",
  vie_note=("La dépense actuelle de retraite est dans la position de départ ; le second terme ne mesure que l'évolution projetée des "
            "dépenses liées à l'âge. L'édition précédente est un autre jeu de projections, non une évolution observée. Projections "
            "conditionnelles aux hypothèses de la Commission (Ageing Report 2024, réforme des retraites de 2023 appliquée)."),
  montre=[("actifs", "Depuis trente ans, la part des déficits français qui a correspondu, dans les comptes, à un accroissement net des "
                     "actifs publics a baissé de décennie en décennie ; la plus grande part a correspondu à des dépenses courantes que les "
                     "recettes courantes ne couvraient pas. Ces parts ne mesurent ni la rentabilité ni l'utilité des dépenses."),
          ("detention", "La part de la dette française détenue par des non-résidents dépasse la moitié la plupart des années depuis le "
                        "début des années 2000 ; les achats de la Banque de France l'ont fait reculer, le reflux de ses avoirs la fait "
                        "remonter. C'est la résidence des détenteurs enregistrés, non une mesure du fardeau."),
          ("vieillissement", "Selon l'indicateur S2 de la Commission, l'ajustement qui stabiliserait la dette française correspond surtout "
                             "à la position budgétaire de départ, où figure la dépense de retraite actuelle ; l'effet des variations "
                             "futures des dépenses liées à l'âge y est faible dans le scénario de base, nettement plus fort dans le "
                             "scénario de risque.")]),
 "en": dict(
  pc="%",
  act_titre="France: what public deficits corresponded to",
  act_titre_m="France: what public deficits corresponded to",
  act_desc=("For each period, general government net borrowing, set at 100, split between three uses: net acquisitions of assets "
            "(investment minus depreciation of capital, land, inventories), net capital transfers to other sectors, current spending "
            "not covered by current revenue. "),
  act_desc_per="%s: %s points of GDP of cumulated net borrowing, of which %s%% in assets, %s%% in transfers, %s%% in current spending.",
  act_leg=("assets (net investment, land)", "capital transfers", "uncovered current spending"),
  act_axe=("cumulated net borrowing,", "points of GDP"),
  act_part_m="assets: %s%%",
  act_besoin_m="%s points of GDP, cumulated",
  act_src="Eurostat gov_10a_main (B9, B8N, P51G, P51C, NP, P52_P53, D9), gov_10dd_edpt1 (GDP), general government, France, %s-%s",
  act_src_m="Eurostat (gov_10a_main, gov_10dd_edpt1), general government, France, %s-%s",
  act_note=("Accounting identity: it shows what deficits corresponded to, not what they produced. These shares measure neither the "
            "return on spending nor its usefulness to future generations; education does not count as an asset."),
  det_titre="Who holds French public debt",
  det_titre_m="Who holds French public debt",
  det_desc=("Shares of general government debt held by non-residents, by the Banque de France and by other residents, from %s to %s. "
            "Non-residents: %s%% in %s, %s%% at the peak in %s, %s%% in %s. Banque de France: %s%% in %s, %s%% in %s, %s%% in %s. "
            "Other residents: %s%% in %s."),
  det_lib=("non-residents", "other residents", "Banque de France"),
  det_moitie="half of the debt",
  det_en_m="in %s:",
  det_src="ECB, government finance statistics (GFS): Maastricht debt of general government by holding sector, France, %s-%s",
  det_src_m="ECB, government finance statistics (GFS), France, %s-%s",
  det_note=("Residence of the registered holder (fund, custodian), not of the final saver; nominal value. The non-resident share does "
            "not, on its own, measure the net transfer of wealth abroad."),
  vie_titre="France: the adjustment that would stabilise the debt (S2 indicator) and its two components",
  vie_titre_m="France: the S2 adjustment that would stabilise the debt",
  vie_groupes=("Baseline scenario", "Lower productivity", "Health and long-term care risk", "Previous edition (%s)"),
  vie_prec="another set of projections",
  vie_total="S2: %s",
  vie_desc=("Permanent adjustment of the structural primary balance, from %s, that would stabilise the debt over an infinite horizon, in "
            "points of GDP, split between the initial budgetary position and the effect of future changes in age-related spending, "
            "according to the European Commission (%s), in three scenarios and in the previous edition. "),
  vie_desc_g="%s: S2 %s, of which initial position %s and effect of future changes in age-related spending %s.",
  vie_leg=("initial budgetary position (current spending included, pensions too)",
           "effect of future changes in age-related spending"),
  vie_leg_m=("initial position (current spending, pensions included)", "effect of future changes in age-related spending"),
  vie_axe="points of GDP",
  vie_src="European Commission, %s, country tables: S2 indicator and its two components, France",
  vie_src_m="European Commission, %s, France",
  vie_note=("Current pension spending is in the initial position; the second term measures only the projected change in age-related "
            "spending. The previous edition is another set of projections, not an observed change. Projections conditional on the "
            "Commission's assumptions (2024 Ageing Report, 2023 pension reform applied)."),
  montre=[("actifs", "Over thirty years, the share of French deficits that corresponded, in the accounts, to a net increase in public "
                     "assets has fallen decade by decade; the largest share corresponded to current spending that current revenue did "
                     "not cover. These shares measure neither the return on spending nor its usefulness."),
          ("detention", "The share of French debt held by non-residents has been above half in most years since the early 2000s; "
                        "Banque de France purchases pushed it down, and the ebb of its holdings is pushing it back up. This is the "
                        "residence of registered holders, not a measure of the burden."),
          ("vieillissement", "According to the Commission's S2 indicator, the adjustment that would stabilise French debt mainly "
                             "reflects the initial budgetary position, which includes current pension spending; the effect of future "
                             "changes in age-related spending is small in the baseline scenario and much larger in the risk scenario.")]),
}


def esc(s: str) -> str:
    return html.escape(s, quote=False)


def entete(h, ident, titre, desc):
    return ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 %d %d" role="img" font-variant-numeric="tabular-nums" '
            'aria-labelledby="%s-t %s-d" font-family="%s">' % (W, h, ident, ident, FONT),
            '<title id="%s-t">%s</title><desc id="%s-d">%s</desc>' % (ident, esc(titre), ident, esc(desc)),
            '<rect width="%d" height="%d" fill="#ffffff"/>' % (W, h),
            '<text x="0" y="18" font-size="15" font-weight="600" fill="%s">%s</text>' % (INK, esc(titre))]


def pied(e, y0, source, note):
    return pied_figure.pied(e, y0, W, [(source, INK2, "pied-source"), (note, INK2, "pied-note"),
                                       (LICENCES[LANG], MUTED, "pied-licence")], GRID)


def pied_m(e, y0, source, note):
    return pied_figure.pied_mobile(e, y0, [(source, INK2, "pied-source"), (note, INK2, "pied-note"),
                                           (LICENCES_M[LANG], MUTED, "pied-licence")], GRID)


def legende(e, items, y=36, x=0):
    for col, lib_ in items:
        e.append('<rect x="%d" y="%d" width="10" height="10" fill="%s"/>' % (x, y, col))
        e.append('<text x="%d" y="%d" font-size="11" fill="%s">%s</text>' % (x + 14, y + 9, INK2, esc(lib_)))
        x += 14 + 6.0 * len(lib_) + 18


def desc_de(svg: str) -> str:
    return html.unescape(re.search(r"<desc[^>]*>(.*?)</desc>", svg, re.S).group(1))


def rang_actifs(total, dec, A):
    return [(A["d%d_lib" % k], d) for k, d in enumerate(dec, 1)] + [("%s-%s" % (A["a0"], A["fin"]), total)]


def fig_actifs(total, dec, A):
    """Signature : 100 % du besoin de financement de chaque période, partagé entre ses trois emplois."""
    X0, X1 = 132, W - 64
    T = TXT[LANG]
    rang = rang_actifs(total, dec, A)
    desc = (T["act_desc"]
            + " ".join(T["act_desc_per"]
                       % (lib_, fr(d["besoin"]), fr(d["part_actifs"], 0), fr(d["part_transferts"], 0), fr(d["part_desepargne"], 0)) for lib_, d in rang))
    e = entete(0, "ga", T["act_titre"], desc)
    legende(e, [(BLEU, T["act_leg"][0]), (GRIS_CLAIR, T["act_leg"][1]), (ORANGE, T["act_leg"][2])])
    e.append('<text x="%d" y="72" font-size="10" fill="%s" text-anchor="end">%s</text>' % (W, MUTED, T["act_axe"][0]))
    e.append('<text x="%d" y="84" font-size="10" fill="%s" text-anchor="end">%s</text>' % (W, MUTED, T["act_axe"][1]))
    top, hb, pas = 96, 34, 50
    y = top
    for k, (lib_, d) in enumerate(rang):
        y = top + pas * k + (14 if k == 3 else 0)
        if k == 3:
            e.append('<line x1="0" y1="%.1f" x2="%d" y2="%.1f" stroke="%s"/>' % (y - 12, W, y - 12, GRID))
        e.append('<text x="0" y="%.1f" font-size="12" font-weight="600" fill="%s">%s</text>' % (y + hb / 2 + 4, INK, esc(lib_)))
        x = X0
        for cle, col, txt in (("part_actifs", BLEU, "#ffffff"), ("part_transferts", GRIS_CLAIR, INK), ("part_desepargne", ORANGE, "#ffffff")):
            w = (X1 - X0) * max(d[cle], 0) / 100
            e.append('<rect x="%.1f" y="%.1f" width="%.1f" height="%d" fill="%s"/>' % (x, y, max(w, 0.6), hb, col))
            if w >= 30:
                e.append('<text x="%.1f" y="%.1f" font-size="11" font-weight="600" fill="%s" text-anchor="middle">%s%s</text>'
                         % (x + w / 2, y + hb / 2 + 4, txt, fr(d[cle], 0), T["pc"]))
            x += w
        e.append('<text x="%d" y="%.1f" font-size="11" font-weight="600" fill="%s" text-anchor="end">%s</text>' % (W, y + hb / 2 + 4, INK2, fr(d["besoin"])))
    return pied(e, y + hb + 12, T["act_src"] % (A["a0"], A["fin"]), T["act_note"])


def fig_actifs_m(total, dec, A, svg_d):
    """Téléphone : une barre par période, sa part en actifs écrite au-dessus ; légende sur trois lignes."""
    T = TXT[LANG]
    e, y = pied_figure.tete_mobile("gam", T["act_titre_m"], desc_de(svg_d), INK, FONT)
    for col, lib_ in zip((BLEU, GRIS_CLAIR, ORANGE), T["act_leg"]):
        y += 20
        e.append('<rect x="0" y="%d" width="12" height="12" fill="%s"/>' % (y - 11, col))
        e.append('<text x="18" y="%d" font-size="%s" fill="%s">%s</text>' % (y, FS_MIN, INK2, esc(lib_)))
    y += 14
    for k, (lib_, d) in enumerate(rang_actifs(total, dec, A)):
        y += 26 + (8 if k == 3 else 0)
        if k == 3:
            e.append('<line x1="0" y1="%d" x2="%d" y2="%d" stroke="%s"/>' % (y - 22, WM, y - 22, GRID))
        e.append('<text x="0" y="%d" font-size="15" font-weight="600" fill="%s">%s</text>' % (y, INK, esc(lib_)))
        e.append('<text x="%d" y="%d" font-size="15" font-weight="600" fill="%s" text-anchor="end">%s</text>'
                 % (WM, y, BLEU, esc(T["act_part_m"] % fr(d["part_actifs"], 0))))
        y += 18
        e.append('<text x="0" y="%d" font-size="%s" fill="%s">%s</text>' % (y, FS_MIN, MUTED, esc(T["act_besoin_m"] % fr(d["besoin"]))))
        y += 8
        x = 0.0
        for cle, col, txt in (("part_actifs", BLEU, "#ffffff"), ("part_transferts", GRIS_CLAIR, INK), ("part_desepargne", ORANGE, "#ffffff")):
            w = WM * max(d[cle], 0) / 100
            e.append('<rect x="%.1f" y="%d" width="%.1f" height="28" fill="%s"/>' % (x, y, max(w, 0.6), col))
            if w >= 44 and cle != "part_actifs":
                e.append('<text x="%.1f" y="%d" font-size="%s" font-weight="600" fill="%s" text-anchor="middle">%s%s</text>'
                         % (x + w / 2, y + 19, FS_MIN, txt, fr(d[cle], 0), T["pc"]))
            x += w
        y += 28
    return pied_m(e, y + 16, T["act_src_m"] % (A["a0"], A["fin"]), T["act_note"])


def desc_detention(A):
    T = TXT[LANG]
    return T["det_desc"] % (A["h0"], A["h_fin"], A["nr0"], A["h0"], A["nr_max"], A["nr_max_annee"], A["nr_fin"], A["h_fin"],
                            A["bc0"], A["h0"], A["bc_max"], A["bc_max_annee"], A["bc_fin"], A["h_fin"], A["ar_fin"], A["h_fin"])


def fig_detention(det, A):
    X0, X1, TOP, BAS = 40, W - 128, 60, 276
    T = TXT[LANG]
    e = entete(0, "gd", T["det_titre"], desc_detention(A))
    n = len(det)
    X = lambda k: X0 + (X1 - X0) * k / (n - 1)
    Y = lambda v: BAS - (BAS - TOP) * v / 100
    bandes = (("part_banque_centrale", BLEU), ("part_autres_residents", GRIS_CLAIR), ("part_non_residents", ORANGE))
    bas_cumul = [0.0] * n
    milieux = {}
    for cle, col in bandes:
        haut = [bas_cumul[k] + det[k][cle] for k in range(n)]
        pts = ["%.1f,%.1f" % (X(k), Y(haut[k])) for k in range(n)] + ["%.1f,%.1f" % (X(k), Y(bas_cumul[k])) for k in reversed(range(n))]
        e.append('<polygon points="%s" fill="%s"/>' % (" ".join(pts), col))
        milieux[cle] = (bas_cumul[-1] + haut[-1]) / 2
        bas_cumul = haut
    for g_ in range(0, 101, 25):
        e.append('<line x1="%d" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="0.6" stroke-opacity="0.8"/>' % (X0, Y(g_), X1, Y(g_), "#ffffff"))
        e.append('<text x="%d" y="%.1f" font-size="10" fill="%s" text-anchor="end">%d%s</text>' % (X0 - 6, Y(g_) + 3, MUTED, g_, T["pc"]))
    e.append('<line x1="%d" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="1.2" stroke-dasharray="4 3"/>' % (X0, Y(50), X1, Y(50), INK))
    for k, r in enumerate(det):
        if r["annee"] % 5 == 0 or k == n - 1:
            e.append('<text x="%.1f" y="%d" font-size="10" fill="%s" text-anchor="middle">%d</text>' % (X(k), BAS + 14, INK2, r["annee"]))
    last = det[-1]
    for cle, lib_ in zip(("part_non_residents", "part_autres_residents", "part_banque_centrale"), T["det_lib"]):
        e.append('<text x="%d" y="%.1f" font-size="11" fill="%s">%s</text>' % (X1 + 8, Y(milieux[cle]) - 2, INK2, esc(lib_)))
        e.append('<text x="%d" y="%.1f" font-size="11" font-weight="600" fill="%s">%s%s</text>' % (X1 + 8, Y(milieux[cle]) + 11, INK, fr(last[cle], 0), T["pc"]))
    e.append('<text x="%d" y="%.1f" font-size="10" fill="%s">%s</text>' % (X0 + 4, Y(50) - 4, INK, T["det_moitie"]))
    return pied(e, BAS + 26, T["det_src"] % (A["h0"], A["h_fin"]), T["det_note"])


def fig_detention_m(det, A, svg_d):
    """Téléphone : aires empilées sur toute la largeur, valeurs de la dernière année en liste sous le graphique."""
    T = TXT[LANG]
    e, y = pied_figure.tete_mobile("gdm", T["det_titre_m"], desc_de(svg_d), INK, FONT)
    X0, X1, TOP = 50, WM - 4, y + 16
    BAS = TOP + 190
    n = len(det)
    X = lambda k: X0 + (X1 - X0) * k / (n - 1)
    Y = lambda v: BAS - (BAS - TOP) * v / 100
    bas_cumul = [0.0] * n
    for cle, col in (("part_banque_centrale", BLEU), ("part_autres_residents", GRIS_CLAIR), ("part_non_residents", ORANGE)):
        haut = [bas_cumul[k] + det[k][cle] for k in range(n)]
        pts = ["%.1f,%.1f" % (X(k), Y(haut[k])) for k in range(n)] + ["%.1f,%.1f" % (X(k), Y(bas_cumul[k])) for k in reversed(range(n))]
        e.append('<polygon points="%s" fill="%s"/>' % (" ".join(pts), col))
        bas_cumul = haut
    for g_ in (0, 50, 100):
        e.append('<text x="%d" y="%.1f" font-size="%s" fill="%s" text-anchor="end">%d%s</text>' % (X0 - 4, Y(g_) + 5, FS_MIN, MUTED, g_, T["pc"]))
    e.append('<line x1="%d" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="1.4" stroke-dasharray="4 3"/>' % (X0, Y(50), X1, Y(50), INK))
    for k, r in enumerate(det):
        if k in (0, n - 1) or r["annee"] == 2010:
            e.append('<text x="%.1f" y="%d" font-size="%s" fill="%s" text-anchor="%s">%d</text>'
                     % (X(k), BAS + 18, FS_MIN, MUTED, "start" if k == 0 else ("end" if k == n - 1 else "middle"), r["annee"]))
    y = BAS + 44
    e.append('<text x="0" y="%d" font-size="%s" fill="%s">%s</text>' % (y, FS_MIN, INK2, esc(T["det_en_m"] % A["h_fin"])))
    last = det[-1]
    for cle, col, lib_ in (("part_non_residents", ORANGE, T["det_lib"][0]), ("part_autres_residents", GRIS_CLAIR, T["det_lib"][1]),
                           ("part_banque_centrale", BLEU, T["det_lib"][2])):
        y += 22
        e.append('<rect x="0" y="%d" width="12" height="12" fill="%s"/>' % (y - 11, col))
        e.append('<text x="18" y="%d" font-size="15" fill="%s">%s</text>' % (y, INK2, esc(lib_)))
        e.append('<text x="%d" y="%d" font-size="15" font-weight="600" fill="%s" text-anchor="end">%s%s</text>' % (WM, y, INK, fr(last[cle]), T["pc"]))
    return pied_m(e, y + 16, T["det_src_m"] % (A["h0"], A["h_fin"]), T["det_note"])


def groupes_s2(S2fr, A):
    T = TXT[LANG]
    g1, g2, g3, g4 = T["vie_groupes"]
    return [(g1, S2fr["S2"]["base"], S2fr["IBP"]["base"], S2fr["CoA"]["base"]),
            (g2, S2fr["S2"]["productivite"], S2fr["IBP"]["productivite"], S2fr["CoA"]["productivite"]),
            (g3, S2fr["S2"]["risque"], S2fr["IBP"]["risque"], S2fr["CoA"]["risque"]),
            (g4 % A["dsm_prec_an"], S2fr["S2"]["base_precedente"], S2fr["IBP"]["base_precedente"], S2fr["CoA"]["base_precedente"])]


def fig_vieillissement(S2fr, C, A):
    X0, X1 = 230, W - 40
    T = TXT[LANG]
    groupes = groupes_s2(S2fr, A)
    desc = (T["vie_desc"] % (A["s2_an"], A["dsm_lib"])
            + " ".join(T["vie_desc_g"] % (g, fr(s), fr(a), fr(b)) for g, s, a, b in groupes))
    e = entete(0, "gv", T["vie_titre"], desc)
    legende(e, [(ORANGE, T["vie_leg"][0])], y=34)
    legende(e, [(GRIS, T["vie_leg"][1])], y=50)
    vmax = int(max(max(a, b) for _, _, a, b in groupes)) + 1
    Xv = lambda v: X0 + (X1 - X0) * v / vmax
    top, pas, hb = 80, 50, 15
    for g_ in range(0, vmax + 1):
        e.append('<line x1="%.1f" y1="%d" x2="%.1f" y2="%d" stroke="%s" stroke-width="%s"/>' % (Xv(g_), top - 6, Xv(g_), top + pas * len(groupes) - 14, GRID, 1.4 if g_ == 0 else 0.6))
        e.append('<text x="%.1f" y="%d" font-size="10" fill="%s" text-anchor="middle">%d</text>' % (Xv(g_), top + pas * len(groupes), MUTED, g_))
    for k, (lib_, s, a, b) in enumerate(groupes):
        y = top + pas * k
        e.append('<text x="0" y="%d" font-size="11.5" font-weight="600" fill="%s">%s</text>' % (y + 12, INK, esc(lib_)))
        sous = T["vie_total"] % fr(s) + (" · " + T["vie_prec"] if k == 3 else "")
        e.append('<text x="0" y="%d" font-size="11" fill="%s">%s</text>' % (y + 27, INK2, esc(sous)))
        for j, (v, col) in enumerate(((a, ORANGE), (b, GRIS))):
            yy = y + j * (hb + 2)
            x0_, x1_ = sorted((Xv(0), Xv(v)))
            e.append('<rect x="%.1f" y="%d" width="%.1f" height="%d" fill="%s"/>' % (x0_, yy, max(x1_ - x0_, 0.6), hb, col))
            e.append('<text x="%.1f" y="%d" font-size="10" font-weight="600" fill="%s">%s</text>' % (max(x1_, Xv(0)) + 4, yy + 11, INK, fr(v)))
    yb = top + pas * len(groupes)
    e.append('<text x="%d" y="%d" font-size="10" fill="%s" text-anchor="end">%s</text>' % (X1, yb + 15, MUTED, T["vie_axe"]))
    return pied(e, yb + 24, T["vie_src"] % A["dsm_lib"], T["vie_note"])


def fig_vieillissement_m(S2fr, C, A, svg_d):
    """Téléphone : par scénario, le total S2 puis ses deux composantes, valeurs écrites au bout des barres."""
    T = TXT[LANG]
    e, y = pied_figure.tete_mobile("gvm", T["vie_titre_m"], desc_de(svg_d), INK, FONT)
    for col, lib_ in ((ORANGE, T["vie_leg_m"][0]), (GRIS, T["vie_leg_m"][1])):
        lignes = pied_figure.coupe(lib_, 32)
        for k, l in enumerate(lignes):
            y += 18
            if k == 0:
                e.append('<rect x="0" y="%d" width="12" height="12" fill="%s"/>' % (y - 11, col))
            e.append('<text x="18" y="%d" font-size="%s" fill="%s">%s</text>' % (y, FS_MIN, INK2, esc(l)))
        y += 4
    groupes = groupes_s2(S2fr, A)
    vmax = int(max(max(a, b) for _, _, a, b in groupes)) + 1
    X0, X1 = 0, WM - 40
    Xv = lambda v: X0 + (X1 - X0) * v / vmax
    for k, (lib_, s, a, b) in enumerate(groupes):
        y += 30
        e.append('<text x="0" y="%d" font-size="15" font-weight="600" fill="%s">%s</text>' % (y, INK, esc(lib_)))
        y += 19
        sous = T["vie_total"] % fr(s) + (" · " + T["vie_prec"] if k == 3 else "")
        e.append('<text x="0" y="%d" font-size="%s" font-weight="600" fill="%s">%s</text>' % (y, FS_MIN, INK2, esc(sous)))
        y += 8
        for v, col in ((a, ORANGE), (b, GRIS)):
            e.append('<rect x="%.1f" y="%d" width="%.1f" height="16" fill="%s"/>' % (Xv(0), y, max(Xv(v) - Xv(0), 0.6), col))
            e.append('<text x="%.1f" y="%d" font-size="%s" font-weight="600" fill="%s">%s</text>' % (Xv(max(v, 0)) + 5, y + 13, FS_MIN, INK, fr(v)))
            y += 19
    e.append('<text x="0" y="%d" font-size="%s" fill="%s">%s</text>' % (y + 14, FS_MIN, MUTED, esc(T["vie_axe"])))
    return pied_m(e, y + 28, T["vie_src_m"] % A["dsm_lib"], T["vie_note"])


def fiches(figs):
    out = {}
    for lang in ("fr", "en"):
        out[lang] = []
        for ident, montre in TXT[lang]["montre"]:
            f = "dette-generations-%s%s" % (ident, SUFFIXE[lang])
            svg = figs[f + ".svg"]
            titre = html.unescape(re.search(r"<title[^>]*>(.*?)</title>", svg).group(1))
            out[lang].append(dict(id=ident, fichier=f, titre=titre, montre=montre,
                                  source=pied_figure.lire(svg, "pied-source"), precaution=pied_figure.lire(svg, "pied-note")))
    return out


def csv_texte(total, dec, annees, cp, pat, det, det_pays, S2, C, etat, ss):
    """Format long : les tableaux n'ont pas les mêmes colonnes."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["tableau", "pays", "periode", "variable", "valeur", "unite"])
    cles = ("besoin", "actifs", "investissement_net", "transferts", "aides_investissement", "desepargne")
    for d in [total] + dec:
        for k in cles:
            w.writerow(["france_emplois_du_besoin_de_financement", "FR", "%d-%d" % (d["debut"], d["fin"]), k, "%.3f" % d[k], "points de PIB cumulés"])
    for r in annees:
        for k in cles:
            w.writerow(["france_annuel", "FR", r["annee"], k, "%.3f" % r[k], "% du PIB"])
    for r in ss["annees"]:
        for k in ("net_apu", "net_loc", "net_autres", "fbcf_loc"):
            w.writerow(["france_investissement_net_sous_secteurs", "FR", r["annee"], k, "%.3f" % r[k], "% du PIB"])
    for d in cp:
        for k in ("besoin", "actifs", "transferts", "desepargne"):
            w.writerow(["ue_emplois_du_besoin_de_financement", d["pays"], "%d-%d" % (d["debut"], d["fin"]), k, "%.3f" % d[k], "points de PIB cumulés"])
    for r in pat:
        for k in ("actifs_produits", "actifs_non_produits", "terrains", "valeur_financiere_nette", "patrimoine_net"):
            w.writerow(["france_patrimoine_apu", "FR", r["annee"], k, "%.3f" % r[k], "% du PIB"])
    for r in det:
        for k in ("dette", "non_residents", "banque_centrale", "autres_residents"):
            w.writerow(["france_detention", "FR", r["annee"], k, "%.3f" % r[k], "% du PIB"])
    for x in etat["fins"]:
        w.writerow(["france_etat_negociable_non_residents", "FR", x["annee"], "part_non_residents_fin_annee", "%.1f" % x["part"], "% de la dette négociable de l'État"])
    for p, x in sorted(det_pays.items()):
        w.writerow(["ue_part_non_residents", p, x["annee"], "part_non_residents", "%.3f" % x["part_non_residents"], "% de la dette"])
    for g, x in sorted(S2.items()):
        for j in JEUX:
            for k in ("S2", "IBP", "CoA", "pensions", "sante", "dependance", "education"):
                w.writerow(["commission_s2", g, j, k, "%.3f" % x[k][j], "points de PIB"])
    for k in ("depenses0", "depenses1", "ressources0", "ressources1", "solde0", "solde1"):
        w.writerow(["cor_retraites", "FR", C["a0"] if k.endswith("0") else C["a1"], k[:-1], "%.3f" % C[k], "% du PIB"])
    return buf.getvalue()


# ------------------------------------------------------------------ main
def mutation_donnees(cp, dec) -> str | None:
    """Témoins positifs des gardes de données nouvelles : un pays nommé change de côté ; l'un des deux mouvements de la
    baisse par décennie cesse de tenir (actifs de la deuxième décennie +10 % ; besoin de la troisième +10 %)."""
    m = next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith("--mutation=")), None)
    if m == "cote":
        x = next(y for y in cp if y["pays"] == CITE_SUP[0])
        x["part_actifs"] = 5.0
    elif m == "deux_temps":
        dec[1]["actifs"] = dec[0]["actifs"] * 1.10
    elif m == "deux_temps_2":
        dec[2]["besoin"] = dec[1]["besoin"] * 1.10
    elif m not in (None, "bilan", "bilan_mot", "vieillissement", "proscrit", "recul_local"):
        fail("mutation inconnue : %s" % m)
    return m


def main() -> int:
    check = "--check" in sys.argv[1:] or any(a.startswith("--mutation") for a in sys.argv[1:])
    if not check:
        try:
            import cairosvg  # noqa: F401
        except ImportError:
            fail("cairosvg absent : SVG et PNG se produisent ensemble ou pas du tout (pip install cairosvg)")
    fin, total, dec, annees, cp, pat, det, det_pays, det_controle, S2, C, cons, etat, ss = calcul()
    m = mutation_donnees(cp, dec)
    use("en")
    A_en = affichage(fin, total, dec, annees, cp, pat, det, det_pays, S2, C, cons, etat, ss)[0]
    use("fr")
    A, aux = affichage(fin, total, dec, annees, cp, pat, det, det_pays, S2, C, cons, etat, ss)
    if set(A) != set(A_en):
        fail("blocs affichage et affichage_en : cles differentes (%s)" % sorted(set(A) ^ set(A_en)))
    vides = sorted(k for k in A if bool(A[k]) != bool(A_en[k]))
    if vides:
        fail("blocs affichage et affichage_en : valeur vide dans une seule langue (%s)" % vides)
    n = gardes(fin, total, dec, annees, cp, pat, det, det_pays, S2, C, aux)
    textes = mutation_texte({"fr": PAGE_FR.read_text(encoding="utf-8"), "en": PAGE_EN.read_text(encoding="utf-8")})
    txt_fig = {lang: "\n".join(v if isinstance(v, str) else " ".join(map(str, v)) for k, v in TXT[lang].items() if k != "pc")
               for lang in ("fr", "en")}
    nt = gardes_texte(textes, txt_fig)
    if m:
        fail("mutation %s : aucune garde n'a mordu" % m)
    log("France %s-%s : besoin %.1f pts = actifs %.1f (%.0f %%) + transferts %.1f + depenses courantes %.1f ; non-residents %.1f %% (%s) ; "
        "S2 %.2f = IBP %.2f + CoA %.2f ; %d annees-pays ecartees sur %d"
        % (A["a0"], A["fin"], total["besoin"], total["actifs"], total["part_actifs"], total["transferts"], total["desepargne"],
           det[-1]["part_non_residents"], A["h_fin"], S2["FR"]["S2"]["base"], S2["FR"]["IBP"]["base"], S2["FR"]["CoA"]["base"],
           cons["ecartees"], cons["annees_pays_calculables"]))
    figs, tableaux = {}, {}
    for lang, AA in (("fr", A), ("en", A_en)):
        use(lang)
        x_ = SUFFIXE[lang]
        figs["dette-generations-actifs%s.svg" % x_] = fig_actifs(total, dec, AA)
        figs["dette-generations-detention%s.svg" % x_] = fig_detention(det, AA)
        figs["dette-generations-vieillissement%s.svg" % x_] = fig_vieillissement(S2["FR"], C, AA)
        figs["dette-generations-actifs%s-m.svg" % x_] = fig_actifs_m(total, dec, AA, figs["dette-generations-actifs%s.svg" % x_])
        figs["dette-generations-detention%s-m.svg" % x_] = fig_detention_m(det, AA, figs["dette-generations-detention%s.svg" % x_])
        figs["dette-generations-vieillissement%s-m.svg" % x_] = fig_vieillissement_m(S2["FR"], C, AA, figs["dette-generations-vieillissement%s.svg" % x_])
        for fid in ("actifs", "detention", "vieillissement"):
            ec = pied_figure.parite(figs["dette-generations-%s%s.svg" % (fid, x_)], figs["dette-generations-%s%s-m.svg" % (fid, x_)], FS_MIN)
            if ec:
                fail("parite mobile %s %s : %s" % (fid, lang, " ; ".join(ec)))
            lic = [html.unescape(t) for t in re.findall(r'<text class="pied-licence"[^>]*>(.*?)</text>', figs["dette-generations-%s%s-m.svg" % (fid, x_)])]
            if not (len(lic) >= 2 and lic[-1] == "/" + URL_M[lang].split("/", 1)[1] and (lic[-2] + lic[-1]).endswith(URL_M[lang])):
                fail("pied telephone %s %s : l'URL de la ressource n'est pas entiere sur ses deux lignes (%s)" % (fid, lang, lic))
            for f_, larg in (("dette-generations-%s%s.svg" % (fid, x_), W), ("dette-generations-%s%s-m.svg" % (fid, x_), WM)):
                hc = pied_figure.hors_cadre(figs[f_], larg)
                if hc:
                    log("AVERTISSEMENT (estimation, le rendu tranche) %s : %s" % (f_, " | ".join(hc)))
        # Libellés du tableau dans la langue de la page (le shortcode gen-tableau les cite entre guillemets dans sa note).
        au_dela, sans_objet = ("au-delà du besoin", "sans objet") if lang == "fr" else ("above net borrowing", "not applicable")
        tableaux[lang] = [dict(pays=x["pays"], nom=nom_pays(x["pays"]), besoin=fr(x["besoin"]), actifs=fr(x["actifs"]), transferts=fr(x["transferts"]),
                               desepargne=fr(x["desepargne"]), part=(au_dela if x["part_actifs"] > 100 else fr(x["part_actifs"], 0) + pc())
                               if x["besoin"] > 0 else sans_objet)
                          for x in sorted(cp, key=lambda x: -(x["part_actifs"] if x["besoin"] > 0 else 1e9))]
    for page, sf in ((PAGE_FR, ""), (PAGE_EN, "-en")):
        ec = pied_figure.controler_page(page.read_text(encoding="utf-8"),
                                        {f: v for f, v in figs.items() if re.fullmatch(r"dette-generations-[a-z]+%s(-m)?\.svg" % sf, f)})
        if ec:
            fail("page %s : %s" % (page.name, " ; ".join(ec)))
    if check:   # après les figures : parité téléphone et hauteurs déclarées se contrôlent aussi sans écrire
        log("--check : %d gardes de donnees, %d gardes de texte passees (%d cles d'affichage), parite des 6 figures telephone "
            "et hauteurs des pages controlees, rien ecrit." % (n, nt, len(A)))
        return 0
    use("fr")
    tableau = tableaux["fr"]
    payload = {"meta": {"releve_le": datetime.now(timezone.utc).strftime("%Y-%m-%d"), "page": "https://" + PAGE_URL, "licence": "CC BY 4.0",
                        "perimetre": "administrations publiques (S.13), SEC 2010, monnaie nationale, PIB de la notification de déficit et de dette",
                        "identite": "-B9 = -B8N + (P51G - P51C + NP + P52_P53) + (D9PAY - D9REC)",
                        "pieces": {k: dict(v, sha256=hashlib.sha256((PIECES / v["fichier"]).read_bytes()).hexdigest()) for k, v in SOURCES.items()},
                        "definitions": {"besoin": "besoin de financement des administrations publiques (-B9), en % du PIB",
                                        "actifs": "acquisitions nettes d'actifs non financiers : investissement (P51G) moins consommation de capital fixe (P51C), plus terrains et autres actifs non produits (NP) et variation des stocks (P52_P53)",
                                        "investissement_net": "P51G - P51C seul (critère du test décisif)",
                                        "sous_secteurs": "investissement net (P51G - P51C) des administrations publiques (S13), des administrations locales (S1313) et des autres administrations (différence), en % du PIB ; fbcf_loc : FBCF des administrations locales ; ccf_reprise : consommation de capital fixe de la dernière année identique à celle de l'année précédente",
                                        "transferts": "transferts en capital versés moins reçus (D9PAY - D9REC) ; les recettes comprennent les impôts en capital (droits de succession)",
                                        "aides_investissement": "aides à l'investissement versées à d'autres secteurs (D92PAY), comprises dans les transferts versés",
                                        "desepargne": "épargne nette changée de signe (-B8N) : dépenses courantes, consommation de capital comprise, non couvertes par les recettes courantes",
                                        "temoin": "ratios comparés à ceux qu'Eurostat publie en %% du PIB ; tolérance %.2f point" % TOL,
                                        "detention": "BCE, GFS : dette de Maastricht par zone (W1 reste du monde, W2 résidents) et secteur détenteur (S121 banque centrale)",
                                        "S2": "ajustement permanent du solde primaire structurel, à partir de 2027, qui stabiliserait la dette sur un horizon infini, sans niveau de dette ni date fixés (à la différence de S1) ; IBP : position budgétaire initiale, écart entre le solde primaire structurel de départ et le solde qui stabiliserait le ratio, dépenses actuelles comprises ; CoA : effet des variations futures des dépenses liées à l'âge (pensions nettes des prélèvements, santé, dépendance, éducation), que la Commission nomme « coût du vieillissement » (Debt Sustainability Monitor 2025, § 3.1 et annexe A6)"}},
               "france": {"periodes": [total] + dec, "annees": annees, "patrimoine": pat, "detention": det, "sous_secteurs": ss},
               "ue": {"emplois": cp, "tableau": tableau, "tableau_en": tableaux["en"], "part_non_residents": det_pays, "s2": S2},
               "cor": C, "etat_negociable": etat, "controles": {"detention_bce_eurostat_annees": det_controle}, "conservation": cons,
               "affichage": A, "affichage_en": A_en}
    releve = payload["meta"]["releve_le"]
    licence = "CC BY 4.0 — compilation Stéphane Lalut ; sources Eurostat, Insee, BCE, Commission européenne, COR"
    if OUT_DATA.exists():
        try:
            prev = json.loads(OUT_DATA.read_text(encoding="utf-8"))
            p2 = dict(prev); p2.pop("releve_le", None); rj.sans_licence(p2, licence); p2["meta"] = dict(p2["meta"], releve_le=None)
            n2 = json.loads(json.dumps(dict(payload, meta=dict(payload["meta"], releve_le=None)), ensure_ascii=False))
            if prev.get("releve_le") and json.dumps(p2, sort_keys=True, ensure_ascii=False) == json.dumps(n2, sort_keys=True, ensure_ascii=False) \
                    and all((OUT_IMG / f).exists() and (OUT_IMG / f).read_text(encoding="utf-8") == s for f, s in figs.items()) \
                    and all((OUT_IMG / f.replace(".svg", ".png")).exists() for f in figs) \
                    and OUT_CSV.exists() and OUT_CSV.read_text(encoding="utf-8-sig") == csv_texte(total, dec, annees, cp, pat, det, det_pays, S2, C, etat, ss) \
                    and OUT_FIGURES.exists() and json.loads(OUT_FIGURES.read_text(encoding="utf-8")) == fiches(figs):
                log("Donnees et figures identiques : rien ecrit (releve_le conserve : %s)." % prev["releve_le"])
                return 0
            # Données identiques hors blocs de texte et textes de méthode : la date de relevé reste celle des DONNÉES, jamais
            # celle d'une retouche de texte ou de figure (11/10/2026 ; organe commun releve_jeu.py depuis le 17/10/2026).
            releve = rj.releve(OUT_DATA, payload, releve)
            payload["meta"]["releve_le"] = releve
        except (ValueError, KeyError):
            pass
    payload = {"releve_le": releve, "_licence": licence, **payload}
    txt = json.dumps(payload, ensure_ascii=False, indent=1)
    OUT_DATA.write_text(txt, encoding="utf-8")
    OUT_STATIC.write_text(txt, encoding="utf-8")
    OUT_CSV.write_text(csv_texte(total, dec, annees, cp, pat, det, det_pays, S2, C, etat, ss), encoding="utf-8-sig", newline="\n")
    import cairosvg
    for f, s in figs.items():
        (OUT_IMG / f).write_text(s, encoding="utf-8")
        cairosvg.svg2png(url=str(OUT_IMG / f), write_to=str(OUT_IMG / f.replace(".svg", ".png")), output_width=1440, background_color="white")
    OUT_FIGURES.write_text(json.dumps(fiches(figs), ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    log("Ecrit : data/ et static/dette_generations.json, static/dette_generations.csv, data/figures_generations.json, %d figures SVG + PNG" % len(figs))
    return 0


if __name__ == "__main__":
    sys.exit(main())
