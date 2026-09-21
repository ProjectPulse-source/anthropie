#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Figures et données de la page « Qui paie vraiment la dette publique ? ».

    python scripts/generer_figures_qui_paie.py

Arbitrage PRO-20260921-DIPTYQUE, étape 3 : trois figures, et non quatre par
symétrie avec la page du coût.

  F1  qui-paie-detention      Qui emprunte ? Qui détient les titres de l'État ?
                              DEUX panneaux, sans flux entre eux : A en valeur
                              nominale (INSEE), B en valeur de marché (Banque de
                              France via l'AFT). Les emboîter inviterait à
                              multiplier des parts de marché par des montants
                              nominaux.
  F2  qui-paie-redistribution Prélèvements et transferts publics par dixième de
                              niveau de vie (Insee, comptes nationaux distribués
                              2023). Décrit qui contribue et qui reçoit
                              aujourd'hui ; ne mesure pas l'incidence de la dette.
  F3  qui-paie-mecanismes     Schéma NON quantitatif des canaux de répartition.

SOURCES — une seule chacune, jamais recopiée à la main :
  - F1 : le registre du livre (03_LIVRES/dette-publique/01_SOURCE/
    REGISTRE_DONNEES.yaml, bloc infographie_detention) -- la même source que la
    p. 101 du livre. Le site n'en garde qu'un instantané, régénéré ici ; la
    révision annuelle du livre (maj_edition.py) relance ce script.
  - F2 : le fichier de l'Insee archivé dans scripts/sources/, contrôlé par son
    empreinte. Un fichier remplacé sans que l'empreinte suive arrête tout.

Sorties : static/img/qui-paie-*.svg et .png, data/figures_qui_paie.json (cartes
du bloc « Réutiliser »), data/ et static/qui_paie_donnees.json (données publiées,
et bloc `affichage` lu par le shortcode qp-val : aucun chiffre en dur dans la
prose de la page).

Les PNG sont rendus dans la MÊME exécution que les SVG, et le script s'arrête si
cairosvg manque : SVG et PNG ne peuvent donc pas diverger, et aucun contrôle de
dérive n'est nécessaire (à la différence des figures de la page du coût, rendues
en CI par une étape tolérante).

Les phrases qualitatives de la page (« plus de la moitié », « varient beaucoup
moins ») sont vérifiées ici contre les données : si l'une cesse d'être vraie, le
script s'arrête au lieu de laisser la prose mentir.
"""
from __future__ import annotations

import hashlib
import json
import os
import sys
from datetime import date
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

REPO = Path(__file__).resolve().parent.parent
REGISTRE = Path(os.environ.get(
    "REGISTRE_DETTE",
    r"D:\PRO\03_LIVRES\dette-publique\01_SOURCE\REGISTRE_DONNEES.yaml"))
XLSX = REPO / "scripts" / "sources" / "insee_IA118_comptes_distribues_2023.xlsx"
XLSX_SHA256 = "6f5f27809be12e192afdca5008022a0e0c63f832869d29be8dc837abbdeec0eb"
XLSX_URL = "https://www.insee.fr/fr/statistiques/8974371"

IMG = REPO / "static" / "img"
URL_PAGE = "stephane-lalut.com/qui-paie-la-dette-publique/"

# Palette dataviz de référence (slots 1-3, validés en clair par
# validate_palette.js le 2026-09-21 ; l'aqua avertit en contraste : d'où les
# étiquettes de valeur visibles et la vue en tableau sur la page).
C1, C2, C3 = "#2a78d6", "#eb6834", "#1baf7a"
INK, INK2, MUTED, GRID, AXIS = "#2b2a28", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"
FONT = "system-ui, -apple-system, Segoe UI, sans-serif"
NBSP = "\u00a0"
W = 720


def fail(msg: str) -> None:
    print("ECHEC : " + msg)
    sys.exit(1)


def fr(v: float, dec: int = 0) -> str:
    s = ("{:,." + str(dec) + "f}").format(v)
    return s.replace(",", NBSP).replace(".", ",")


def esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def txt(x, y, s, size=11, fill=INK2, anchor="start", weight=None):
    w = ' font-weight="%s"' % weight if weight else ""
    return ('<text x="%.1f" y="%.1f" font-size="%s" fill="%s" text-anchor="%s"%s>%s</text>'
            % (x, y, size, fill, anchor, w, esc(s)))


def cartouche(y0: float, lignes: list[tuple[str, str]]) -> list[str]:
    out = ['<line x1="0" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="1"/>'
           % (y0, W, y0, GRID)]
    for i, (s, col) in enumerate(lignes):
        out.append('<text x="0" y="%.1f" font-family="%s" font-size="9" fill="%s">%s</text>'
                   % (y0 + 13 + i * 12, FONT, col, esc(s)))
    return out


def entete(h: float, ident: str, titre: str, desc: str) -> list[str]:
    return ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 %d %d" role="img" '
            'aria-labelledby="%s-t %s-d" font-family="%s">' % (W, h, ident, ident, FONT),
            '<title id="%s-t">%s</title>' % (ident, esc(titre)),
            '<desc id="%s-d">%s</desc>' % (ident, esc(desc))]


# ------------------------------------------------------------------ données
def lire_registre() -> dict:
    try:
        import yaml
    except ImportError:
        fail("PyYAML absent : le registre du livre ne peut pas être lu")
    if not REGISTRE.is_file():
        fail("registre introuvable : %s (variable REGISTRE_DETTE pour un autre chemin)" % REGISTRE)
    brut = REGISTRE.read_bytes()
    reg = yaml.safe_load(brut.decode("utf-8"))
    info = reg["infographie_detention"]
    total_txt = reg["grandeurs"]["dette_fin_annee_mdeur"]["valeur_imprimee"]
    periode_a = reg["grandeurs"]["dette_fin_annee_mdeur"]["periode"]
    total = float(total_txt.replace(" ", "").replace(NBSP, "").replace(",", "."))
    ss, det = info["sous_secteurs"], info["detenteurs_etat"]
    somme = sum(s["mdeur"] for s in ss)
    if abs(somme - total) > 1.0:
        fail("sous-secteurs %s Md€ contre un total de %s Md€" % (somme, total_txt))
    if abs(sum(d["pct"] for d in det) - 100) > 0.05:
        fail("les parts des détenteurs ne somment pas à 100")
    ff = reg["grandeurs"]["foyers_fiscaux"]
    foyers = float(str(ff["valeur_imprimee"]).replace(",", "."))
    if not 30 < foyers < 50:
        fail("foyers fiscaux hors plage plausible : %s millions" % foyers)
    return {"ss": ss, "det": det, "total": total, "total_txt": total_txt,
            "periode_a": periode_a, "sources": info["sources"],
            "foyers": {"millions": foyers, "periode": ff["periode"],
                       "periode_en": ff.get("periode_en", ""),
                       "source": ff["source_citee"], "url": ff.get("url_source", "")},
            "sha256": hashlib.sha256(brut).hexdigest()}


def lire_insee() -> dict:
    if not XLSX.is_file():
        fail("fichier Insee absent : %s" % XLSX)
    sha = hashlib.sha256(XLSX.read_bytes()).hexdigest()
    if sha != XLSX_SHA256:
        fail("empreinte du fichier Insee changée (%s) : relire la source, puis mettre "
             "à jour XLSX_SHA256" % sha[:12])
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    def ligne(feuille: str, debut: str) -> list:
        for r in wb[feuille].iter_rows(values_only=True):
            if r[0] and str(r[0]).startswith(debut):
                return [x for x in r[1:] if x is not None]
        fail("%s : ligne « %s » introuvable" % (feuille, debut))

    ws = wb["Figure 1c"]
    entetes = [x for x in next(r for r in ws.iter_rows(values_only=True)
                               if r[0] == "Nature des revenus")[1:] if x is not None]
    if entetes != ["D%d" % i for i in range(1, 11)] + ["Ensemble"]:
        fail("Figure 1c : colonnes inattendues %s" % entetes)
    prel = ligne("Figure 1c", "Prélèvements")
    esp = ligne("Figure 1c", "Prestations sociales en espèces")
    nat = ligne("Figure 1c", "Transferts non monétaires")
    solde_uc = ligne("Figure 2a", "Solde de la redistribution publique nationale")[-1]
    solde_md = ligne("Figure 2b", "Solde de la redistribution publique nationale")[-1]
    for nom, l in (("prélèvements", prel), ("prestations", esp), ("transferts", nat)):
        if len(l) != 11:
            fail("Figure 1c : %s sur %d colonnes" % (nom, len(l)))
    return {"prel": prel, "esp": esp, "nat": nat, "solde_uc": solde_uc,
            "solde_md": solde_md, "sha256": sha}


# ------------------------------------------------------------------ F1
def svg_detention(r: dict) -> tuple[str, int]:
    ss, det, total = r["ss"], r["det"], r["total"]
    e = []
    y = 50
    corps = [txt(0, 22, "Qui emprunte ? Qui détient les titres de l'État ? Deux questions, deux champs",
                 13, INK2),
             txt(0, y, "A · Qui emprunte ? Contribution des administrations à la dette publique, "
                 + r["periode_a"], 12, INK, weight="600"),
             txt(0, y + 16, "Valeur nominale, en milliards d'euros · total " + r["total_txt"]
                 + NBSP + "Md€", 11, MUTED)]
    x0, larg = 190, 400
    y = y + 34
    for s in ss:
        w = max(2.0, larg * s["mdeur"] / max(x["mdeur"] for x in ss))
        pct = 100 * s["mdeur"] / total
        corps.append(txt(x0 - 10, y + 11, s["libelle"], 11, INK2, "end"))
        corps.append('<rect x="%d" y="%.1f" width="%.1f" height="14" rx="3" fill="%s"/>'
                     % (x0, y, w, C1))
        corps.append(txt(x0 + w + 8, y + 11, "%s%sMd€ · %s%s%%" % (fr(s["mdeur"]), NBSP,
                                                                   fr(pct), NBSP), 11, INK2))
        y += 22
    y += 16
    corps.append('<line x1="0" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="1"/>'
                 % (y, W, y, AXIS))
    y += 28
    corps.append(txt(0, y, "B · Qui détient les titres négociables de l'État ? 1er trimestre 2026",
                     12, INK, weight="600"))
    corps.append(txt(0, y + 16, "Valeur de marché, en % · classement par résidence du détenteur, "
                     "non par nationalité", 11, MUTED))
    y += 34
    noms = {"Autres (français)": "Autres porteurs français *",
            "Établissements de crédit français": "Banques françaises",
            "Assureurs français": "Assureurs français",
            "OPCVM français": "Fonds (OPCVM) français"}
    for d in det:
        w = max(2.0, larg * d["pct"] / max(x["pct"] for x in det))
        corps.append(txt(x0 - 10, y + 11, noms.get(d["libelle"], d["libelle"]), 11, INK2, "end"))
        corps.append('<rect x="%d" y="%.1f" width="%.1f" height="14" rx="3" fill="%s"/>'
                     % (x0, y, w, C2))
        corps.append(txt(x0 + w + 8, y + 11, "%s%s%%" % (fr(d["pct"], 1), NBSP), 11, INK2))
        y += 22
    corps.append(txt(0, y + 10, "* dont la Banque de France (programmes de l'Eurosystème) : "
                     "part non publiée par la source.", 10, MUTED))
    y += 30
    h = int(y + 13 + 12 * 3 + 8)
    desc = ("Deux panneaux séparés, sans lien de proportion entre eux. A, en valeur nominale : "
            + "; ".join("%s %s Md€" % (s["libelle"], fr(s["mdeur"])) for s in ss)
            + ", sur %s Md€, %s. B, en valeur de marché, porteurs des titres négociables de "
              "l'État au 1er trimestre 2026 : " % (r["total_txt"], r["periode_a"])
            + "; ".join("%s %s %%" % (noms.get(d["libelle"], d["libelle"]).rstrip(" *"),
                                       fr(d["pct"], 1)) for d in det) + ".")
    e += entete(h, "qp-det", "Qui emprunte ? Qui détient les titres de l'État ?", desc)
    e += corps
    e += cartouche(y, [
        ("A : INSEE, Informations rapides n° 79 (27/03/2026), dette de Maastricht par "
         "sous-secteur · B : Banque de France, via l'AFT", INK2),
        ("Deux champs distincts. A : valeur nominale ; B : valeur de marché. "
         "La détention ne mesure pas la charge finale.", INK2),
        ("Compilation Stéphane Lalut, CC BY 4.0 · " + URL_PAGE, MUTED)])
    e.append("</svg>")
    return "\n".join(e) + "\n", h


# ------------------------------------------------------------------ F2
def svg_redistribution(d: dict) -> tuple[str, int]:
    prel = [v / 1000 for v in d["prel"][:10]]
    esp = [v / 1000 for v in d["esp"][:10]]
    nat = [v / 1000 for v in d["nat"][:10]]
    vmax, vmin = 40, -100
    top, bas = 78, 360
    k = (bas - top) / (vmax - vmin)

    def Y(v):
        return top + (vmax - v) * k

    ml, mr = 46, 12
    pas = (W - ml - mr) / 10
    bw = 34
    c = []
    c.append(txt(0, 22, "Prélèvements et transferts publics par dixième de niveau de vie, 2023, "
                 "en milliers d'euros par UC", 13, INK2))
    # Positions FIXES, vues au rendu : une largeur estimée par caractère a fait
    # sortir la troisième étiquette du cadre (21/09).
    for lx, col, lab in ((0, C1, "Prestations en espèces"),
                         (180, C3, "Transferts non monétaires (services publics…)"),
                         (478, C2, "Prélèvements (impôts, cotisations)")):
        c.append('<rect x="%d" y="37" width="12" height="12" rx="2" fill="%s"/>' % (lx, col))
        c.append(txt(lx + 17, 47, lab, 11, INK2))
    for g in range(-100, 41, 20):
        yy = Y(g)
        c.append('<line x1="%d" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="1"/>'
                 % (ml, yy, W - mr, yy, AXIS if g == 0 else GRID))
        c.append(txt(ml - 6, yy + 4, fr(g) if g >= 0 else "−" + fr(-g), 11, MUTED, "end"))
    for i in range(10):
        cx = ml + pas * (i + 0.5)
        x = cx - bw / 2
        y0 = Y(0)
        h1 = esp[i] * k
        h2 = nat[i] * k
        c.append('<rect x="%.1f" y="%.1f" width="%d" height="%.1f" fill="%s"/>'
                 % (x, y0 - h1, bw, h1, C1))
        c.append('<rect x="%.1f" y="%.1f" width="%d" height="%.1f" rx="2" fill="%s"/>'
                 % (x, y0 - h1 - 2 - h2, bw, h2, C3))
        hp = -prel[i] * k
        c.append('<rect x="%.1f" y="%.1f" width="%d" height="%.1f" rx="2" fill="%s"/>'
                 % (x, y0 + 2, bw, hp - 2, C2))
        c.append(txt(cx, y0 - h1 - h2 - 7, fr(esp[i] + nat[i], 1), 10, INK2, "middle"))
        c.append(txt(cx, y0 + hp + 13, "−" + fr(-prel[i], 1), 10, INK2, "middle"))
        c.append(txt(cx, bas + 16, "D%d" % (i + 1), 11, MUTED, "middle"))
    c.append(txt(ml, bas + 30, "← 10 % les plus modestes", 10, MUTED))
    c.append(txt(W - mr, bas + 30, "10 % les plus aisés →", 10, MUTED, "end"))
    y = bas + 44
    h = int(y + 13 + 12 * 3 + 8)
    desc = ("Barres par dixième de niveau de vie, en 2023, en milliers d'euros par unité de "
            "consommation. Au-dessus de zéro, les transferts reçus (prestations en espèces et "
            "transferts non monétaires) : de %s pour le premier dixième à %s pour le dernier. "
            "Sous zéro, les prélèvements : de %s à %s." % (
                fr(esp[0] + nat[0], 1), fr(esp[9] + nat[9], 1), fr(-prel[0], 1), fr(-prel[9], 1)))
    e = entete(h, "qp-red", "Prélèvements et transferts publics par dixième de niveau de vie, 2023",
               desc) + c
    e += cartouche(y, [
        ("Insee, comptes nationaux distribués 2023 (Insee Analyses n° 118, 16/04/2026, figure 1c) "
         "· France, euros par UC", INK2),
        ("Répartition avec conventions d'imputation ; ne mesure pas l'incidence spécifique "
         "de la dette.", INK2),
        ("Compilation Stéphane Lalut, CC BY 4.0 · " + URL_PAGE, MUTED)])
    e.append("</svg>")
    return "\n".join(e) + "\n", h


# ------------------------------------------------------------------ F3
def svg_mecanismes() -> tuple[str, int]:
    c = [txt(0, 22, "Par quels canaux la charge de la dette peut-elle être répartie ?", 13, INK2)]

    def boite(x, y, w, h, lignes, trait=AXIS, tiret=False, fond="#ffffff"):
        d = ' stroke-dasharray="4 3"' if tiret else ""
        out = ['<rect x="%d" y="%d" width="%d" height="%d" rx="6" fill="%s" stroke="%s" '
               'stroke-width="1.5"%s/>' % (x, y, w, h, fond, trait, d)]
        for i, (s, taille, col, gras) in enumerate(lignes):
            out.append(txt(x + 12, y + 20 + i * 16, s, taille, col, weight="600" if gras else None))
        return out

    def fleche(x1, y1, x2, y2):
        return ['<line x1="%d" y1="%d" x2="%d" y2="%d" stroke="%s" stroke-width="1.5" '
                'stroke-dasharray="5 4" marker-end="url(#qp-fl)"/>' % (x1, y1, x2, y2, MUTED)]

    c.append('<defs><marker id="qp-fl" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" '
             'markerHeight="7" orient="auto"><path d="M0,0 L10,5 L0,10 z" fill="%s"/></marker>'
             '</defs>' % MUTED)
    c += boite(1, 150, 200, 74, [("Service de la dette", 12, INK, True),
                                 ("intérêts, et titres qui", 11, INK2, False),
                                 ("arrivent à échéance", 11, INK2, False)], trait=C1)
    canaux = [
        ("Prélèvements", "contribuables : impôts, cotisations"),
        ("Dépenses et prestations", "usagers, bénéficiaires : réduites, gelées, reportées"),
        ("Inflation", "détenteurs de créances et de revenus mal indexés"),
        ("Restructuration (cas extrême)", "porteurs des titres"),
    ]
    y = 44
    for tit, qui in canaux:
        c += boite(300, y, 418, 50, [(tit, 12, INK, True), (qui, 11, INK2, False)])
        c += fleche(200, 187, 296, y + 25)
        y += 64
    c += boite(1, 44, 200, 70, [("Refinancement", 12, INK, True),
                                ("reporte l'échéance ;", 11, INK2, False),
                                ("ne désigne aucun perdant", 11, INK2, False)], tiret=True)
    c += fleche(100, 150, 100, 118)
    c += boite(1, 316, 717, 52, [("En regard : ce que la dette a financé", 12, INK, True),
                                 ("services, prestations, investissements, soutien en crise — "
                                  "bénéfices présents et futurs", 11, INK2, False)],
               tiret=True, fond="#f7f7f4")
    y = 392
    h = int(y + 13 + 12 * 3 + 8)
    desc = ("Schéma sans quantités. Le service de la dette peut être assuré par quatre canaux, qui "
            "se combinent : prélèvements (contribuables), dépenses et prestations (usagers, "
            "bénéficiaires), inflation (détenteurs de créances et de revenus mal indexés), "
            "restructuration, cas extrême (porteurs des titres). Le refinancement reporte "
            "l'échéance sans désigner de perdant. En regard figure ce que la dette a financé.")
    e = entete(h, "qp-mec", "Par quels canaux la charge de la dette peut-elle être répartie ?",
               desc) + c
    e += cartouche(y, [
        ("Synthèse de l'auteur · schéma non quantitatif, flèches d'égale épaisseur", INK2),
        ("Schéma de mécanismes possibles. Aucun poids relatif ni effet causal n'est mesuré ici.",
         INK2),
        ("Stéphane Lalut, CC BY 4.0 · " + URL_PAGE, MUTED)])
    e.append("</svg>")
    return "\n".join(e) + "\n", h


# ------------------------------------------------------------------ texte
def affichage(r: dict, d: dict) -> dict:
    ss = {s["libelle"]: s["mdeur"] for s in r["ss"]}
    det = {x["libelle"]: x["pct"] for x in r["det"]}
    etat_pct = 100 * ss["État"] / r["total"]
    nonres = det["Non-résidents"]
    bafs = det["Établissements de crédit français"] + det["Assureurs français"] + det["OPCVM français"]
    recus = [d["esp"][i] + d["nat"][i] for i in range(10)]
    prel = [-v for v in d["prel"][:10]]
    # Phrases qualitatives de la page : vérifiées ici, sinon arrêt.
    if not nonres > 50:
        fail("la page dit « plus de la moitié » des titres de l'État aux non-résidents : %s %%" % nonres)
    if not etat_pct > 50:
        fail("la page dit que l'État porte l'essentiel de la dette : %s %%" % etat_pct)
    if not (max(recus) / min(recus)) < (max(prel) / min(prel)) / 3:
        fail("la page dit que les transferts reçus varient « beaucoup moins » que les prélèvements")
    if not d["solde_uc"] < 0:
        fail("la page dit que la puissance publique a versé plus qu'elle n'a prélevé")
    return {
        "releve_le": "21 septembre 2026",
        "total_mdeur": r["total_txt"], "periode_a": r["periode_a"],
        "etat_mdeur": fr(ss["État"]), "etat_pct": fr(etat_pct),
        "nonres_pct": fr(nonres, 1), "bafs_pct": fr(bafs, 1),
        "autres_fr_pct": fr(det["Autres (français)"], 1),
        "cd_annee": "2023",
        "d1_prel": fr(prel[0]), "d10_prel": fr(prel[9]),
        "ratio_prel": fr(prel[9] / prel[0]),
        "recu_min": fr(min(recus)), "recu_max": fr(max(recus)),
        "solde_uc": fr(-d["solde_uc"]), "solde_mdeur": fr(-d["solde_md"]),
    }


def main() -> int:
    try:
        import cairosvg
    except ImportError:
        fail("cairosvg absent : SVG et PNG se produisent ensemble ou pas du tout")
    r, d = lire_registre(), lire_insee()
    aff = affichage(r, d)
    figs = [("qui-paie-detention", svg_detention(r)),
            ("qui-paie-redistribution", svg_redistribution(d)),
            ("qui-paie-mecanismes", svg_mecanismes())]
    for nom, (svg, _h) in figs:
        p = IMG / (nom + ".svg")
        p.write_text(svg, encoding="utf-8")
        cairosvg.svg2png(url=str(p), write_to=str(p.with_suffix(".png")), output_width=1440,
                         background_color="white")
    cartes = {"fr": [
        {"id": "detention", "fichier": "qui-paie-detention",
         "titre": "Qui emprunte ? Qui détient les titres de l'État ?",
         "montre": "Les administrations qui empruntent, puis les porteurs des titres de l'État : "
                   "deux champs distincts, sans lien de proportion.",
         "source": "INSEE, IR n° 79 (27/03/2026), %s  ·  Banque de France via l'AFT, "
                   "1er trimestre 2026" % r["periode_a"],
         "precaution": "A : valeur nominale ; B : valeur de marché. La détention ne mesure pas "
                       "la charge finale."},
        {"id": "redistribution", "fichier": "qui-paie-redistribution",
         "titre": "Prélèvements et transferts publics par dixième de niveau de vie, 2023",
         "montre": "Les prélèvements croissent fortement avec le niveau de vie ; les transferts "
                   "reçus, prestations et services publics, varient beaucoup moins.",
         "source": "Insee, comptes nationaux distribués 2023 (Insee Analyses n° 118, figure 1c)",
         "precaution": "Répartition avec conventions d'imputation ; ne mesure pas l'incidence "
                       "spécifique de la dette."},
        {"id": "mecanismes", "fichier": "qui-paie-mecanismes",
         "titre": "Par quels canaux la charge peut-elle être répartie ?",
         "montre": "Quatre canaux possibles, le refinancement qui reporte sans désigner de "
                   "perdant, et en regard ce que la dette a financé.",
         "source": "Synthèse de l'auteur, schéma non quantitatif",
         "precaution": "Aucun poids relatif ni effet causal n'est mesuré ici."},
    ]}
    (REPO / "data" / "figures_qui_paie.json").write_text(
        json.dumps(cartes, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    donnees = {
        "_licence": ("Compilation Stéphane Lalut, CC BY 4.0 (assemblage, grandeurs dérivées, mise "
                     "en cohérence). Données d'origine : INSEE (Licence Ouverte Etalab), "
                     "Banque de France via l'Agence France Trésor, sous leurs propres conditions."),
        "_genere_le": date.today().isoformat(),
        "affichage": aff,
        "detention": {
            "A_sous_secteurs": {"unite": "milliards d'euros, valeur nominale",
                                "periode": r["periode_a"], "total": r["total"],
                                "valeurs": [{"libelle": s["libelle"], "mdeur": s["mdeur"]}
                                            for s in r["ss"]]},
            "B_detenteurs_titres_etat": {"unite": "% des titres négociables de l'État, valeur de "
                                                  "marché", "periode": "1er trimestre 2026",
                                         "valeurs": [{"libelle": x["libelle"], "pct": x["pct"]}
                                                     for x in r["det"]]},
            "sources": r["sources"],
            "note": "Deux champs distincts : ne pas multiplier les parts de B par les montants de A.",
            "registre_sha256": r["sha256"],
        },
        # Denominateur de l'ordre de grandeur « par foyer » de la page du cout
        # (shortcode interets-par-foyer) : meme entree du registre que le livre.
        "echelle_foyer": r["foyers"],
        "redistribution": {
            "unite": "euros par unité de consommation (UC), 2023",
            "colonnes": ["D%d" % i for i in range(1, 11)] + ["Ensemble"],
            "prelevements": d["prel"], "prestations_especes": d["esp"],
            "transferts_non_monetaires": d["nat"],
            "solde_finance_par_endettement_uc": d["solde_uc"],
            "solde_finance_par_endettement_mdeur": d["solde_md"],
            "convention": ("Insee : le supplément financé par endettement est imputé par "
                           "convention pour moitié à de moindres prélèvements, pour moitié à des "
                           "transferts supplémentaires."),
            "source": {"publication": "Insee Analyses n° 118, 16/04/2026, figure 1c et 2a-2b",
                       "url": XLSX_URL, "sha256_fichier": d["sha256"]},
        },
    }
    corps = json.dumps(donnees, ensure_ascii=False, indent=1) + "\n"
    (REPO / "data" / "qui_paie_donnees.json").write_text(corps, encoding="utf-8")
    (REPO / "static" / "qui_paie_donnees.json").write_text(corps, encoding="utf-8")
    for nom, (_s, h) in figs:
        print("OK  %s.svg (720 x %d) + PNG 1440 px" % (nom, h))
    print("OK  data/figures_qui_paie.json, data/ et static/qui_paie_donnees.json")
    print("    registre %s…  ·  Insee %s…" % (r["sha256"][:12], d["sha256"][:12]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
