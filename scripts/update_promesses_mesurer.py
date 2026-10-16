#!/usr/bin/env python3
"""Volet « Mesurer » du dossier Promesses : deux mesures du chômage, deux populations, et des règles qui changent.

Protocole : D:/PRO/06_PROMOTION/RECHERCHE_DOSSIER_PROMESSES_2026-10-05/FICHES_PREUVE.md, volet 5, E1 et E2 (v2,
gelée au commit dfd0ec9 ; arbitrage PRO-20261005-164029). Refonte de phase B (12/10/2026) : arbitrage
06_PROMOTION/.../REFONTE_PILOTE/MESURER/ARBITRAGE_CE_MESURER_PHASE_A.md.

E1 — restitution de l'ESTIMATION Insee (Insee Références « Emploi, chômage, revenus du travail », éd. 2026, dossier 1,
figure 2, moyenne 2024) : trois ensembles disjoints tirés des EFFECTIFS, jamais de parts arrondies. Énoncé autorisé :
« deux définitions, des populations qui se recoupent partiellement ».
E2 — séries publiées par la Dares (catégorie A et « A hors RSA et hors jeunes en parcours », France hors Mayotte,
CVS-CJO, moyenne trimestrielle des inscrits en fin de mois), sans aucun écart calculé comme effet.
E3 (phase B) — série longue publiée par l'Insee (même dossier, figure 1 : moyennes annuelles 2012-2025 des deux
mesures), ruptures de règle annotées, aucune fenêtre de mandat.

RELEVÉ FIGÉ (phase B, 12/10/2026) : le générateur ne lit QUE les fichiers archivés de scripts/sources_mesurer/, dont
l'empreinte est dans SHA256SUMS (garde M0). Aucune donnée publiée après le relevé n'entre sans le geste explicite
`--relever` (Insee xlsx, CSV et métadonnées Dares relus en direct, archive et SHA256SUMS réécrits) ; `--check` n'écrit
jamais rien, pas même l'archive (défaut D16 du diagnostic : l'ancienne fonction `charger` réécrivait l'archive sous
--check).

Gardes (toutes ARRÊTENT, rien n'est écrit) :
  M0  relevé figé : chaque fichier lu a l'empreinte inscrite dans SHA256SUMS ;
  M1  champ et année imprimés dans le tableur = ceux du protocole (France hors Mayotte, 15-64 ans, logement ordinaire, 2024) ;
  M2  additivité des effectifs de la figure 2 (lignes et colonnes) à l'arrondi publié près (± 1 millier par cellule) ;
  M3  « se recoupent partiellement » : intersection et deux ensembles exclusifs strictement positifs au-delà de la
      précision de restitution (> 2 milliers, soit deux arrondis) ;
  M4  CONTRÔLE DE DÉFINITION : les trois ensembles restitués depuis la figure 2 égalent ceux que l'Insee imprime dans
      un autre tableau, la figure 8 (lignes (1), (2), (3)) ; écart nul exigé ;
  M5  cohérence algébrique (déclarée comme telle, jamais appelée témoin) : les parts publiées 49 % et 62 % se
      retrouvent depuis les effectifs, à l'arrondi ;
  M6  séries Dares : présence des deux séries, champ « France », CVS-CJO ; la série « hors RSA » commence en 2018-T1 ;
  M7  citations : chaque passage cité (définitions, Insee, Dares, ASP, Unédic, Dreets) est retrouvé mot pour mot dans
      son fichier archivé ; les nombres cités (105 000 en 2022, 36 000 en 2024, 36 000 en 2025) sont LUS dans ces
      passages ; les nombres de chaque traduction égalent ceux de l'original ;
  M8  série longue (phase B) : les écarts entre les deux mesures calculés depuis la feuille « Figure 1 » égalent ceux
      que l'Insee imprime dans le texte du dossier (2012, 2019, 2020, 2021, 2024) et les variations 2025 ;
  M9  phrase « la série hors RSA porte aussi la rupture de 2025 » (phase B, défaut D3) : les DEUX séries montent au
      1er trimestre 2025 et reculent au 2e, et restent au-dessus de leur niveau de fin 2024 (« de plus ») ; sinon la
      phrase n'est plus soutenue.
Le TÉMOIN d'extraction indépendant (lecture du PDF par une chaîne écrite séparément) est exécuté hors de ce script :
D:/PRO/06_PROMOTION/RECHERCHE_DOSSIER_PROMESSES_2026-10-05/test_decisif/mesurer_e1/.
Autotest de mutation à chaque exécution, chaque mutation devant s'arrêter sur SA garde : octet altéré -> M0 ; commun
altéré -> M4 ; parts permutées -> M5 ; champ altéré -> M1 ; citation altérée -> M7 ; ligne de la série longue
décalée -> M8 ; saut du 1er trimestre 2025 effacé de la série hors RSA -> M9 ; version téléphone faussée -> parité.

Écrit : data/promesses_mesurer.json (+ static/), static/promesses_mesurer.csv (UTF-8 BOM, format long),
data/figures_mesurer.json, static/img/mesurer-{recoupement,regle,serie-longue}{,-en}{,-m}.svg/.png.
Rien n'est écrit à sorties identiques (toutes comparées, fiches « Réutiliser » comprises).
"""
from __future__ import annotations

import copy
import csv
import hashlib
import html
import io
import json
import re
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import langue_generee as lg  # noqa: E402  (scripts/langue_generee.py)
import lecture_tableur as lt  # noqa: E402  -- forme des cellules (organe commun)
import pied_figure  # noqa: E402  (scripts/pied_figure.py)

SRCDIR = ROOT / "scripts" / "sources_mesurer"
SUMS = SRCDIR / "SHA256SUMS"
URL_XLSX = "https://www.insee.fr/fr/statistiques/fichier/8733119/ECRT2026-D1.xlsx"
URL_DARES = ("https://data.dares.travail-emploi.gouv.fr/api/explore/v2.1/catalog/datasets/"
             "dares_defm_stock_france_cvs_trim/exports/csv?use_labels=true")
URL_DARES_META = "https://data.dares.travail-emploi.gouv.fr/api/explore/v2.1/catalog/datasets/dares_defm_stock_france_cvs_trim"
EN_DIRECT = {"ECRT2026-D1.xlsx": URL_XLSX, "dares_defm_stock_france_cvs_trim.csv": URL_DARES,
             "dares_defm_stock_france_cvs_trim_meta.json": URL_DARES_META}
PAGE_URL = "stephane-lalut.com/comment-savoir-si-une-promesse-est-tenue/"
PAGE_URL_EN = "stephane-lalut.com/en/how-can-we-tell-whether-a-promise-was-kept/"
OUT_DATA = ROOT / "data" / "promesses_mesurer.json"
OUT_STATIC = ROOT / "static" / "promesses_mesurer.json"
OUT_CSV = ROOT / "static" / "promesses_mesurer.csv"
OUT_FIGURES = ROOT / "data" / "figures_mesurer.json"
OUT_IMG = ROOT / "static" / "img"
CHAMP_ATTENDU = "France hors Mayotte, personnes âgées de 15 à 64 ans, vivant dans un logement ordinaire."
STATUTS = ["En emploi", "Chômeur", "Inactif dans le halo", "Inactif hors halo"]
CATS = ["Catégorie A", "Catégorie B", "Catégorie C", "Catégorie D", "Catégorie E", "Non inscrits"]


class Arret(Exception):
    pass


def log(m: str) -> None:
    print(m.encode("ascii", "replace").decode("ascii"))


def fr(n: float, dec: int = 0) -> str:
    return ("{:,.%df}" % dec).format(n).replace(",", "\u202f").replace(".", ",")


def nb(s: str) -> str:
    return s.replace("\u00a0", " ").strip()


# ------------------------------------------------------------------ relevé figé (M0)
def lire_sommes() -> tuple[dict, str]:
    if not SUMS.is_file():
        raise Arret("M0 : %s absent, aucun releve fige" % SUMS.name)
    sommes, releve = {}, ""
    for l in SUMS.read_text(encoding="utf-8").splitlines():
        m = re.match(r"# releve: (\d{4}-\d\d-\d\d)", l)
        if m:
            releve = m.group(1)
        elif l and not l.startswith("#"):
            h, nom = l.split(None, 1)
            sommes[nom.strip()] = h
    if not releve:
        raise Arret("M0 : date de releve absente de SHA256SUMS")
    return sommes, releve


def verifier_releve(octets: dict, sommes: dict) -> str:
    for nom, b in octets.items():
        if nom not in sommes:
            raise Arret("M0 : %s lu mais absent de SHA256SUMS" % nom)
        if hashlib.sha256(b).hexdigest() != sommes[nom]:
            raise Arret("M0 : empreinte de %s differente du releve fige" % nom)
    return "M0 : %d fichiers lus, empreintes conformes au releve fige" % len(octets)


def relever() -> int:
    """Geste explicite : relit les trois sources en direct, réécrit l'archive et SHA256SUMS. Jamais sous --check."""
    import datetime
    for nom, url in EN_DIRECT.items():
        req = urllib.request.Request(url, headers={"User-Agent": "stephane-lalut.com generateur (contact via site)"})
        with urllib.request.urlopen(req, timeout=90) as r:
            (SRCDIR / nom).write_bytes(r.read())
    sommes, _ = lire_sommes()
    tete = [l for l in SUMS.read_text(encoding="utf-8").splitlines() if l.startswith("#") and not l.startswith("# releve:")]
    lignes = tete + ["# releve: %s (sources en direct relues par --relever)" % datetime.date.today().isoformat()]
    for nom in sommes:
        lignes.append("%s  %s" % (hashlib.sha256((SRCDIR / nom).read_bytes()).hexdigest(), nom))
    SUMS.write_text("\n".join(lignes) + "\n", encoding="utf-8")
    log("--relever : archive et SHA256SUMS reecrits ; relancer le generateur, puis relire la page (donnees nouvelles).")
    return 0


# ------------------------------------------------------------------ lecture
def lire_insee(b: bytes) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
    f2 = [list(r) for r in wb["Figure 2"].iter_rows(values_only=True) if any(c is not None for c in r)]
    titre = f2[0][0]
    entete = lt.entetes(f2[1])[1:8]
    if entete != CATS + ["Ensemble"]:
        raise Arret("figure 2 : en-tete inattendue %s" % entete)
    eff, i = {}, 3
    for k in range(5):
        r = f2[i + k]
        eff[r[0]] = dict(zip(CATS + ["Ensemble"], r[1:8]))
    textes = [r[0] for r in f2 if r and isinstance(r[0], str)]
    champ = next((nb(t[len("Champ :"):]) for t in textes if nb(t).startswith("Champ :")), "")
    source = next((nb(t) for t in textes if t.startswith("Source")), "")
    parts_cat = {}
    j = next(k for k, r in enumerate(f2) if r[0] and str(r[0]).startswith("Part dans la catégorie"))
    for k in range(1, 6):
        r = f2[j + k]
        parts_cat[r[0]] = dict(zip(CATS + ["Ensemble"], r[1:8]))
    jj = next(k for k, r in enumerate(f2) if r[0] and str(r[0]).startswith("Part dans le statut"))
    parts_stat = {}
    for k in range(1, 6):
        r = f2[jj + k]
        parts_stat[r[0]] = dict(zip(CATS + ["Ensemble"], r[1:8]))
    f8 = [[c for c in r if c is not None] for r in wb["Figure 8"].iter_rows(values_only=True)]
    f8d = {}
    for r in f8:
        if r and isinstance(r[0], str) and len(r) >= 3 and lt.valeur(r[1]) is not None:
            if lt.valeur(r[2]) is None:
                raise Arret("figure 8, %s : valeur 2021 manquante ou illisible (%r)" % (nb(r[0]), r[2]))
            f8d[nb(r[0])] = {"2024": lt.valeur(r[1]), "2021": lt.valeur(r[2])}
    # Figure 1 (phase B) : moyennes annuelles 2012-2025 des deux mesures, en millions
    f1 = [[c for c in r if c is not None] for r in wb["Figure 1"].iter_rows(values_only=True)]
    annees = next(r[1:] for r in f1 if r and r[0] == "Catégorie")
    lg_ = {}
    for r in f1:
        if r and r[0] in ("Chômeurs au sens du BIT", "Inscrits en catégorie A à France Travail"):
            lg_[r[0]] = dict(zip([str(a) for a in annees], r[1:]))
    champ1 = next((nb(r[0]) for r in f1 if r and isinstance(r[0], str) and nb(r[0]).startswith("Champ")), "")
    return {"titre": nb(titre), "eff": eff, "champ": champ, "source": source, "parts_cat": parts_cat,
            "parts_stat": parts_stat, "f8": f8d,
            "f1": {"bit": lg_.get("Chômeurs au sens du BIT", {}), "a": lg_.get("Inscrits en catégorie A à France Travail", {})},
            "f1_champ": champ1}


def lire_dares(b: bytes) -> dict:
    rows = list(csv.DictReader(io.StringIO(b.decode("utf-8-sig")), delimiter=";"))
    out = {}
    for cat, cle in (("A", "a"), ("A hors RSA et hors jeunes en parcours", "a_hors")):
        s = {r["Date"]: int(r["Nombre de demandeurs d'emploi"]) for r in rows
             if r["Champ"] == "France" and r["Catégorie"] == cat and r["Sexe"] == "Total" and r["Tranche d'âge"] == "Total"
             and r["Ancienneté"] == "Total" and r["Type de données"] == "cvs-cjo" and r["Nombre de demandeurs d'emploi"]}
        out[cle] = dict(sorted(s.items()))
    return out


def texte_source(nom: str, b: bytes) -> str:
    """Texte d'une source archivée, espaces normalisés (insécables compris) ; métadonnées Dares : description sans balises."""
    t = b.decode("utf-8")
    if nom.endswith("_meta.json"):
        t = html.unescape(re.sub(r"<[^>]+>", " ", json.loads(t)["metas"]["default"]["description"]))
    return re.sub(r"\s+", " ", t.replace("\u00a0", " ").replace("\u202f", " "))


# ------------------------------------------------------------------ calcul et gardes
def calculer(I: dict, D: dict) -> tuple[dict, list[str]]:
    g = []
    if "2024" not in I["titre"] or I["champ"] != CHAMP_ATTENDU:
        raise Arret("M1 : annee ou champ de la figure 2 differents du protocole (%r / %r)" % (I["titre"][-30:], I["champ"]))
    g.append("M1 : figure 2, moyenne 2024, champ « %s »" % I["champ"])
    E = I["eff"]
    for s in STATUTS + ["Ensemble"]:
        if abs(sum(E[s][c] for c in CATS) - E[s]["Ensemble"]) > len(CATS):
            raise Arret("M2 : ligne %s non additive" % s)
    for c in CATS + ["Ensemble"]:
        if abs(sum(E[s][c] for s in STATUTS) - E["Ensemble"][c]) > len(STATUTS):
            raise Arret("M2 : colonne %s non additive" % c)
    g.append("M2 : effectifs additifs a l'arrondi publie pres")
    commun = E["Chômeur"]["Catégorie A"]
    bit_seul = E["Chômeur"]["Ensemble"] - commun
    a_seul = E["Ensemble"]["Catégorie A"] - commun
    if not (commun > 2 and bit_seul > 2 and a_seul > 2):
        raise Arret("M3 : « se recoupent partiellement » faux (commun %s, BIT seul %s, A seul %s)" % (commun, bit_seul, a_seul))
    g.append("M3 : recoupement partiel (commun %d, BIT hors A %d, A hors BIT %d milliers)" % (commun, bit_seul, a_seul))
    f8 = I["f8"]
    k1 = next(k for k in f8 if k.startswith("Inscrits en catégorie A et non chômeurs"))
    k2 = next(k for k in f8 if k.startswith("Inscrits en catégorie A et chômeurs"))
    k3 = next(k for k in f8 if k.startswith("Chômeurs et non inscrits en catégorie A"))
    if (f8[k1]["2024"], f8[k2]["2024"], f8[k3]["2024"]) != (a_seul, commun, bit_seul):
        raise Arret("M4 : restitution figure 2 (%d, %d, %d) != figure 8 (%s, %s, %s)"
                    % (a_seul, commun, bit_seul, f8[k1]["2024"], f8[k2]["2024"], f8[k3]["2024"]))
    g.append("M4 : les trois ensembles restitues egalent la figure 8 imprimee")
    p_a = 100 * commun / E["Ensemble"]["Catégorie A"]
    p_b = 100 * commun / E["Chômeur"]["Ensemble"]
    if round(p_a) != I["parts_cat"]["Chômeur"]["Catégorie A"] or round(p_b) != I["parts_stat"]["Chômeur"]["Catégorie A"]:
        raise Arret("M5 : parts publiees non retrouvees (%.1f/%s ; %.1f/%s)"
                    % (p_a, I["parts_cat"]["Chômeur"]["Catégorie A"], p_b, I["parts_stat"]["Chômeur"]["Catégorie A"]))
    g.append("M5 (coherence algebrique, non temoin) : 49 et 62 % retrouves")
    if not D["a"] or not D["a_hors"] or min(D["a_hors"]) != "2018-T1":
        raise Arret("M6 : series Dares absentes ou debut inattendu (%s)" % (min(D["a_hors"]) if D["a_hors"] else None))
    g.append("M6 : series Dares A (%s-%s) et A hors RSA (%s-%s)" % (min(D["a"]), max(D["a"]), min(D["a_hors"]), max(D["a_hors"])))
    # M9 : la phrase « la série hors RSA porte aussi la rupture de 2025 » (D3)
    sauts = {}
    for cle in ("a", "a_hors"):
        s = D[cle]
        if not all(k in s for k in ("2024-T4", "2025-T1", "2025-T2")):
            raise Arret("M9 : trimestres 2024-T4 a 2025-T2 absents de la serie %s" % cle)
        t1, t2 = s["2025-T1"] - s["2024-T4"], s["2025-T2"] - s["2025-T1"]
        if not (t1 > 0 and t2 < 0):
            raise Arret("M9 : la serie %s ne monte plus au 1er trimestre 2025 puis ne recule plus au 2e (%+d, %+d)" % (cle, t1, t2))
        if not s[max(s)] > s["2024-T4"]:  # la page écrit « de plus » depuis la fin de 2024 (jetons *_var, sans signe)
            raise Arret("M9 : la serie %s n'est plus au-dessus de son niveau du 4e trimestre 2024" % cle)
        sauts[cle] = (t1, t2)
    g.append("M9 : les deux series montent au 1er trimestre 2025 et reculent au 2e (A %+d / %+d ; hors RSA %+d / %+d)"
             % (sauts["a"] + sauts["a_hors"]))
    r = {"commun": commun, "bit_seul": bit_seul, "a_seul": a_seul, "bit_total": E["Chômeur"]["Ensemble"],
         "a_total": E["Ensemble"]["Catégorie A"], "p_a": p_a, "p_b": p_b,
         "bit_non_inscrits": E["Chômeur"]["Non inscrits"],
         "bit_bc": E["Chômeur"]["Catégorie B"] + E["Chômeur"]["Catégorie C"],
         "bit_de": E["Chômeur"]["Catégorie D"] + E["Chômeur"]["Catégorie E"],
         "a_halo": E["Inactif dans le halo"]["Catégorie A"], "a_hors_halo": E["Inactif hors halo"]["Catégorie A"],
         "a_emploi": E["En emploi"]["Catégorie A"],
         "f8_2021": {"commun": f8[k2]["2021"], "a_seul": f8[k1]["2021"], "bit_seul": f8[k3]["2021"]},
         "dares_a": D["a"], "dares_a_hors": D["a_hors"], "eff": E, "sauts_2025": sauts}
    return r, g


def serie_longue(I: dict, texte_d1: str) -> tuple[dict, str]:
    """M8 : écarts de la feuille « Figure 1 » contre ceux que l'Insee imprime dans le texte du même dossier."""
    bit, a = I["f1"]["bit"], I["f1"]["a"]
    ans = [str(y) for y in range(2012, 2026)]
    if [k for k in ans if k not in bit or k not in a]:
        raise Arret("M8 : annees 2012-2025 incompletes dans la feuille Figure 1")
    if not I["f1_champ"].startswith("Champ : France hors Mayotte, personnes âgées de 15 à 64 ans"):
        raise Arret("M8 : champ de la figure 1 inattendu (%r)" % I["f1_champ"])
    ec = {k: round(a[k] - bit[k], 1) for k in ans}
    num = lambda s: float(s.replace(",", "."))  # noqa: E731
    pats = [(r"L’écart, qui est de (\d,\d) million en 2024, était de (\d,\d) million en 2012", ("2024", "2012")),
            (r"augmenté jusqu’à (\d,\d) million en 2019 et même (\d,\d) million en 2020", ("2019", "2020")),
            (r"En 2021, il atteignait (\d,\d) million", ("2021",))]
    lu = {}
    for p, cles in pats:
        m = re.search(p, texte_d1)
        if not m:
            raise Arret("M8 : passage du dossier introuvable (%s)" % p[:40])
        for c, v in zip(cles, m.groups()):
            lu[c] = num(v)
    for c, v in lu.items():
        if abs(ec[c] - v) > 1e-9:
            raise Arret("M8 : ecart %s = %.1f dans la feuille Figure 1, %.1f dans le texte de l'Insee" % (c, ec[c], v))
    m = re.search(r"chômeurs au sens du BIT augmente de (\d,\d) million par rapport à 2024 et le nombre d’inscrits à "
                  r"France Travail en catégorie A de près de (\d,\d) million", texte_d1)
    if not m:
        raise Arret("M8 : passage de l'encadre 1 (2025) introuvable")
    d_bit, d_a = round(bit["2025"] - bit["2024"], 1), round(a["2025"] - a["2024"], 1)
    if (d_bit, d_a) != (num(m.group(1)), num(m.group(2))):
        raise Arret("M8 : variations 2025 (%.1f, %.1f) differentes du texte (%s, %s)" % (d_bit, d_a, m.group(1), m.group(2)))
    return ({"bit": {k: bit[k] for k in ans}, "a": {k: a[k] for k in ans}, "ecart": ec},
            "M8 : serie longue 2012-2025 ; ecarts 2012, 2019, 2020, 2021, 2024 et variations 2025 egaux au texte de l'Insee")


# ------------------------------------------------------------------ citations (M7)
# Chaque passage cité est retrouvé mot pour mot dans son fichier archivé ; les clés « cit_* » et les définitions sont
# émises par jeton (prose FR et EN), jamais recopiées dans la page (défaut D1 : définition tronquée hors jeton).
D1 = "insee_ecrt2026_d1_texte.txt"
VE = "insee_ve2026_encadre5_extrait.txt"
IR = "insee_ir192_2026-08-07_extrait.txt"
UN = "unedic_comprendre_bit_2026-09-04_extrait.txt"
META = "dares_defm_stock_france_cvs_trim_meta.json"
ASP = "asp_delibere_2024-11-14_texte.txt"
DR = "dreets_normandie_2024t4_extrait.txt"
CIT = {
    "chomeur_bit": (D1, "une personne âgée de 15 ans ou plus sans emploi, disponible pour en occuper un dans les quinze jours et qui a activement cherché un emploi dans le mois précédent (ou en a trouvé un qui commence dans moins de trois mois)"),
    "categorie_a": (D1, "les personnes sans activité au cours du mois et tenues de rechercher un emploi"),
    "cit_dares_ssm": (D1, "service statistique ministériel en charge du marché du travail"),
    "cit_alea": (D1, "l’aléa associé à des données d’enquête"),
    "cit_insee_titre": (D1, "l’écart s’est réduit entre 2021 et 2024 mais demeure élevé"),
    "cit_divergent": (D1, "les deux indicateurs divergent de nouveau"),
    "cit_2019": (D1, "des mesures susceptibles d’affecter le comportement d’inscription à France Travail indépendamment de la situation sur le marché du travail"),
    "cit_2021": (D1, "des pondérations rétropolées ont été utilisées pour corriger au mieux l’impact de ces changements"),
    "cit_2022": (META, "un impact pérenne à la baisse sur le nombre d'inscrits en catégorie A"),
    "cit_2024": (DR, "entre octobre et novembre 2024, une évolution dans le formulaire d’actualisation a conduit à augmenter le nombre de demandeurs d’emploi en catégorie A"),
    "cit_ir192": (IR, "les deux publics cibles de la loi pour le plein emploi contribuent pour près de la moitié de la hausse du taux de chômage depuis fin 2024 (+0,44 point sur +0,98 point au total)"),
    "cit_ir192_prudence": (IR, "Ces contributions doivent toutefois être considérées avec prudence, du fait des aléas d’enquête et de la faible part des populations concernées"),
    "cit_ir192_reserve": (IR, "Il s’agit par ailleurs de contributions comptables, qui ne permettent pas de distinguer les effets de la loi pour le plein emploi des effets conjoncturels affectant ces publics"),
    "cit_unedic": (UN, "cette loi expliquerait près de la moitié de la hausse du taux de chômage depuis fin 2024"),
    "cit_ve_titre": (VE, "En 2025, la loi pour le plein emploi a des effets relativement limités sur les indicateurs du marché du travail au sens du BIT"),
    "cit_ve_contrib": (VE, "une faible contribution haussière des bénéficiaires du RSA sur les taux d’activité et de chômage"),
    "cit_ve_recherche": (VE, "Les bénéficiaires du RSA recherchent plus souvent un emploi"),
    "cit_ve_limites": (VE, "l’aléa d’échantillonnage conduisant à une certaine variabilité et la difficulté à distinguer précisément les publics concernés"),
    "cit_asp": (ASP, "ce qui rendra difficilement interprétables les séries statistiques de demandeurs d’emploi jusqu’ici labellisées"),
    "cit_asp_2021": (ASP, "avait été renouvelée et étendue le 21 mai 2021 pour une période de cinq ans"),
    "cit_asp_suspend": (ASP, "la labellisation de ces séries pour la période restant à courir entre le 1er janvier 2025 et le 20 mai 2026"),
    "cit_asp_attente": (ASP, "dans l’attente d’être à même d’envisager un processus de labellisation des nouvelles séries, lorsque le régime pérenne sera atteint"),
    "cit_dares_label": (META, "a suspendu, à partir du 1er janvier 2025, la labellisation des séries statistiques sur les demandeurs d'emploi produites par France Travail et la Dares"),
}
TR = {
    "chomeur_bit": "a person aged 15 or over without a job, available to take one within fifteen days and who has actively looked for a job in the previous month (or has found one starting within three months)",
    "categorie_a": "people who did no work during the month and are required to look for work",
    "cit_dares_ssm": "the ministerial statistical service responsible for the labour market",
    "cit_alea": "the random error associated with survey data",
    "cit_insee_titre": "the gap narrowed between 2021 and 2024 but remains large",
    "cit_divergent": "the two indicators diverge again",
    "cit_2019": "measures likely to affect registration behaviour with France Travail regardless of the situation on the labour market",
    "cit_2021": "back-cast weights were used to correct as far as possible for the impact of these changes",
    "cit_2022": "a lasting downward impact on the number of registrants in category A",
    "cit_2024": "between October and November 2024, a change in the monthly updating form led to an increase in the number of jobseekers in category A",
    "cit_ir192": "the two target groups of the full employment act account for nearly half of the rise in the unemployment rate since end-2024 (%s out of %s in total)"
                 % (lg.points(0.44, 2, "en", True), lg.points(0.98, 2, "en", True)),
    "cit_ir192_prudence": "these contributions should however be treated with caution, because of survey error and the small share of the populations concerned",
    "cit_ir192_reserve": "these are moreover accounting contributions, which do not make it possible to distinguish the effects of the full employment act from the cyclical effects affecting these groups",
    "cit_unedic": "this law would explain nearly half of the rise in the unemployment rate since end-2024",
    "cit_ve_titre": "In 2025, the full employment act has relatively limited effects on labour market indicators in the ILO sense",
    "cit_ve_contrib": "a small upward contribution of RSA recipients to the activity and unemployment rates",
    "cit_ve_recherche": "RSA recipients look for work more often",
    "cit_ve_limites": "sampling error, which causes some variability, and the difficulty of identifying the groups concerned precisely",
    "cit_asp": "which will make the previously labelled statistical series on jobseekers difficult to interpret",
    "cit_asp_2021": "had been renewed and extended on 21 May 2021 for a period of five years",
    "cit_asp_suspend": "the label of these series for the period remaining between 1 January 2025 and 20 May 2026",
    "cit_asp_attente": "until it is in a position to consider a labelling process for the new series, once the permanent regime has been reached",
    "cit_dares_label": "suspended, from 1 January 2025, the label of the statistical series on jobseekers produced by France Travail and the Dares",
}
# Passages qui fondent un libellé ou un champ sans être affichés (libellé de la figure 2, Mayotte).
PREUVES = {
    "fin_de_mois": (META, "nombre de personnes inscrites en fin de mois"),
    "moyenne_trimestre": (DR, "s'établit en moyenne sur le trimestre"),
    "dares_hors_mayotte": (META, "à l’exception de Mayotte"),
    "ir192_avec_mayotte": ("insee_ir192_2026-08-07_avertissement_extrait.txt",
                           "le champ des principaux indicateurs conjoncturels du marché du travail porte sur le champ « France », étendant ainsi le champ précédent « France hors Mayotte »"),
}
# Nombres cités LUS dans les sources (catégorie « lue par le script » de la phrase de méthode).
NOMBRES = {
    "rupture_2022": (META, r"diminution cumulée de (\d+ \d{3}) inscrits en catégorie A entre janvier et novembre 2022"),
    "formulaire_2024": (DR, r"augmentation d’environ (\d+ \d{3}) demandeurs d’emploi en catégorie A"),
    "incident_2025": (META, r"(\d+ \d{3}) personnes avaient été comptabilisées à tort en catégorie A"),
}


def _nums(s: str, lang: str) -> list[float]:
    if lang == "fr":
        return sorted(float(x.replace(" ", "").replace("\u00a0", "").replace(",", "."))
                      for x in re.findall(r"\d+(?:[ \u00a0]\d{3})*(?:,\d+)?", s))
    return sorted(float(x.replace(",", "")) for x in re.findall(r"\d+(?:,\d{3})*(?:\.\d+)?", s))


def verifier_citations(textes: dict) -> tuple[str, dict, dict]:
    out, nombres = {}, {}
    for k, (f, c) in list(CIT.items()) + list(PREUVES.items()):
        if c not in textes[f]:
            raise Arret("M7 : citation %s introuvable dans %s" % (k, f))
        out[k] = c
    for k, (f, p) in NOMBRES.items():
        m = re.search(p, textes[f])
        if not m:
            raise Arret("M7 : nombre %s introuvable dans %s" % (k, f))
        nombres[k] = int(m.group(1).replace(" ", ""))
    for k, tr in TR.items():
        if _nums(tr, "en") != _nums(out[k], "fr"):
            raise Arret("M7 : nombres de la traduction %s differents de l'original (%s / %s)"
                        % (k, _nums(tr, "en"), _nums(out[k], "fr")))
    if set(TR) != set(CIT):
        raise Arret("M7 : traductions et citations ne couvrent pas les memes cles")
    return ("M7 : %d citations et %d preuves retrouvees mot pour mot ; %d nombres lus dans les sources ; traductions controlees"
            % (len(CIT), len(PREUVES), len(NOMBRES))), {k: out[k] for k in CIT}, nombres


# ------------------------------------------------------------------ affichage (jetons mes-val et {mes.*})
def _variations(r: dict, lang: str) -> dict:
    da, dh = r["dares_a"], r["dares_a_hors"]
    fin = max(da)
    (a1, a2), (h1, h2) = r["sauts_2025"]["a"], r["sauts_2025"]["a_hors"]
    return {"dares_a_var": lg.sans_signe(da[fin] - da["2024-T4"], 0, lang),
            "dares_a_hors_var": lg.sans_signe(dh[max(dh)] - dh["2024-T4"], 0, lang),
            "dares_a_t1_2025": lg.signe(a1, 0, lang), "dares_a_t2_2025": lg.signe(a2, 0, lang),
            "dares_hors_t1_2025": lg.signe(h1, 0, lang), "dares_hors_t2_2025": lg.signe(h2, 0, lang)}


def affichage_en(r: dict, L: dict, N: dict) -> dict:
    def m(v):  # thousands -> "1.43 million" or "863,000"
        return ("%.2f million" % (v / 1000)) if v >= 1000 else "{:,.0f}".format(v * 1000)
    da, dh = r["dares_a"], r["dares_a_hors"]
    t_fin = max(da)
    e = lambda v, d=0: ("{:,.%df}" % d).format(v)  # noqa: E731
    out = {
        "commun": m(r["commun"]), "bit_seul": m(r["bit_seul"]), "a_seul": m(r["a_seul"]),
        "commun_env": "about " + m(r["commun"]), "a_seul_env": "about " + m(r["a_seul"]),
        "bit_total": m(r["bit_total"]), "a_total": m(r["a_total"]),
        "p_a": e(r["p_a"]), "p_b": e(r["p_b"]),
        "bit_non_inscrits": m(r["bit_non_inscrits"]), "bit_bc": m(r["bit_bc"]), "bit_de": m(r["bit_de"]),
        "a_halo": m(r["a_halo"]), "a_hors_halo": m(r["a_hors_halo"]), "a_emploi": m(r["a_emploi"]),
        "a_sans_emploi": m(r["a_halo"] + r["a_hors_halo"]),
        "dares_fin": "%s quarter of %s" % ({"1": "first", "2": "second", "3": "third", "4": "fourth"}[t_fin[-1]], t_fin[:4]),
        "dares_a_fin": m(da[t_fin] / 1000), "dares_a_hors_fin": m(dh[max(dh)] / 1000),
        "dares_a_2024t4": m(da["2024-T4"] / 1000), "dares_a_hors_2024t4": m(dh["2024-T4"] / 1000),
        "rupture_2022": e(N["rupture_2022"]), "formulaire_2024": e(N["formulaire_2024"]), "incident_2025": e(N["incident_2025"]),
    }
    for k in ("2012", "2019", "2020", "2021", "2024", "2025"):
        out["ecart_" + k] = "%.1f million" % L["ecart"][k]
    out.update(_variations(r, "en"))
    return out


def affichage(r: dict, L: dict, N: dict) -> dict:
    def m(v):  # milliers -> « 1,43 million » ou « 863 000 »
        return (fr(v / 1000, 2) + " million" + ("s" if v >= 2000 else "")) if v >= 1000 else fr(v * 1000, 0)
    da, dh = r["dares_a"], r["dares_a_hors"]
    t_fin = max(da)
    out = {
        "commun": m(r["commun"]), "bit_seul": m(r["bit_seul"]), "a_seul": m(r["a_seul"]),
        "commun_env": "environ " + m(r["commun"]), "a_seul_env": "environ " + m(r["a_seul"]),
        "bit_total": m(r["bit_total"]), "a_total": m(r["a_total"]),
        "p_a": fr(r["p_a"], 0), "p_b": fr(r["p_b"], 0),
        "bit_non_inscrits": m(r["bit_non_inscrits"]), "bit_bc": m(r["bit_bc"]), "bit_de": m(r["bit_de"]),
        "a_halo": m(r["a_halo"]), "a_hors_halo": m(r["a_hors_halo"]), "a_emploi": m(r["a_emploi"]),
        "a_sans_emploi": m(r["a_halo"] + r["a_hors_halo"]),
        "dares_fin": "%s trimestre %s" % ({"1": "1er", "2": "2e", "3": "3e", "4": "4e"}[t_fin[-1]], t_fin[:4]),
        "dares_a_fin": m(da[t_fin] / 1000), "dares_a_hors_fin": m(dh[max(dh)] / 1000),
        "dares_a_2024t4": m(da["2024-T4"] / 1000), "dares_a_hors_2024t4": m(dh["2024-T4"] / 1000),
        "rupture_2022": fr(N["rupture_2022"]), "formulaire_2024": fr(N["formulaire_2024"]), "incident_2025": fr(N["incident_2025"]),
    }
    for k in ("2012", "2019", "2020", "2021", "2024", "2025"):
        out["ecart_" + k] = fr(L["ecart"][k], 1) + " million"
    out.update(_variations(r, "fr"))
    return out


# ------------------------------------------------------------------ figures (charte des dossiers du site)
W = 720
BLEU, GRIS, GRIS_CLAIR = "#184f95", "#8a8f98", "#c9cdd3"
INK, INK2, MUTED, GRID = "#0A0A0E", "#52514e", "#898781", "#e1e0d9"
FONT = "system-ui, -apple-system, Segoe UI, sans-serif"
LICENCES = {"fr": "Compilation Stéphane Lalut, CC BY 4.0 · " + PAGE_URL, "en": "Compiled by Stéphane Lalut, CC BY 4.0 · " + PAGE_URL_EN}
LANG = "fr"
FS_MIN = pied_figure.POLICE_M


def nbl(v, dec: int = 0) -> str:  # nombre selon la langue de la figure
    return ("{:,.%df}" % dec).format(v) if LANG == "en" else fr(v, dec)


def esc(s) -> str:
    return html.escape(str(s), quote=True)


def txt(x, y, s, size=11, fill=INK2, anchor="start", weight=None) -> str:
    w = ' font-weight="%s"' % weight if weight else ""
    return '<text x="%.1f" y="%.1f" font-size="%s" fill="%s" text-anchor="%s"%s>%s</text>' % (x, y, size, fill, anchor, w, esc(s))


def cadre(H: int, ident: str, titre: str, sous: str, desc: str) -> list[str]:
    return ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 %d %d" role="img" font-variant-numeric="tabular-nums" '
            'aria-labelledby="%s-t %s-d" font-family="%s">' % (W, H + 48, ident, ident, FONT),
            '<title id="%s-t">%s</title><desc id="%s-d">%s</desc>' % (ident, esc(titre), ident, esc(desc)),
            '<rect width="%d" height="%d" fill="#ffffff"/>' % (W, H + 48),
            txt(0, 18, titre, 13, INK, weight="600"), txt(0, 36, sous, 12, INK2)]


LICENCES_M = {"fr": "Compilation Stéphane Lalut, CC BY 4.0 · stephane-lalut.com",
              "en": "Compiled by Stéphane Lalut, CC BY 4.0 · stephane-lalut.com"}


def blocs_pied(src: str, note: str, mobile: bool = False) -> list:
    """Version téléphone : licence courte (l'adresse complète de la page ne tient pas en 300 de large)."""
    return [(src, INK2, "pied-source"), (note, INK2, "pied-note"), ((LICENCES_M if mobile else LICENCES)[LANG], MUTED, "pied-licence")]


T = {
    "fr": {"src1": "Insee Références, Emploi, chômage, revenus du travail, éd. 2026, dossier 1, figures 2 et 8 (appariement Dares-Insee-France Travail)",
           "note1": "Estimation de l'Insee, moyenne 2024, France hors Mayotte, 15-64 ans, logement ordinaire ; deux définitions distinctes ; marge d'enquête non publiée.",
           "src2": "Dares et France Travail, moyenne trimestrielle des inscrits en fin de mois, France hors Mayotte, données CVS-CJO (série « hors RSA » depuis 2018)",
           "note2": "Champ et période différents de la figure du croisement 2024. Séries non labellisées depuis le 01/01/2025 (label de 2021 suspendu par l'ASP, qui courait jusqu'au 20/05/2026).",
           "src3": "Insee Références, Emploi, chômage, revenus du travail, éd. 2026, dossier 1, figure 1 (Insee, enquête Emploi, séries longues ; Dares-France Travail, STMT)",
           "note3": "Moyennes annuelles publiées, en millions, France hors Mayotte, 15-64 ans. Repères : Insee (1, 2) et Dares (3, 4).",
           "t1": "Chômeurs au sens du BIT et inscrits en catégorie A, 2024",
           "s1": "deux définitions, des populations qui se recoupent en partie (en milliers de personnes)",
           "d1": "Barre en trois segments, en milliers de personnes, moyenne 2024 (estimation Insee) : {bit_seul} chômeurs au sens du BIT "
                 "non inscrits en catégorie A, {commun} dans les deux à la fois, {a_seul} inscrits en catégorie A non chômeurs au sens "
                 "du BIT. Chômeurs au sens du BIT : {bit_total}, dont {p_b} % inscrits en catégorie A. Inscrits en catégorie A : "
                 "{a_total}, dont {p_a} % chômeurs au sens du BIT.",
           "bit": "Chômeurs au sens du BIT : %s, dont %s %% inscrits en A",
           "insc": "Inscrits en catégorie A : %s, dont %s %% chômeurs au sens du BIT",
           "dont1": "dont : non inscrits %s ; inscrits en B ou C %s ;", "dont2": "inscrits en D ou E %s",
           "dont3": "dont : halo autour du chômage %s ;", "dont4": "inactifs hors halo %s ; en emploi %s", "deux": "dans les deux",
           "m_bit": "Chômeurs BIT non inscrits en A", "m_deux": "Dans les deux à la fois", "m_a": "Inscrits en A non chômeurs BIT",
           "m_tot_bit": "Chômeurs BIT : %s, dont %s %% inscrits en A", "m_tot_a": "Inscrits en A : %s, dont %s %% chômeurs BIT",
           "m_unite": "En milliers de personnes, moyenne 2024",
           "t2": "Inscrits en catégorie A : série publiée et série hors bénéficiaires du RSA et jeunes en parcours",
           "s2": "France hors Mayotte, moyenne trimestrielle des inscrits en fin de mois, CVS-CJO ; même échelle pour les deux panneaux",
           "d2": "Deux panneaux de 2018 à {an}, sur la même échelle : à gauche la catégorie A publiée, à droite la série hors "
                 "bénéficiaires du RSA et hors jeunes en parcours. Une bande marque 2022 (nouveau formulaire d'actualisation), une "
                 "ligne le 1er janvier 2025 (loi pour le plein emploi). Au {fin}, {va} pour la catégorie A publiée et {vh} hors RSA.",
           "p1": "Catégorie A publiée", "p2": "Hors bénéficiaires du RSA et jeunes en parcours", "p2m": "Hors RSA et jeunes en parcours",
           "r2022": "2022 : formulaire", "r2025": "2025 : loi",
           "ecart": "Les deux séries sautent au 1er trimestre 2025 et reculent au 2e (actualisation automatique, janvier-mars 2025) ; leur écart ne mesure pas l'effet de la réforme.",
           "t3": "Chômeurs au sens du BIT et inscrits en catégorie A, 2012-2025",
           "s3": "moyennes annuelles publiées par l'Insee, en millions de personnes",
           "d3": "Deux courbes annuelles de 2012 à 2025, en millions. Chômeurs au sens du BIT : {bit}. Inscrits en catégorie A : {a}. "
                 "Repères : 1, réforme de l'assurance chômage (novembre 2019 à décembre 2021) ; 2, enquête Emploi rénovée en 2021 ; "
                 "3, nouveau formulaire d'actualisation en 2022 ; 4, formulaire d'octobre-novembre 2024 et loi pour le plein emploi en 2025.",
           "l_bit": "Chômeurs au sens du BIT", "l_a": "Inscrits en catégorie A",
           "rep": ["1 Assurance chômage, 2019-2021 (inscription)", "2 Enquête Emploi rénovée, 2021 (BIT)",
                   "3 Formulaire d'actualisation, 2022 (A)", "4 Formulaire fin 2024, loi 2025 (A)"],
           "M": lambda v: fr(v / 1e6, 1) + " M", "Mi": lambda v: fr(v, 1)},
    "en": {"src1": "Insee Références, Emploi, chômage, revenus du travail, 2026 ed., dossier 1 (in-depth study), figures 2 and 8 (Dares-Insee-France Travail matching)",
           "note1": "Insee estimate, 2024 average, France excluding Mayotte, aged 15-64, ordinary households; two distinct definitions; survey margin not published.",
           "src2": "Dares and France Travail, quarterly average of end-of-month registrants, France excluding Mayotte, seasonally adjusted (series 'excluding RSA' since 2018)",
           "note2": "Scope and period differ from the 2024 cross-tabulation figure. Series not labelled since 1 Jan. 2025 (2021 label suspended by the ASP, which ran until 20 May 2026).",
           "src3": "Insee Références, Emploi, chômage, revenus du travail, 2026 ed., dossier 1, figure 1 (Insee, Labour Force Survey, long series; Dares-France Travail, STMT)",
           "note3": "Published annual averages, in millions, France excluding Mayotte, aged 15-64. Markers: Insee (1, 2) and Dares (3, 4).",
           "t1": "ILO-defined unemployment and category A registration in France, 2024",
           "s1": "two definitions, populations that partly overlap (thousands of people)",
           "d1": "Bar in three segments, in thousands of people, 2024 average (Insee estimate): {bit_seul} people unemployed under the "
                 "ILO definition not registered in category A, {commun} in both at once, {a_seul} category A registrants who are not "
                 "unemployed under the ILO definition. Unemployed (ILO): {bit_total}, of whom {p_b}% registered in category A. Category "
                 "A registrants: {a_total}, of whom {p_a}% unemployed under the ILO definition.",
           "bit": "Unemployed (ILO definition): %s, of whom %s%% registered in A",
           "insc": "Registered in category A: %s, of whom %s%% unemployed (ILO)",
           "dont1": "of whom: not registered %s; registered in B or C %s;", "dont2": "registered in D or E %s",
           "dont3": "of whom: halo around unemployment %s;", "dont4": "inactive outside the halo %s; in work %s", "deux": "in both",
           "m_bit": "ILO unemployed, not in A", "m_deux": "In both at once", "m_a": "In category A, not ILO unemployed",
           "m_tot_bit": "ILO unemployed: %s, %s%% in A", "m_tot_a": "In category A: %s, %s%% ILO unemployed",
           "m_unite": "Thousands of people, 2024 average",
           "t2": "Category A registrants: published series and series without RSA recipients or youth programmes",
           "s2": "France excluding Mayotte, quarterly average of end-of-month registrants, seasonally adjusted; same scale for both panels",
           "d2": "Two panels from 2018 to {an}, on the same scale: left, published category A; right, the series excluding RSA "
                 "recipients and young people in programmes. A band marks 2022 (new monthly updating form), a line 1 January 2025 "
                 "(full employment act). In the {fin}, {va} for published category A and {vh} excluding RSA.",
           "p1": "Published category A", "p2": "Excluding RSA recipients and young people in programmes", "p2m": "Excluding RSA and youth programmes",
           "r2022": "2022: form", "r2025": "2025: law",
           "ecart": "Both series jump in the first quarter of 2025 and fall back in the second (automatic updating, Jan.-Mar. 2025); their gap does not measure the effect of the reform.",
           "t3": "ILO unemployment and category A registrants in France, 2012-2025",
           "s3": "annual averages published by Insee, in millions of people",
           "d3": "Two annual lines from 2012 to 2025, in millions. Unemployed under the ILO definition: {bit}. Category A registrants: "
                 "{a}. Markers: 1, unemployment insurance reform (November 2019 to December 2021); 2, Labour Force Survey redesigned "
                 "in 2021; 3, new monthly updating form in 2022; 4, form of October-November 2024 and full employment act in 2025.",
           "l_bit": "Unemployed (ILO)", "l_a": "Category A registrants",
           "rep": ["1 Unemployment insurance, 2019-2021 (registration)", "2 Labour Force Survey redesigned, 2021 (ILO)",
                   "3 Monthly updating form, 2022 (A)", "4 Form in late 2024, law in 2025 (A)"],
           "M": lambda v: "%.1fm" % (v / 1e6), "Mi": lambda v: "%.1f" % v},
}


def _fmt_fig1(r: dict) -> dict:
    return {"bit_seul": nbl(r["bit_seul"]), "commun": nbl(r["commun"]), "a_seul": nbl(r["a_seul"]),
            "bit_total": nbl(r["bit_total"]), "a_total": nbl(r["a_total"]), "p_a": nbl(r["p_a"]), "p_b": nbl(r["p_b"])}


def fig_recoupement(r: dict) -> str:
    H = 270
    x0, x1 = 0, W
    tot = r["bit_seul"] + r["commun"] + r["a_seul"]
    X = lambda v: x0 + v / tot * (x1 - x0)  # noqa: E731
    y, h = 118, 46
    L = T[LANG]
    F = _fmt_fig1(r)
    e = cadre(H, "mr", L["t1"], L["s1"], L["d1"].format(**F))
    a, b = X(r["bit_seul"]), X(r["bit_seul"] + r["commun"])
    e.append('<rect x="%.1f" y="%d" width="%.1f" height="%d" fill="%s"/>' % (x0, y, a - x0, h, GRIS_CLAIR))
    e.append('<rect x="%.1f" y="%d" width="%.1f" height="%d" fill="%s"/>' % (a, y, b - a, h, BLEU))
    e.append('<rect x="%.1f" y="%d" width="%.1f" height="%d" fill="%s"/>' % (b, y, x1 - b, h, GRIS_CLAIR))
    e.append(txt((x0 + a) / 2, y + 28, F["bit_seul"], 15, INK, "middle", "600"))
    e.append(txt((a + b) / 2, y + 28, F["commun"], 15, "#ffffff", "middle", "600"))
    e.append(txt((b + x1) / 2, y + 28, F["a_seul"], 15, INK, "middle", "600"))
    e.append('<path d="M%.1f %d v-8 H%.1f v8" fill="none" stroke="%s" stroke-width="1.5"/>' % (x0 + 1, y - 6, b, INK2))
    e.append(txt((x0 + b) / 2, y - 20, L["bit"] % (F["bit_total"], F["p_b"]), 12, INK, "middle", "600"))
    e.append('<path d="M%.1f %d v8 H%.1f v-8" fill="none" stroke="%s" stroke-width="1.5"/>' % (a, y + h + 6, x1 - 1, INK2))
    e.append(txt((a + x1) / 2, y + h + 30, L["insc"] % (F["a_total"], F["p_a"]), 12, INK, "middle", "600"))
    yy = y + h + 62
    e.append(txt(x0, yy, L["dont1"] % (nbl(r["bit_non_inscrits"]), nbl(r["bit_bc"])), 10, INK2))
    e.append(txt(x0, yy + 13, L["dont2"] % nbl(r["bit_de"]), 10, INK2))
    e.append(txt(x1, yy, L["dont3"] % nbl(r["a_halo"]), 10, INK2, "end"))
    e.append(txt(x1, yy + 13, L["dont4"] % (nbl(r["a_hors_halo"]), nbl(r["a_emploi"])), 10, INK2, "end"))
    e.append(txt((a + b) / 2, yy, L["deux"], 10, BLEU, "middle", "600"))
    return pied_figure.pied(e, H + 2, W, blocs_pied(L["src1"], L["note1"]), GRID)


def fig_recoupement_m(r: dict) -> str:
    L, F, P = T[LANG], _fmt_fig1(r), pied_figure.POLICE_M
    e, y = pied_figure.tete_mobile("mrm", L["t1"], L["d1"].format(**F), INK, FONT)
    y += 8
    e.append(txt(0, y, L["m_unite"], P, MUTED))
    y += 18
    vmax = max(r["bit_seul"], r["commun"], r["a_seul"])
    for lib, v, col in ((L["m_bit"], r["bit_seul"], GRIS_CLAIR), (L["m_deux"], r["commun"], BLEU), (L["m_a"], r["a_seul"], GRIS_CLAIR)):
        y += 16
        e.append(txt(0, y, lib, P, INK2))
        y += 8
        wbar = 220 * v / vmax
        e.append('<rect x="0" y="%d" width="%.1f" height="22" fill="%s"/>' % (y, wbar, col))
        e.append(txt(wbar + 6, y + 16, nbl(v), 15, INK, "start", "600"))
        y += 30
    for gab, tot, part in ((L["m_tot_bit"], F["bit_total"], F["p_b"]), (L["m_tot_a"], F["a_total"], F["p_a"])):
        tete, _, queue = (gab % (tot, part)).partition(", ")
        y += 22
        e.append(txt(0, y, tete + ",", P, INK, "start", "600"))
        y += 18
        e.append(txt(0, y, queue, P, INK2))
    return pied_figure.pied_mobile(e, y + 14, blocs_pied(L["src1"], L["note1"], True), GRID)


def _echelle(r: dict) -> tuple[list, int, int]:
    da, dh = r["dares_a"], r["dares_a_hors"]
    cles = list(dh)  # 2018-T1 -> fin
    tout = [v for s in (da, dh) for k, v in s.items() if k in cles]
    return cles, int(min(tout) // 200_000) * 200_000, int(-(-max(tout) // 200_000)) * 200_000


def _fmt_fig2(r: dict) -> dict:
    L = T[LANG]
    da, dh = r["dares_a"], r["dares_a_hors"]
    fin = max(dh)
    lib_fin = ("%s trimestre %s" % ({"1": "1er", "2": "2e", "3": "3e", "4": "4e"}[fin[-1]], fin[:4]) if LANG == "fr"
               else "%s quarter of %s" % ({"1": "first", "2": "second", "3": "third", "4": "fourth"}[fin[-1]], fin[:4]))
    return {"an": fin[:4], "fin": lib_fin, "va": L["M"](da[fin]), "vh": L["M"](dh[fin])}


def _panneau(e: list, s: dict, cles: list, vmin: int, vmax: int, x0: float, pw: float, y0: float, y1: float,
             col: str, nom: str, fs: float, axe: bool, an_pas: int, mobile: bool = False) -> None:
    L = T[LANG]
    Y = lambda v: y1 - (v - vmin) / (vmax - vmin) * (y1 - y0)  # noqa: E731
    X = lambda i: x0 + i / (len(cles) - 1) * pw  # noqa: E731
    e.append(txt(0 if mobile else x0, y0 - 26, nom, fs, col if col == BLEU else INK2, "start", "600"))
    i22a, i22b = cles.index("2022-T1"), cles.index("2022-T4")
    e.append('<rect x="%.1f" y="%.1f" width="%.1f" height="%.1f" fill="%s" opacity="0.6"/>'
             % (X(i22a), y0 - 8, X(i22b) - X(i22a), y1 - y0 + 8, "#eeede8"))
    for v in range(vmin, vmax + 1, 400_000):
        e.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s"/>' % (x0, Y(v), x0 + pw, Y(v), GRID))
        if axe:
            e.append(txt(x0 - 6, Y(v) + 4, L["M"](v), fs - 1, MUTED, "end"))
    for i, c in enumerate(cles):
        if c.endswith("T1") and int(c[:4]) % an_pas == 0:
            e.append(txt(X(i), y1 + 15, c[:4], fs - 1, MUTED, "middle"))
    i25 = cles.index("2025-T1") - 0.5
    e.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s" stroke-dasharray="3 3"/>' % (X(i25), y0 - 8, X(i25), y1, INK2))
    pts = " ".join("%.1f,%.1f" % (X(i), Y(s[c])) for i, c in enumerate(cles))
    e.append('<polyline points="%s" fill="none" stroke="%s" stroke-width="2.4"%s/>'
             % (pts, col, "" if col == BLEU else ' stroke-dasharray="5 3"'))
    # téléphone : l'année seule (les deux libellés se chevauchaient à 300 de large)
    e.append(txt((X(i22a) + X(i22b)) / 2, y0 - 10, "2022" if mobile else L["r2022"], fs - 2, INK2, "middle"))
    e.append(txt(X(i25) + 3, y0 - 10, "2025" if mobile else L["r2025"], fs - 2, INK2, "start"))


def fig_regle(r: dict) -> str:
    """Deux PANNEAUX sur la même échelle (contre-expertise PRO-20261005-183942, P6) : deux courbes superposées autour
    d'une date se lisent comme une quasi-expérience ; deux panneaux séparés disent deux populations. Phase B : bande
    2022 (formulaire d'actualisation) et note sur le saut du 1er trimestre 2025, commun aux deux panneaux (D2, D3)."""
    H = 360
    cles, vmin, vmax = _echelle(r)
    L = T[LANG]
    e = cadre(H, "mg", L["t2"], L["s2"], L["d2"].format(**_fmt_fig2(r)))
    pw, gap, x0g = 300, 50, 56
    for k, (s, col, nom) in enumerate(((r["dares_a"], BLEU, L["p1"]), (r["dares_a_hors"], GRIS, L["p2"]))):
        _panneau(e, s, cles, vmin, vmax, x0g + k * (pw + gap), pw, 96, 300, col, nom, 11, k == 0, 2)
    e.append(txt(0, 336, L["ecart"], 10, INK2))
    return pied_figure.pied(e, H + 2, W, blocs_pied(L["src2"], L["note2"]), GRID)


def fig_regle_m(r: dict) -> str:
    L, P = T[LANG], pied_figure.POLICE_M
    cles, vmin, vmax = _echelle(r)
    F = _fmt_fig2(r)
    e, y = pied_figure.tete_mobile("mgm", L["t2"], L["d2"].format(**F), INK, FONT)
    for s, col, nom, val in ((r["dares_a"], BLEU, L["p1"], F["va"]), (r["dares_a_hors"], GRIS, L["p2m"], F["vh"])):
        y0 = y + 60
        _panneau(e, s, cles, vmin, vmax, 52, 236, y0, y0 + 150, col, nom, P + 2, True, 4, mobile=True)
        e.append(txt(298, y0 + 150 + 34, "%s : %s" % (F["fin"], val) if LANG == "fr" else "%s: %s" % (F["fin"], val), P, INK, "end", "600"))
        y = y0 + 150 + 40
    for l in pied_figure.coupe(L["ecart"], 38):
        y += 18
        e.append(txt(0, y, l, P, INK2))
    return pied_figure.pied_mobile(e, y + 12, blocs_pied(L["src2"], L["note2"], True), GRID)


REPERES = [("2019", "2021", 1), ("2021", "2021", 2), ("2022", "2022", 3), ("2024", "2025", 4)]


def _fmt_fig3(Lg: dict) -> dict:
    Mi = T[LANG]["Mi"]
    sep = "; " if LANG == "en" else " ; "
    f = lambda s: sep.join("%s %s" % (k, Mi(v)) for k, v in s.items())  # noqa: E731
    return {"bit": f(Lg["bit"]), "a": f(Lg["a"])}


def _courbes(e: list, Lg: dict, x0: float, x1: float, y0: float, y1: float, fs: float, mobile: bool) -> None:
    L = T[LANG]
    ans = list(Lg["bit"])
    vmin, vmax = 2.0, 4.0
    X = lambda i: x0 + i / (len(ans) - 1) * (x1 - x0)  # noqa: E731
    Y = lambda v: y1 - (v - vmin) / (vmax - vmin) * (y1 - y0)  # noqa: E731
    for v in (2.0, 2.5, 3.0, 3.5, 4.0):
        e.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s"/>' % (x0, Y(v), x1, Y(v), GRID))
        e.append(txt(x0 - 6, Y(v) + 4, L["Mi"](v), fs - 1, MUTED, "end"))
    for i, a in enumerate(ans):
        if not mobile or int(a) % 4 == 0:
            e.append(txt(X(i), y1 + (22 if mobile else 16), a, fs - 1, MUTED, "middle"))
    for a0, a1, n in REPERES:
        xa, xb = X(ans.index(a0)), X(ans.index(a1))
        if xb > xa:
            e.append('<rect x="%.1f" y="%.1f" width="%.1f" height="%.1f" fill="#eeede8" opacity="0.7"/>' % (xa, y0 - 4, xb - xa, y1 - y0 + 4))
        e.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s" stroke-dasharray="2 3"/>' % (xa, y0 - 4, xa, y1, MUTED))
        cy = y0 - 16 - (22 if mobile and n % 2 == 0 else 0)  # téléphone : repères voisins sur deux hauteurs
        e.append('<circle cx="%.1f" cy="%.1f" r="%d" fill="#ffffff" stroke="%s"/>' % ((xa + xb) / 2, cy, 10 if mobile else 8, INK2))
        e.append(txt((xa + xb) / 2, cy + 4, str(n), fs - 2, INK2, "middle"))
    for cle, col, dash in (("a", BLEU, ""), ("bit", INK2, ' stroke-dasharray="5 3"')):
        s = Lg[cle]
        pts = " ".join("%.1f,%.1f" % (X(i), Y(s[a])) for i, a in enumerate(ans))
        e.append('<polyline points="%s" fill="none" stroke="%s" stroke-width="2.4"%s/>' % (pts, col, dash))
        if not mobile:
            e.append(txt(X(0) - 4, Y(s[ans[0]]) - 8, L["Mi"](s[ans[0]]), fs, col, "start", "600"))
        e.append(txt(X(len(ans) - 1), Y(s[ans[-1]]) - 8, L["Mi"](s[ans[-1]]), fs, col, "end", "600"))


def fig_serie_longue(Lg: dict) -> str:
    """Troisième figure (arbitrage de phase A : série longue Insee, ruptures annotées, AUCUNE fenêtre de mandat)."""
    H = 330
    L = T[LANG]
    e = cadre(H, "ml", L["t3"], L["s3"], L["d3"].format(**_fmt_fig3(Lg)))
    e.append('<line x1="0" y1="56" x2="18" y2="56" stroke="%s" stroke-width="2.4"/>' % BLEU)
    e.append(txt(24, 60, L["l_a"], 11, BLEU, "start", "600"))
    e.append('<line x1="220" y1="56" x2="238" y2="56" stroke="%s" stroke-width="2.4" stroke-dasharray="5 3"/>' % INK2)
    e.append(txt(244, 60, L["l_bit"], 11, INK2, "start", "600"))
    _courbes(e, Lg, 40, 700, 100, 260, 11, False)
    for k, rep in enumerate(L["rep"]):
        e.append(txt((k % 2) * 360, 296 + 15 * (k // 2), rep, 10, INK2))
    return pied_figure.pied(e, H + 2, W, blocs_pied(L["src3"], L["note3"]), GRID)


def fig_serie_longue_m(Lg: dict) -> str:
    L, P = T[LANG], pied_figure.POLICE_M
    e, y = pied_figure.tete_mobile("mlm", L["t3"], L["d3"].format(**_fmt_fig3(Lg)), INK, FONT)
    y += 10
    e.append('<line x1="0" y1="%d" x2="16" y2="%d" stroke="%s" stroke-width="2.4"/>' % (y, y, BLEU))
    e.append(txt(22, y + 5, L["l_a"], P, BLEU))
    y += 20
    e.append('<line x1="0" y1="%d" x2="16" y2="%d" stroke="%s" stroke-width="2.4" stroke-dasharray="5 3"/>' % (y, y, INK2))
    e.append(txt(22, y + 5, L["l_bit"], P, INK2))
    y0 = y + 70
    _courbes(e, Lg, 34, 292, y0, y0 + 170, P + 2, True)
    y = y0 + 170 + 26
    for rep in L["rep"]:
        for l in pied_figure.coupe(rep, 38):
            y += 18
            e.append(txt(0, y, l, P, INK2))
    return pied_figure.pied_mobile(e, y + 12, blocs_pied(L["src3"], L["note3"], True), GRID)


MONTRE = {
    "fr": {"recoupement": "En 2024, selon l'estimation de l'Insee, environ {commun} de personnes sont à la fois chômeurs au sens du BIT "
                          "et inscrites en catégorie A ; environ {a_seul} d'inscrits en catégorie A ne sont pas chômeurs au sens du BIT, "
                          "et {bit_seul} chômeurs au sens du BIT ne sont pas inscrits en catégorie A.",
           "regle": "Les deux séries d'inscrits en catégorie A portent les changements de règles : formulaire d'actualisation en "
                    "2022, inscription d'office et actualisation automatique en 2025. Elles décrivent deux populations, et leur écart "
                    "ne mesure pas l'effet de la réforme.",
           "serie-longue": "De 2012 à 2025, l'écart entre inscrits en catégorie A et chômeurs au sens du BIT passe de {ecart_2012} en "
                           "2012 à {ecart_2019} en 2019, {ecart_2021} en 2021 et {ecart_2024} en 2024, selon les moyennes annuelles publiées "
                           "par l'Insee ; les règles d'inscription et l'enquête ont changé dans l'intervalle."},
    "en": {"recoupement": "In 2024, according to the Insee estimate, about {commun} people are both unemployed in the ILO sense and "
                          "registered in category A; about {a_seul} category A registrants are not unemployed under the ILO definition, and {bit_seul} "
                          "people unemployed under the ILO definition are not registered in category A.",
           "regle": "Both series of category A registrants carry the rule changes: monthly updating form in 2022, automatic "
                    "registration and automatic updating in 2025. They describe two populations, and the gap between them does not "
                    "measure the effect of the reform.",
           "serie-longue": "From 2012 to 2025, the gap between category A registrants and the ILO unemployed goes from {ecart_2012} in "
                           "2012 to {ecart_2019} in 2019, {ecart_2021} in 2021 and {ecart_2024} in 2024, according to the annual averages "
                           "published by Insee; registration rules and the survey changed in the meantime."},
}
FIGS = [("recoupement", "mesurer-recoupement"), ("regle", "mesurer-regle"), ("serie-longue", "mesurer-serie-longue")]


def fiches(figs: dict, A: dict, A_en: dict) -> dict:
    res = {}
    for lang, suf, AA in (("fr", "", A), ("en", "-en", A_en)):
        out = []
        for ident, fichier in FIGS:
            svg = figs[fichier + suf + ".svg"]
            titre = html.unescape(re.search(r"<title[^>]*>(.*?)</title>", svg).group(1))
            src, note = pied_figure.lire(svg, "pied-source"), pied_figure.lire(svg, "pied-note")
            if not src or not note or pied_figure.lire(svg, "pied-licence") != LICENCES[lang]:
                raise Arret("fiche %s%s : pied illisible" % (fichier, suf))
            out.append(dict(id=ident, fichier=fichier + suf, titre=titre, montre=MONTRE[lang][ident].format(**AA),
                            source=src, precaution=note))
        res[lang] = out
    return res


def parite(figs: dict) -> list[str]:
    out = []
    for _, fichier in FIGS:
        for suf in ("", "-en"):
            ec = pied_figure.parite(figs[fichier + suf + ".svg"], figs[fichier + suf + "-m.svg"], FS_MIN, (BLEU,))
            if ec:
                raise Arret("parite telephone %s%s : %s" % (fichier, suf, "; ".join(ec)))
            out.append(fichier + suf)
    return out


def csv_texte(r: dict, Lg: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["serie", "periode", "statut_bit", "categorie_inscription", "valeur", "unite", "champ", "source"])
    for s in STATUTS + ["Ensemble"]:
        for c in CATS + ["Ensemble"]:
            w.writerow(["croisement_bit_inscription", "2024", s, c, r["eff"][s][c], "milliers de personnes",
                        CHAMP_ATTENDU, "Insee Références ECRT 2026, dossier 1, figure 2 (estimation)"])
    for k, v in r["dares_a"].items():
        w.writerow(["inscrits_categorie_A", k, "", "A", v, "personnes, moyenne trimestrielle des inscrits en fin de mois",
                    "France hors Mayotte, CVS-CJO", "Dares-France Travail"])
    for k, v in r["dares_a_hors"].items():
        w.writerow(["inscrits_categorie_A_hors_RSA_et_jeunes_en_parcours", k, "", "A hors RSA et jeunes en parcours", v,
                    "personnes, moyenne trimestrielle des inscrits en fin de mois", "France hors Mayotte, CVS-CJO", "Dares-France Travail"])
    for cle, nom, cat in (("bit", "serie_longue_chomeurs_bit", ""), ("a", "serie_longue_inscrits_categorie_A", "A")):
        for k, v in Lg[cle].items():
            w.writerow([nom, k, "Chômeur" if cle == "bit" else "", cat, v, "millions de personnes, moyenne annuelle",
                        "France hors Mayotte, 15-64 ans", "Insee Références ECRT 2026, dossier 1, figure 1"])
    return buf.getvalue()


# ------------------------------------------------------------------ autotest : chaque mutation doit s'arrêter sur SA garde
def autotest(I: dict, D: dict, octets: dict, sommes: dict, textes: dict) -> list[str]:
    out = []

    def attendre(nom, pref, f):
        try:
            f()
        except Arret as e:
            if not str(e).startswith(pref):
                raise Arret("mutation %s : arretee par une autre garde (%s), controle %s ABSENT" % (nom, str(e)[:50], pref))
            out.append("%s -> %s" % (nom, str(e)[:70]))
            return
        raise Arret("mutation %s : aucune garde n'a mordu, controle %s ABSENT" % (nom, pref))

    o = dict(octets)
    o["dares_defm_stock_france_cvs_trim.csv"] = o["dares_defm_stock_france_cvs_trim.csv"][:-1] + b"9"
    attendre("octet altere", "M0", lambda: verifier_releve(o, sommes))
    for nom, mut, pref in (
            # +10 sur le commun ET sur les totaux de sa ligne et de sa colonne : l'additivité (M2) reste vraie,
            # seule la confrontation à la figure 8 (M4) peut voir l'erreur
            ("commun altere", lambda m: [m["eff"][s].__setitem__(c, m["eff"][s][c] + 10) for s, c in
                                         (("Chômeur", "Catégorie A"), ("Chômeur", "Ensemble"), ("Ensemble", "Catégorie A"),
                                          ("Ensemble", "Ensemble"))], "M4"),
            ("parts permutees", lambda m: (m["parts_cat"]["Chômeur"].__setitem__("Catégorie A", 62),
                                           m["parts_stat"]["Chômeur"].__setitem__("Catégorie A", 49)), "M5"),
            ("champ altere", lambda m: m.__setitem__("champ", m["champ"].replace("15 à 64", "15 ans ou plus, 15 à 74")), "M1")):
        m = copy.deepcopy(I)
        mut(m)
        attendre(nom, pref, lambda m=m: calculer(m, D))
    d = copy.deepcopy(D)
    d["a_hors"]["2025-T1"] = d["a_hors"]["2024-T4"]
    attendre("saut 2025-T1 efface de la serie hors RSA", "M9", lambda: calculer(I, d))
    m = copy.deepcopy(I)
    ans = sorted(m["f1"]["a"])
    m["f1"]["a"] = dict(zip(ans, [m["f1"]["a"][ans[0]]] + [m["f1"]["a"][k] for k in ans[:-1]]))
    attendre("serie longue decalee d'un an", "M8", lambda: serie_longue(m, textes[D1]))
    sauve = CIT["cit_ve_titre"]
    CIT["cit_ve_titre"] = (sauve[0], sauve[1].replace("relativement limités", "limités"))
    try:
        attendre("citation Insee alteree", "M7", lambda: verifier_citations(textes))
    finally:
        CIT["cit_ve_titre"] = sauve
    return out


def main() -> int:
    try:
        if "--relever" in sys.argv[1:]:
            if "--check" in sys.argv[1:]:
                raise Arret("--relever et --check sont incompatibles : --check n'ecrit rien")
            return relever()
        return _main()
    except Arret as e:
        log("ARRET : " + str(e))
        log("Aucun fichier ecrit.")
        return 1


def _main() -> int:
    check = "--check" in sys.argv[1:]
    sommes, releve = lire_sommes()
    octets = {nom: (SRCDIR / nom).read_bytes() for nom in sommes}
    g = [verifier_releve(octets, sommes)]
    textes = {nom: texte_source(nom, b) for nom, b in octets.items() if nom.endswith((".txt", ".json"))}
    bm = octets["dares_defm_stock_france_cvs_trim_meta.json"]
    maj_dares = json.loads(bm.decode("utf-8"))["metas"]["default"]["modified"][:10]
    I, D = lire_insee(octets["ECRT2026-D1.xlsx"]), lire_dares(octets["dares_defm_stock_france_cvs_trim.csv"])
    r, g2 = calculer(I, D)
    g += g2
    Lg, m8 = serie_longue(I, textes[D1])
    g.append(m8)
    m7, cit, N = verifier_citations(textes)
    g.append(m7)
    for m in autotest(I, D, octets, sommes, textes):
        log("autotest : la mutation a mordu : " + m)
    A, A_en = affichage(r, Lg, N), affichage_en(r, Lg, N)
    A.update(cit)
    A_en.update(TR)
    if set(A) != set(A_en):
        raise Arret("blocs affichage et affichage_en : cles differentes (%s)" % sorted(set(A) ^ set(A_en)))
    global LANG
    figs = {}
    for LANG, suf in (("fr", ""), ("en", "-en")):
        figs["mesurer-recoupement%s.svg" % suf] = fig_recoupement(r)
        figs["mesurer-recoupement%s-m.svg" % suf] = fig_recoupement_m(r)
        figs["mesurer-regle%s.svg" % suf] = fig_regle(r)
        figs["mesurer-regle%s-m.svg" % suf] = fig_regle_m(r)
        figs["mesurer-serie-longue%s.svg" % suf] = fig_serie_longue(Lg)
        figs["mesurer-serie-longue%s-m.svg" % suf] = fig_serie_longue_m(Lg)
    LANG = "fr"
    vues = parite(figs)
    # autotest de la parité : une valeur faussée dans une version téléphone doit arrêter
    fausse = dict(figs)
    fausse["mesurer-recoupement-m.svg"] = fausse["mesurer-recoupement-m.svg"].replace(">1\u202f427<", ">1\u202f472<")
    if fausse["mesurer-recoupement-m.svg"] == figs["mesurer-recoupement-m.svg"]:
        raise Arret("autotest parite : la valeur a fausser n'est pas dessinee, mutation sans objet")
    try:
        parite(fausse)
        raise Arret("mutation parite : aucune garde n'a mordu, controle ABSENT")
    except Arret as e:
        if not str(e).startswith("parite telephone"):
            raise
        log("autotest : la mutation a mordu : valeur faussee en version telephone -> " + str(e)[:70])
    g.append("parite : %d versions telephone conformes a leur figure detaillee" % len(vues))
    if check:
        log("--check : %d gardes tenues, %d cles d'affichage, rien ecrit (archive comprise)." % (len(g), len(A)))
        return 0
    try:
        import cairosvg
    except ImportError:
        raise Arret("cairosvg absent : SVG et PNG ensemble ou pas du tout")
    fi = fiches(figs, A, A_en)
    csvt = csv_texte(r, Lg)
    lib_fin_fr, lib_fin_en = A["dares_fin"], A_en["dares_fin"]
    payload = {
        "releve_le": releve,
        "_licence": "CC BY 4.0 — compilation Stéphane Lalut ; sources Insee (Insee Références 2026), Dares-France Travail",
        "meta": {"page": "https://" + PAGE_URL, "protocole": "fiches de preuve v2, volet Mesurer, E1 et E2 ; phase B du 12/10/2026 (E3, série longue)",
                 "sources": {"insee": {"url": URL_XLSX, "sha256": sommes["ECRT2026-D1.xlsx"]},
                             "dares": {"url": URL_DARES, "sha256": sommes["dares_defm_stock_france_cvs_trim.csv"],
                                       "metadonnees": URL_DARES_META, "mise_a_jour": maj_dares},
                             "releve_fige": "%s (scripts/sources_mesurer/SHA256SUMS)" % releve},
                 "definitions": {
                     "chomeur_bit": "« %s » (Insee, ECRT 2026, dossier 1)" % cit["chomeur_bit"],
                     "categorie_a": "« %s » (Insee, ibid.)" % cit["categorie_a"],
                     "estimation": "appariement du Fichier historique statistique de France Travail et de l'enquête Emploi de l'Insee",
                     "series_dares": "moyenne trimestrielle des inscrits en fin de mois, France hors Mayotte, CVS-CJO"},
                 "ruptures": [
                     {"date": "2019-11/2021-12", "objet": "réforme de l'assurance chômage", "touche": "inscription (Insee)"},
                     {"date": "2021", "objet": "enquête Emploi rénovée, pondérations rétropolées", "touche": "chômage BIT (Insee)"},
                     {"date": "2022-01/2022-11", "objet": "nouveau formulaire d'actualisation", "touche": "catégorie A, -%d (Dares)" % N["rupture_2022"]},
                     {"date": "2024-10/2024-11", "objet": "évolution du formulaire d'actualisation", "touche": "catégorie A, +%d environ au 4e trimestre 2024 (Dares-France Travail, Dreets Normandie)" % N["formulaire_2024"]},
                     {"date": "2025-01", "objet": "loi pour le plein emploi : inscription d'office", "touche": "catégorie A et série hors RSA (Dares)"},
                     {"date": "2025-01/2025-03", "objet": "actualisation automatique des nouveaux inscrits non indemnisables, restreinte en avril", "touche": "catégorie A et série hors RSA (Dares)"},
                     {"date": "2025-03/2025-04", "objet": "incident : %d personnes comptées à tort en A, reclassées en G" % N["incident_2025"], "touche": "catégorie A (Dares)"},
                     {"date": "2025-06", "objet": "décret relatif aux sanctions", "touche": "radiations (Dares, Insee)"}]},
        "gardes": g,
        "calcul": {k: v for k, v in r.items() if k not in ("dares_a", "dares_a_hors")},
        "series": {"inscrits_A": r["dares_a"], "inscrits_A_hors_RSA_et_jeunes_en_parcours": r["dares_a_hors"]},
        "series_longues_insee": {"unite": "millions de personnes, moyenne annuelle", "chomeurs_bit": Lg["bit"],
                                 "inscrits_A": Lg["a"], "ecart": Lg["ecart"]},
        # « Citer » : « relevé du <date> » = date à laquelle les sources ont été relues et figées (défaut D16 : c'était la
        # date de mise à jour de la Dares, non celle de la lecture)
        "affichage": dict(A, releve_le="%s (séries Dares mises à jour le %s, jusqu'au %s ; Insee, édition du 02/07/2026)"
                          % ("/".join(reversed(releve.split("-"))), "/".join(reversed(maj_dares.split("-"))), lib_fin_fr)),
        "affichage_en": dict(A_en, releve_le="%s (Dares series updated on %s, up to the %s; Insee, edition of 2 July 2026)"
                             % (releve, maj_dares, lib_fin_en)),
    }
    txt_json = json.dumps(payload, ensure_ascii=False, indent=1, default=str)
    same = (OUT_DATA.exists() and OUT_DATA.read_text(encoding="utf-8") == txt_json
            and OUT_STATIC.exists() and OUT_STATIC.read_text(encoding="utf-8") == txt_json
            and OUT_CSV.exists() and OUT_CSV.read_text(encoding="utf-8-sig") == csvt
            and all((OUT_IMG / n).exists() and (OUT_IMG / n).read_text(encoding="utf-8") == s for n, s in figs.items())
            and all((OUT_IMG / n.replace(".svg", ".png")).exists() for n in figs)
            and OUT_FIGURES.exists() and json.loads(OUT_FIGURES.read_text(encoding="utf-8")) == fi)
    if same:
        log("Donnees et figures identiques : rien ecrit.")
        return 0
    OUT_DATA.write_text(txt_json, encoding="utf-8")
    OUT_STATIC.write_text(txt_json, encoding="utf-8")
    OUT_CSV.write_text(csvt, encoding="utf-8-sig", newline="\n")
    for n, s in figs.items():
        (OUT_IMG / n).write_text(s, encoding="utf-8")
        cairosvg.svg2png(url=str(OUT_IMG / n), write_to=str(OUT_IMG / n.replace(".svg", ".png")),
                         output_width=1440 if not n.endswith("-m.svg") else 900, background_color="white")
    OUT_FIGURES.write_text(json.dumps(fi, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    for x in g:
        log("OK " + x)
    log("Ecrit : data/promesses_mesurer.json, static/promesses_mesurer.{json,csv}, data/figures_mesurer.json, %d figures SVG + PNG (FR, EN, telephone)" % len(figs))
    return 0


if __name__ == "__main__":
    sys.exit(main())
